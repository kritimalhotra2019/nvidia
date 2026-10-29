import asyncio
import json
from langchain_core.messages import HumanMessage, ToolMessage
from squipi_nemo.runner import compact_context, run_job
from squipi_nemo.supervisor import supervise
from squipi_nemo.state import Job
from test_trial import task, clean, ModelFixture, server_for, run_cli
import subprocess
import sys
import threading


def test_automatic_continuation_keeps_work(task):
    attempts = []
    async def attempt():
        job = Job(task)
        try:
            attempts.append(True)
            if len(attempts) == 1:
                clean(job)
                job.save(dict(job.load(), status='paused', pause_reason='request_timeout'))
            else:
                job.save(dict(job.load(), status='completed'))
            return job.status()
        finally:
            job.close()
    result = asyncio.run(supervise(task, attempt, backoff=0))
    assert result['status'] == 'completed'
    assert result['completed_tools'].count('clean_dataset') == 1
    assert len(attempts) == 2


def test_no_progress_stops_boundedly(task):
    calls = []
    async def attempt():
        calls.append(1)
        job = Job(task)
        job.save(dict(job.load(), status='paused', pause_reason='request_timeout'))
        result = job.status()
        job.close()
        return result
    result = asyncio.run(supervise(task, attempt, cycles=10, max_stalls=2, backoff=0))
    assert len(calls) == 2 and result['pause_reason'] == 'no_progress'


def test_explicit_stop_interrupts_model_and_does_not_resume(task):
    calls = []
    class Model:
        async def ainvoke(self, messages):
            calls.append(1)
            await asyncio.sleep(30)
    async def exercise():
        async def attempt():
            return await run_job(task, Model(), {}, request_seconds=20)
        async def stop():
            await asyncio.sleep(.05)
            (task/'stop.requested').touch()
        stop_task = asyncio.create_task(stop())
        result = await supervise(task, attempt, backoff=0)
        await stop_task
        return result
    result = asyncio.run(exercise())
    assert calls == [1] and result['pause_reason'] == 'user_stop'


def test_compaction_retains_all_user_answers_and_result_references():
    body = json.dumps({'_squipi_step_key': 'abc', 'dataset_id': 'cleaned', 'detail': 'X'*5000})
    user = HumanMessage(content='Keep schooling and use the saved green template.')
    messages = [user, ToolMessage(content=body, tool_call_id='call_1')] + [HumanMessage(content='continue')]*5
    shorter = compact_context(messages)
    assert shorter[0] is user
    assert len(shorter[1].content) < 500 and 'abc' in shorter[1].content and 'cleaned' in shorter[1].content
    assert messages[1].content == body


def test_real_nat_work_continues_across_step_budgets(task):
    fixture = ModelFixture(task)
    server = server_for(fixture)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        command = [sys.executable, '-m', 'squipi_nemo.cli', 'work', str(task), '--model', 'controlled-fixture',
                   '--base-url', f'http://127.0.0.1:{server.server_port}/v1', '--steps', '1',
                   '--cycles', '6', '--seconds', '30', '--request-seconds', '10']
        process = subprocess.run(command, capture_output=True, text=True, timeout=80)
        assert process.returncode == 0, process.stderr
        result = json.loads(process.stdout)
        assert result['status'] == 'completed', result
        assert result['completed_tools'].count('clean_dataset') == 1
        assert result['attempts'] >= 4
    finally:
        server.shutdown()
        server.server_close()


def test_real_template_export_waits_for_visual_review(task):
    from langchain_core.messages import AIMessage
    from pptx import Presentation
    manifest=json.loads((task/'task.json').read_text())
    manifest['required_outputs']=['pptx']
    (task/'task.json').write_text(json.dumps(manifest))
    job=Job(task)
    started=job.call('begin_template_presentation',{'template_id':manifest['templates']['pptx']['id'],'slide_numbers':[1]})
    deck=Presentation(started['draft_path'])
    edits=[{'shape_id':s.shape_id,'text':'Controlled test'} for s in deck.slides[0].shapes if s.has_text_frame]
    job.call('edit_presentation_slide',{'draft_id':started['draft_id'],'revision':0,'slide_number':1,'text_edits':edits})
    exported=job.call('finalize_template_presentation',{'draft_id':started['draft_id'],'revision':1,'title':'Review test'})
    job.close()
    class Model:
        async def ainvoke(self,messages):
            return AIMessage(content='The file is saved; it needs visual review.')
    async def attempt():
        return await run_job(task,Model(),{})
    result=asyncio.run(supervise(task,attempt,backoff=0))
    assert result['status']=='needs_review' and result['pause_reason']=='visual_review'
    assert result['pending_visual_reviews']==[exported['path']]
    assert result['attempts']==1


def test_amendment_replans_unexecuted_calls(task):
    from langchain_core.messages import AIMessage, messages_to_dict
    job = Job(task)
    saved = job.load()
    saved['messages'] = messages_to_dict([HumanMessage(content='Create a report.'), AIMessage(content='', tool_calls=[{'id':'old_export', 'name':'export_document', 'args':{'format':'docx'}}])])
    job.save(saved);job.close()
    observed=[]
    class Model:
        async def ainvoke(self,messages):
            observed.extend(messages)
            raise TimeoutError('test pause')
    result=asyncio.run(run_job(task,Model(),{},message='Do not make the Word file yet.',request_seconds=2))
    assert not result['completed_tools']
    assert any(isinstance(m,ToolMessage) and m.tool_call_id=='old_export' and 'Skipped' in m.content for m in observed)
    assert any(isinstance(m,HumanMessage) and m.content=='Do not make the Word file yet.' for m in observed)
