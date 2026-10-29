"""Exercise real NAT composition, real worker writes, and process-boundary recovery.

The controlled HTTP model is labelled as such: it tests orchestration, not reasoning.
"""
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import subprocess
import sys
import threading
import time
import zipfile

import pandas as pd
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from squipi_nemo.cli import configuration, initialize
from squipi_nemo.state import Job, NeedsReview, digest, exclusive_job
from squipi_nemo.runner import progress_message
from squipi_worker.registry import InvalidArguments, registry


def seed(directory):
    directory.mkdir(parents=True, exist_ok=True)
    data = directory / 'households.xlsx'
    pd.DataFrame([{'name': f'Synthetic Person {i}', 'email': f'person{i}@example.test',
                   'income': 2000 + 100 * i, 'gender': ['Female', 'Male'][i % 2],
                   'schooling_years': 6 + i % 10, 'wellbeing': 3 + i % 7}
                  for i in range(40)]).to_excel(data, index=False)
    template = directory / 'reference.pptx'
    deck = Presentation()
    deck.slide_width = Inches(12)
    deck.slide_height = Inches(6.75)
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = 'OLD EXAMPLE CONTENT'
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('F7F6F2')
    deck.save(template)
    return data, template


@pytest.fixture
def task(tmp_path):
    data, template = seed(tmp_path / 'sources')
    root = tmp_path / 'task'
    initialize(root, 'Remove name and email PII, save cleaned Excel, then analyse the cleaned data '
               'and produce Word and PowerPoint reports with labelled graphs using my presentation template.',
               [data], [template], ['xlsx', 'docx', 'pptx'])
    return root


def clean(job):
    imported = job.call('import_dataset', {'path': job.manifest['inputs'][0]['path']})
    return job.call('clean_dataset', {'dataset_id': imported['dataset_id'], 'pii': 'drop',
                                     'pii_columns': ['name', 'email'], 'steps': []})


def test_successful_writes_reused_across_restart(task):
    job = Job(task)
    result = clean(job)
    args = {'dataset_id': result['dataset_id']}
    saved = job.call('export_excel', args)
    job.close()
    resumed = Job(task)
    assert resumed.call('export_excel', args) == saved
    assert len(resumed.rows()) == 3
    frame = pd.read_excel(saved['path'])
    assert set(frame.columns) == {'income', 'gender', 'schooling_years', 'wellbeing'}
    assert len(frame) == 40 and frame['income'].mean() == 3950
    resumed.close()


def test_modified_artifact_blocks_resume(task):
    job = Job(task)
    result = clean(job)
    Path(result['path']).write_text('changed')
    with pytest.raises(NeedsReview, match='missing or changed'):
        job.check_integrity()
    job.close()


def test_uncertain_write_is_not_repeated(task):
    job = Job(task)
    with job.db:
        job.db.execute('INSERT INTO calls VALUES (?, ?, ?, ?, NULL, NULL, ?)',
                       ('uncertain', 'export_excel', '{}', 'running', time.time()))
    with pytest.raises(NeedsReview, match='Interrupted during'):
        job.check_integrity()
    job.close()


def test_file_and_network_boundaries(task, tmp_path):
    job = Job(task)
    outside = tmp_path / 'unattached.csv'
    outside.write_text('a\n1\n')
    with pytest.raises(PermissionError):
        job.call('import_dataset', {'path': str(outside)})
    with pytest.raises(ValueError, match='outside the local trial'):
        job.call('fetch_web', {'url': 'https://example.com'})
    with pytest.raises(ValueError, match='loopback'):
        configuration(task, 'model', 'https://api.example.com/v1', 10, 10)
    job.close()


def test_single_writer(task):
    with exclusive_job(task):
        with pytest.raises(RuntimeError, match='already running'):
            with exclusive_job(task):
                pass


def test_correctable_outcome_error_preserves_cleaning(task):
    job = Job(task)
    cleaned = clean(job)
    args = {'dataset_id': cleaned['dataset_id'], 'format': 'docx', 'outcome': 'income'}
    with pytest.raises(InvalidArguments, match='not a numeric outcome'):
        job.call('report_dataset', args)
    job.check_integrity()
    job.call('set_column_roles', {'dataset_id': cleaned['dataset_id'], 'roles': {'income': 'numeric'}})
    result = job.call('report_dataset', args)
    assert Path(result['docx_path']).is_file()
    assert sum(row['tool'] == 'clean_dataset' for row in job.rows()) == 1
    job.close()


def test_validation_error_after_deliverable_write_stays_uncertain(task, monkeypatch):
    job = Job(task)
    cleaned = clean(job)
    def partial(ctx, args):
        (ctx.project.root / 'artifacts' / 'partial.docx').write_bytes(b'incomplete')
        raise InvalidArguments('late failure')
    monkeypatch.setattr(registry.get('export_document'), 'handler', partial)
    with pytest.raises(NeedsReview):
        job.call('export_document', {'format': 'docx', 'title': 'Test', 'sections': [{'heading': 'Test'}]})
    with pytest.raises(NeedsReview):
        job.check_integrity()
    job.close()


def test_original_dataset_rejected_for_downstream_report(task):
    job = Job(task)
    cleaned = clean(job)
    with pytest.raises(ValueError, match='saved cleaned dataset'):
        job.call('report_dataset', {'dataset_id': cleaned['from_dataset_id']})
    job.close()


def test_progress_brief_restored_from_disk(task):
    job = Job(task)
    cleaned = clean(job)
    saved = job.call('export_excel', {'dataset_id': cleaned['dataset_id']})
    job.close()
    resumed = Job(task)
    brief = progress_message(resumed).content
    assert cleaned['dataset_id'] in brief and saved['path'] in brief
    assert '"missing_deliverables": ["docx", "pptx"]' in brief
    resumed.close()


def test_excel_formulas_remain_literal(task):
    job = Job(task)
    source = job.project.root / 'formula.csv'
    source.write_text('comment\n=1+1\n')
    data = job.call('import_dataset', {'path': str(source)})
    result = job.call('export_excel', {'dataset_id': data['dataset_id']})
    from openpyxl import load_workbook
    cell = load_workbook(result['path']).active['A2']
    assert cell.value == '=1+1' and cell.data_type == 's'
    job.close()


class ModelFixture:
    def __init__(self, root, stall=False):
        self.root, self.stall = root, stall
        self.requests = []

    def respond(self, request):
        self.requests.append(request)
        messages = request['messages']
        names = {call['id']: call['function']['name'] for m in messages for call in m.get('tool_calls', [])}
        results = {names[m['tool_call_id']]: json.loads(m['content']) for m in messages if m['role'] == 'tool'}
        job = Job(self.root)
        manifest = job.manifest
        job.close()
        if 'import_dataset' not in results:
            return 'import_dataset', {'path': manifest['inputs'][0]['path']}
        if 'clean_dataset' not in results:
            return 'clean_dataset', {'dataset_id': results['import_dataset']['dataset_id'],
                                     'pii': 'drop', 'pii_columns': ['name', 'email'], 'steps': []}
        if 'export_excel' not in results:
            return 'export_excel', {'dataset_id': results['clean_dataset']['dataset_id']}
        if self.stall:
            time.sleep(4)
        if 'report_dataset' not in results:
            return 'report_dataset', {'dataset_id': results['clean_dataset']['dataset_id'],
                                      'format': 'all', 'outcome': 'wellbeing', 'title': 'Synthetic household analysis'}
        return None, None


def server_for(fixture):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def do_POST(self):
            request = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            name, args = fixture.respond(request)
            message = {'role': 'assistant', 'content': 'Completed the requested files.' if not name else ''}
            if name:
                message['tool_calls'] = [{'id': 'call_' + str(len(fixture.requests)), 'type': 'function',
                                          'function': {'name': name, 'arguments': json.dumps(args)}}]
            payload = json.dumps({'id': 'fixture', 'object': 'chat.completion', 'created': 1,
                                  'model': 'controlled-fixture', 'choices': [{'index': 0, 'message': message,
                                      'finish_reason': 'tool_calls' if name else 'stop'}],
                                  'usage': {'prompt_tokens': 1, 'completion_tokens': 1, 'total_tokens': 2}}).encode()
            try:
                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Content-Length', str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)
            except (BrokenPipeError, ConnectionResetError):
                pass
    return ThreadingHTTPServer(('127.0.0.1', 0), Handler)


def run_cli(root, server, steps=12, seconds=60, pause_after=None):
    command = [sys.executable, '-m', 'squipi_nemo.cli', 'run', str(root), '--model', 'controlled-fixture',
               '--base-url', f'http://127.0.0.1:{server.server_port}/v1', '--steps', str(steps), '--seconds', str(seconds)]
    if pause_after:
        command += ['--pause-after', pause_after]
    result = subprocess.run(command, capture_output=True, text=True, timeout=100)
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)


def test_real_nat_resume_after_process_exit(task):
    fixture = ModelFixture(task)
    server = server_for(fixture)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        first = run_cli(task, server, pause_after='export_excel')
        assert first['status'] == 'paused', first
        assert first['completed_tools'] == ['import_dataset', 'clean_dataset', 'export_excel']
        excel = next(path for path in first['outputs'] if path.endswith('.xlsx'))
        checksum = digest(excel)
        second = run_cli(task, server)
        assert second['status'] == 'completed', second
        assert second['attempts'] == 2
        assert second['completed_tools'] == first['completed_tools'] + ['report_dataset']
        assert digest(excel) == checksum
        assert {Path(path).suffix for path in second['outputs']} >= {'.xlsx', '.docx', '.pptx'}
        # Restarted model context contains original instruction, all tool decisions/results,
        # cleaned ID and saved template. It does not need to search old task summaries.
        replay = json.dumps(fixture.requests[3]['messages'])
        assert 'Remove name and email PII' in replay and 'export_excel' in replay
        job = Job(task)
        report = next(row for row in job.rows() if row['tool'] == 'report_dataset')
        assert json.loads(report['arguments'])['pptx_template_id'] == job.manifest['templates']['pptx']['id']
        cleaned = next(row for row in job.rows() if row['tool'] == 'clean_dataset')
        assert json.loads(report['arguments'])['dataset_id'] == json.loads(cleaned['result'])['dataset_id']
        job.close()
        ppt = next(path for path in second['outputs'] if path.endswith('.pptx'))
        assert Presentation(ppt).slide_width == Inches(12)
        with zipfile.ZipFile(ppt) as archive:
            text = '\n'.join(archive.read(name).decode() for name in archive.namelist() if name.endswith('.xml'))
            assert 'OLD EXAMPLE CONTENT' not in text and 'Synthetic Person' not in text
            assert any(name.startswith('ppt/media/') for name in archive.namelist())
    finally:
        server.shutdown()
        server.server_close()


def test_controlled_response_timeout_then_resume(task):
    fixture = ModelFixture(task, stall=True)
    server = server_for(fixture)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        first = run_cli(task, server, seconds=2)
        assert first['status'] == 'paused', first
        assert first['completed_tools'] == ['import_dataset', 'clean_dataset', 'export_excel']
        fixture.stall = False
        second = run_cli(task, server)
        assert second['status'] == 'completed', second
        assert second['completed_tools'].count('clean_dataset') == 1
    finally:
        server.shutdown()
        server.server_close()


def test_model_cannot_complete_with_only_a_promise(task):
    fixture = ModelFixture(task)
    fixture.respond = lambda request: (None, None)
    server = server_for(fixture)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        result = run_cli(task, server, steps=1)
        assert result['status'] == 'paused'
        assert result['missing_outputs'] == ['docx', 'pptx', 'xlsx']
        assert result['completed_tools'] == []
    finally:
        server.shutdown()
        server.server_close()
