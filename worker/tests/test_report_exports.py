"""Actual Office artifacts and selective cleaning, without system Quick Look."""
import json
from pathlib import Path

import docx
import pandas as pd
import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, registry, InvalidArguments
from squipi_worker.security import AllowList
from squipi_worker.tools import export, insights
from squipi_worker.tools.report import _descriptive_sections, _dictionary_units


def import_frame(tmp_path, frame):
    source = tmp_path / "households.xlsx"
    frame.to_excel(source, sheet_name="Households", index=False)
    allow = AllowList()
    allow.allow(source)
    project = Project(tmp_path / "project", allow)
    ctx = JobContext(project, lambda *args: None, lambda: False)
    ds = registry.call("import_dataset", {"path": str(source), "sheet": "Households"}, ctx)
    return ctx, ds["dataset_id"], source


def no_findings(*args, **kwargs):
    return {"findings": [], "data_quality": [], "limits": ["Descriptive results only."]}


def test_word_report_has_descriptives_and_embedded_labelled_graphs_without_findings(tmp_path, monkeypatch):
    frame = pd.DataFrame({
        "household_id": [f"HH-{i}" for i in range(6)],
        "monthly_income_inr": [10000, 20000, 30000, 40000, 50000, None],
        "wellbeing_index_0_100": [40, 50, 60, 70, 80, 90],
        "head_gender": ["Female", "Female", "Male", "Male", "Male", None],
        "street_address": [f"Synthetic address {i}" for i in range(6)],
    })
    ctx, ds, source = import_frame(tmp_path, frame)
    before = source.read_bytes()
    monkeypatch.setattr(insights, "find_insights", no_findings)
    run = export.subprocess.run
    def without_quicklook(command, *args, **kwargs):
        assert command[0] != "qlmanage", "Report should embed PNGs without Quick Look"
        return run(command, *args, **kwargs)
    monkeypatch.setattr(export.subprocess, "run", without_quicklook)
    result = registry.call("report_dataset", {"dataset_id": ds, "format": "docx", "title": "Household descriptives", "filename": "Household descriptive statistics.docx"}, ctx)
    assert Path(result["docx_path"]).name == "Household descriptive statistics.docx"
    assert result["findings"] == 0 and result["charts"] == 3
    assert result["exports"]["docx"] == {"charts_embedded": 3, "charts_referenced_only": [], "template_applied": None}
    assert result["warnings"] == [] and source.read_bytes() == before
    document = docx.Document(result["docx_path"])
    assert len(document.inline_shapes) == 3
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph._p.xpath(".//w:drawing"):
            assert document.paragraphs[index - 1].paragraph_format.keep_with_next, "Chart descriptions must stay with their image"
    tables = [[list(cell.text for cell in row.cells) for row in table.rows] for table in document.tables]
    numeric = next(t for t in tables if t[0] == ["Variable", "Valid N", "Missing", "Mean", "SD", "Median"])
    assert numeric[1] == ["monthly_income_inr", "5", "1", "30,000.00", "15,811.39", "30,000.00"]
    assert not any(row[0] == "household_id" for row in numeric[1:])
    categories = next(t for t in tables if t[0] == ["Category", "Count", "% of non-missing"])
    assert ["Male", "3", "60"] in categories and ["Female", "2", "40"] in categories
    assert ["Missing (excluded from percentages)", "1", "—"] in categories
    html = Path(result["report_path"]).read_text()
    assert "monthly income inr" in html and "wellbeing index 0 100" in html
    assert "<svg" in html and "Synthetic address" not in html
    assert "PII-free" not in result["summary_text"]


@pytest.mark.parametrize("fmt,default", [("docx", "report.docx"), ("pptx", "slides.pptx")])
@pytest.mark.parametrize("name", [None, "Named findings", "Named findings.UPPER"])
def test_office_export_filename_preserves_table_contents_and_defaults(tmp_path, fmt, default, name):
    from pptx import Presentation
    ctx = JobContext(Project(tmp_path / "project", AllowList()), lambda *args: None, lambda: False)
    args = {"format": fmt, "title": "Synthetic findings", "template_id": "none", "sections": [{
        "heading": "Observed counts", "table": {"columns": ["Measure", "Value"], "rows": [["Eligible", 10], ["Responded", 5]]},
    }]}
    if name is not None:
        args["filename"] = name.replace("UPPER", fmt.upper())
    result = registry.call("export_document", args, ctx)
    expected_name = default if name is None else (args["filename"] if Path(args["filename"]).suffix else args["filename"] + "." + fmt)
    assert Path(result["path"]).name == expected_name
    assert Path(result["path"]).parent == Path(result["artifact_dir"])
    assert expected_name in result["summary_text"] and Path(result["path"]).stat().st_size == result["size_bytes"]
    if fmt == "docx":
        tables = docx.Document(result["path"]).tables
    else:
        tables = [shape.table for slide in Presentation(result["path"]).slides for shape in slide.shapes if shape.has_table]
    actual = [[list(cell.text for cell in row.cells) for row in table.rows] for table in tables]
    assert actual == [[["Measure", "Value"], ["Eligible", "10"], ["Responded", "5"]]]
    assert len(list(Path(result["artifact_dir"]).glob("*." + fmt))) == (0 if expected_name.endswith(fmt.upper()) else 1)


def test_named_powerpoint_dataset_report_keeps_values_and_pii_policy(tmp_path, monkeypatch):
    from pptx import Presentation
    frame = pd.DataFrame({"income": [10, 20, 30], "respondent_email": ["a@example.test", "b@example.test", "c@example.test"]})
    ctx, ds, source = import_frame(tmp_path, frame)
    before = source.read_bytes()
    monkeypatch.setattr(insights, "find_insights", no_findings)
    result = registry.call("report_dataset", {"dataset_id": ds, "format": "pptx", "filename": "Income findings", "title": "Income findings"}, ctx)
    assert Path(result["pptx_path"]).name == "Income findings.pptx"
    tables = [shape.table for slide in Presentation(result["pptx_path"]).slides for shape in slide.shapes if shape.has_table]
    table_rows = [[list(cell.text for cell in row.cells) for row in table.rows] for table in tables]
    numeric = next(rows for rows in table_rows if rows[0] == ["Variable", "Valid N", "Missing", "Mean", "SD", "Median"])
    assert numeric[1] == ["income", "3", "0", "20.00", "10.00", "20.00"]
    assert all(email not in json.dumps(table_rows) for email in frame["respondent_email"])
    assert source.read_bytes() == before


@pytest.mark.parametrize("tool_name", ["export_document", "report_dataset"])
@pytest.mark.parametrize("filename", ["../outside.docx", "/tmp/outside.docx", "subdir/outside.docx", "subdir\\outside.docx",
    "C:outside.docx", ".", "..", " ", "line\nbreak.docx", "null\x00.docx", "control\x7f.docx", "control\x85.docx", "wrong.pptx", "wrong.xlsx"])
def test_office_export_rejects_unsafe_or_conflicting_filenames_before_writing(tmp_path, tool_name, filename):
    ctx = JobContext(Project(tmp_path / "project", AllowList()), lambda *args: None, lambda: False)
    args = {"format": "docx", "filename": filename}
    if tool_name == "export_document":
        args.update(title="Test", sections=[{"heading": "Test", "text": "Synthetic."}])
    else:
        args["dataset_id"] = "not-looked-up-before-filename-validation"
    with pytest.raises(InvalidArguments):
        registry.call(tool_name, args, ctx)
    assert not list(ctx.project.root.glob("artifacts/*"))


@pytest.mark.parametrize("fmt", [None, "html", "all"])
def test_dataset_report_filename_requires_one_office_format(tmp_path, fmt):
    ctx = JobContext(Project(tmp_path / "project", AllowList()), lambda *args: None, lambda: False)
    args = {"dataset_id": "not-looked-up-before-filename-validation", "filename": "Named report"}
    if fmt:
        args["format"] = fmt
    with pytest.raises(InvalidArguments, match="separate exports"):
        registry.call("report_dataset", args, ctx)
    assert not list(ctx.project.root.glob("artifacts/*"))


def test_numeric_descriptives_report_single_value_sd_as_undefined():
    sections = _descriptive_sections({"rows_profiled": 3, "columns": [{
        "name": "income", "role": "numeric", "missing": 2,
        "histogram": {"counts": [1]},
        "stats": {"mean": 12345.67, "median": 12345.67, "std": 0, "min": 12345.67, "max": 12345.67},
    }]})
    assert sections[0]["table"]["rows"] == [["income", 1, 2, "12,345.67", None, "12,345.67"]]


def test_report_exposes_exported_graph_failure(tmp_path, monkeypatch):
    ctx, ds, _ = import_frame(tmp_path, pd.DataFrame({"income": [10, 20, 30, 40]}))
    monkeypatch.setattr(insights, "find_insights", no_findings)
    monkeypatch.setattr(export, "export_document", lambda *a, **k: {
        "path": str(tmp_path / "report.docx"), "charts_embedded": 0, "charts_referenced_only": ["unrendered.svg"], "template_applied": None,
    })
    result = registry.call("report_dataset", {"dataset_id": ds, "format": "docx"}, ctx)
    assert result["exports"]["docx"]["charts_referenced_only"] == ["unrendered.svg"]
    assert len(result["warnings"]) == 1 and "1 chart(s)" in result["warnings"][0]
    assert "incomplete" in result["summary_text"]


def test_selective_removal_preserves_every_other_column_cell_and_duplicate(tmp_path):
    frame = pd.DataFrame({
        "household_id": ["HH-01", "HH-02", "HH-02"],
        "respondent_full_name": ["Synthetic One", "Synthetic Two", "Synthetic Two"],
        "head_full_name": ["Head One", "Head Two", "Head Two"],
        "respondent_email": ["one@example.test", "two@example.test", "two@example.test"],
        "respondent_phone": ["+1 202 555 0101", "+1 202 555 0102", "+1 202 555 0102"],
        "national_id": ["TEST-001", "TEST-002", "TEST-002"],
        "street_address": ["  Synthetic 1 Road  ", "Synthetic 2 Road", "Synthetic 2 Road"],
        "respondent_dob": ["01/02/1980", "04/05/1990", "04/05/1990"],
        "interview_notes": ["Contact one@example.test about TEST-001", "Call +1 202 555 0102", "Call +1 202 555 0102"],
        "Unchanged Header": ["  YES ", "yes", "yes"],
        "income_inr": [12345.67, None, None],
    })
    ctx, ds, source = import_frame(tmp_path, frame)
    original_bytes = source.read_bytes()
    columns = ["respondent_full_name", "head_full_name", "respondent_email", "respondent_phone"]
    result = registry.call("clean_dataset", {"dataset_id": ds, "steps": [], "pii": "keep", "drop_columns": columns}, ctx)
    actual = pd.read_csv(result["path"])
    pd.testing.assert_frame_equal(actual, frame.drop(columns=columns), check_dtype=False)
    assert result["cols_after"] == len(frame.columns) - 4 and result["rows_after"] == 3
    assert [change["step"] for change in result["changes"]] == ["drop_columns"]
    changes = json.loads((Path(result["artifact_dir"]) / "changes.json").read_text())
    assert changes["steps"] == [] and changes["pii"] == "keep" and source.read_bytes() == original_bytes


def test_omitted_cleaning_steps_still_apply_defaults(tmp_path):
    ctx, ds, _ = import_frame(tmp_path, pd.DataFrame({"Messy Header": ["  A ", "  A "]}))
    result = registry.call("clean_dataset", {"dataset_id": ds}, ctx)
    changes = {change["step"] for change in result["changes"]}
    assert "normalize_names" in changes and "trim_whitespace" in changes and "drop_duplicates" in changes


def test_dictionary_units_are_source_labels_and_survive_a_cleaned_copy(tmp_path):
    ctx, ds, source = import_frame(tmp_path, pd.DataFrame({"income_inr": [100, 200], "respondent_email": ["a@example.test", "b@example.test"]}))
    with pd.ExcelWriter(source, engine="openpyxl", mode="a") as writer:
        pd.DataFrame({"Variable": ["income_inr"], "Type / unit": ["Number, INR/month"]}).to_excel(writer, sheet_name="Data Dictionary", index=False)
    result = registry.call("clean_dataset", {"dataset_id": ds, "steps": [], "pii": "keep", "drop_columns": ["respondent_email"]}, ctx)
    assert _dictionary_units(ctx, ctx.project.dataset(result["dataset_id"])) == {"income_inr": "Number, INR/month"}


def test_word_layout_keeps_long_labels_readable_and_tables_together(tmp_path):
    from docx.oxml.ns import qn

    allow = AllowList()
    ctx = JobContext(Project(tmp_path / "project", allow), lambda *args: None, lambda: False)
    result = registry.call("export_document", {
        "format": "docx", "title": "Household descriptive report", "subtitle": "Synthetic fixture",
        "sections": [
            {"heading": "Numeric descriptive statistics", "table": {
                "columns": ["Variable", "Valid N", "Missing", "Mean", "SD", "Median"],
                "rows": [["monthly_household_income_inr", 499, 1, "30,000.00", "15,811.39", "30,000.00"]],
            }},
            {"heading": "Column guide", "table": {
                "columns": ["column", "role", "missing", "summary"],
                "rows": [["monthly_household_income_inr", "numeric", "0.2%", "Monthly household income in INR with one observation missing."]],
            }},
        ],
    }, ctx)
    document = docx.Document(result["path"])
    for style_name in ["Title", "Subtitle", "Header", "Heading 1"]:
        style = document.styles[style_name]
        assert str(style.font.color.rgb) == "000000"
        assert not style.element.xpath("./w:pPr/w:pBdr")
        color = style.element.xpath("./w:rPr/w:color")[0]
        assert color.get(qn("w:themeColor")) is None
    numeric, guide = document.tables
    widths = [column.width.twips for column in numeric.columns]
    assert widths[0] == 2880 and min(widths[1:]) >= 1100  # two-inch labels, usable numeric columns
    assert guide.columns[3].width > guide.columns[0].width > guide.columns[1].width
    assert guide.columns[1].width.twips >= 1440  # the word 'categorical' must fit without a dangling letter
    assert document.paragraphs[-1].paragraph_format.space_before.pt >= 6
    for table in document.tables:
        assert table.autofit is False
        assert table._tbl.getprevious().xpath("./w:pPr/w:keepNext"), "A table's heading or description must stay with its first row"
        borders = table._tbl.xpath("./w:tblPr/w:tblBorders/*")
        assert len(borders) == 6 and all(border.get(qn("w:color")) == "D9D9D9" for border in borders)
        assert len(table._tbl.xpath("./w:tblPr/w:tblCellMar/*")) == 4
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        for row in table.rows:
            assert row._tr.xpath("./w:trPr/w:cantSplit")
            assert not row._tr.xpath("./w:trPr/w:trHeight")  # content can wrap without clipping
            for index, cell in enumerate(row.cells):
                assert cell.width == table.columns[index].width
                assert cell._tc.xpath("./w:tcPr/w:vAlign")[0].get(qn("w:val")) == "center"
        for cell in table.rows[0].cells:
            assert cell._tc.xpath("./w:tcPr/w:shd")[0].get(qn("w:fill")) == "F2F2F2"
