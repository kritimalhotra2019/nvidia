"""Regression coverage for complete, bounded reading of persisted document snapshots."""
import json
from pathlib import Path
from unittest.mock import patch

import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import InvalidArguments, JobContext, registry
from squipi_worker.security import AccessDenied, AllowList
from squipi_worker.tools import documents


def make_pdf(path, count=30, blank_pages=()):
    """Create an ordinary multi-page text PDF without optional rendering dependencies."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    for number in range(1, count + 1):
        page = writer.add_blank_page(width=612, height=792)
        if number in blank_pages:
            continue
        lines = [f"PAGE_{number:02d}_EVIDENCE: household research section {number}."]
        lines += [f"Survey row {row:02d}: income, schooling and wellbeing were measured independently."
                  for row in range(1, 41)]
        if number == count:
            lines += ["FINAL_PAGE_SECRET: the final recommendation is add two evening clinics."]
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 9 Tf 40 750 Td 14 TL\n" + "\n".join(f"({line}) Tj T*" for line in lines) + "\nET").encode("ascii"))
        page[NameObject("/Contents")] = stream
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    writer.write(str(path))
    return path


@pytest.fixture
def project(tmp_path):
    return Project(tmp_path / "project", AllowList())


def call(project, name="read_document", **args):
    return registry.call(name, args, JobContext(project, lambda *_: None, lambda: False))


def read_all(project, first):
    chunks = [first]
    assert len(json.dumps(first, ensure_ascii=True)) <= documents.MAX_DOCUMENT_RESULT_CHARS
    while chunks[-1]["next_offset"] is not None:
        previous = chunks[-1]
        chunk = call(project, **previous["next_call"]["arguments"])
        assert chunk["offset"] == previous["end_offset"]
        assert chunk["snapshot_id"] == first["snapshot_id"]
        assert chunk["returned_chars"] == len(chunk["text"])
        assert len(chunk["text"]) <= documents.MAX_CHARS_LIMIT
        assert len(json.dumps(chunk, ensure_ascii=True)) <= documents.MAX_DOCUMENT_RESULT_CHARS
        chunks.append(chunk)
        assert len(chunks) < 2000
    return "".join(chunk["text"] for chunk in chunks), chunks


def test_thirty_page_pdf_reads_final_page_without_reextracting(project, tmp_path):
    path = make_pdf(tmp_path / "thirty-pages.pdf")
    project.allow.allow(path)
    with patch.object(documents, "_pdf", wraps=documents._pdf) as extract:
        first = call(project, path=str(path))
        assert first["pages"] == 30 and first["coverage"]["total_units"] == 30
        assert not first["complete"] and "FINAL_PAGE_SECRET" not in first["text"]
        full, chunks = read_all(project, first)
        assert extract.call_count == 1
    assert len(chunks) > 6
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert all(f"PAGE_{page:02d}_EVIDENCE" in full for page in range(1, 31))
    assert "add two evening clinics" in full
    assert chunks[-1]["coverage"]["last_unit"] == 30
    assert chunks[-1]["complete"] and not chunks[-1]["has_more"]
    assert chunks[-1]["end_offset"] == first["total_chars"]
    found = call(project, "search_document", document_id=first["document_id"], query="FINAL_PAGE_SECRET")
    assert found["matches"][0]["unit_index"] == 30
    assert found["matches"][0]["offset"] > first["end_offset"]
    assert "two evening clinics" in found["matches"][0]["text"]


def test_thirty_slide_pptx_reads_late_text_tables_and_notes_without_reextracting(project, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "thirty-slides.pptx"
    presentation = Presentation()
    for number in range(1, 31):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"SLIDE_{number:02d}_EVIDENCE: household research"
        slide.placeholders[1].text = "\n".join(
            f"Survey observation {row:02d}: income, schooling and wellbeing were measured independently."
            for row in range(1, 41))
        if number == 30:
            slide.placeholders[1].text += "\nFINAL_SLIDE_EVIDENCE: add two evening clinics in the final recommendation."
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(7), Inches(1)).table
            table.cell(0, 0).text, table.cell(0, 1).text = "Indicator", "Verified count"
            table.cell(1, 0).text, table.cell(1, 1).text = "ELIGIBLE_HOUSEHOLDS_30", "500"
            slide.notes_slide.notes_text_frame.text = "SPEAKER_NOTES_TAIL: confirm the final-slide household eligibility count."
    presentation.save(path)
    project.allow.allow(path)
    with patch.object(documents, "_pptx", wraps=documents._pptx) as extract:
        first = call(project, path=str(path))
        assert first["kind"] == "pptx" and first["unit"] == "slides"
        assert first["pages"] == 30 and first["coverage"]["total_units"] == 30
        assert "FINAL_SLIDE_EVIDENCE" not in first["text"]
        full, chunks = read_all(project, first)
        assert extract.call_count == 1
    assert len(chunks) > 6
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert all(f"SLIDE_{number:02d}_EVIDENCE" in full for number in range(1, 31))
    assert "FINAL_SLIDE_EVIDENCE: add two evening clinics" in full
    assert "ELIGIBLE_HOUSEHOLDS_30 | 500" in full
    assert "SPEAKER_NOTES_TAIL" in full
    assert chunks[-1]["coverage"]["last_unit"] == 30
    assert chunks[-1]["structure"][-1]["tables"] == 1
    assert chunks[-1]["complete"] and chunks[-1]["next_offset"] is None
    assert chunks[-1]["end_offset"] == first["total_chars"]
    for query in ("FINAL_SLIDE_EVIDENCE", "ELIGIBLE_HOUSEHOLDS_30", "SPEAKER_NOTES_TAIL"):
        found = call(project, "search_document", document_id=first["document_id"], query=query)
        assert found["matches"][0]["unit_index"] == 30
        assert found["matches"][0]["offset"] > first["end_offset"]


def test_docx_long_sections_and_tables_are_continued_without_missing_tail(project, tmp_path):
    import docx

    path = tmp_path / "long.docx"
    doc = docx.Document()
    doc.add_heading("A long section", level=1)
    doc.add_paragraph("Research findings about education and household resources. " * 700)
    doc.add_heading("Final evidence", level=1)
    doc.add_paragraph("The late result is TAIL_DOCX_42.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Households", "500"
    doc.save(path)
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=3333)
    full, chunks = read_all(project, first)
    assert first["unit"] == "sections" and first["pages"] == 2
    assert len(chunks) > 10 and "TAIL_DOCX_42" in full
    assert "Households | 500" in full
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert chunks[-1]["structure"][-1]["tables"] == 1


def test_unicode_offsets_are_codepoints_and_reassemble_exactly(project, tmp_path):
    path = tmp_path / "unicode.txt"
    path.write_text(("नाम e\u0301 👨‍👩‍👧‍👦 café 東京\n" * 1000) + "END_UNICODE", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=137)
    full, chunks = read_all(project, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert first["offset_unit"] == "unicode_code_points"
    assert sum(chunk["returned_chars"] for chunk in chunks) == len(full)
    found = call(project, "search_document", document_id=first["document_id"], query="END_UNICODE")
    assert found["matches"][0]["offset"] == full.index("END_UNICODE")


def test_cache_survives_restart_and_original_deletion(project, tmp_path):
    path = tmp_path / "durable.txt"
    path.write_text("saved evidence " * 3000, "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    again = call(project, path=str(path))
    assert again["cached"] and again["text_path"] == first["text_path"]
    path.unlink()
    restarted = Project(project.root, AllowList())
    with patch.object(documents, "_text", side_effect=AssertionError("must not re-extract")):
        full, _ = read_all(restarted, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert len(list((project.root / "artifacts").iterdir())) == 1


def test_changed_source_refreshes_and_old_snapshot_cannot_continue(project, tmp_path):
    path = tmp_path / "changed.txt"
    path.write_text("old text " * 3000, "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    path.write_text("new document " * 3000, "utf-8")
    with pytest.raises(InvalidArguments, match="unchanged snapshot"):
        call(project, path=str(path), snapshot_id=first["snapshot_id"], offset=first["next_offset"])
    refreshed = call(project, path=str(path))
    assert refreshed["document_id"] == first["document_id"]
    assert refreshed["snapshot_id"] != first["snapshot_id"]
    with pytest.raises(InvalidArguments, match="snapshot changed"):
        call(project, **first["next_call"]["arguments"])


@pytest.mark.parametrize("invalid", [{}, {"path": "a", "document_id": "b"}, {"document_id": "unknown"},
                                     {"document_id": "unknown", "offset": -1}, {"document_id": "unknown", "offset": True}])
def test_invalid_cursor_arguments_are_rejected(project, invalid):
    with pytest.raises(InvalidArguments):
        call(project, **invalid)


def test_offset_bounds_and_eof_do_not_claim_prior_coverage(project, tmp_path):
    path = tmp_path / "short.txt"
    path.write_text("short text", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    with pytest.raises(InvalidArguments, match="offset must"):
        call(project, document_id=first["document_id"], offset=first["total_chars"] + 1)
    eof = call(project, document_id=first["document_id"], offset=first["total_chars"])
    assert eof["complete"] and eof["text"] == "" and eof["returned_chars"] == 0
    assert eof["coverage"]["first_unit"] is None


def test_snapshot_never_reads_an_unregistered_or_outside_path(project, tmp_path):
    path = tmp_path / "denied.txt"
    path.write_text("private evidence", "utf-8")
    with pytest.raises(AccessDenied):
        call(project, path=str(path))
    project.allow.allow(path)
    first = call(project, path=str(path))
    project.document(first["document_id"]).text_path = str(path)
    with pytest.raises(AccessDenied):
        call(project, document_id=first["document_id"])


def test_corrupt_snapshot_is_detected_and_path_can_reextract(project, tmp_path):
    path = tmp_path / "integrity.txt"
    path.write_text("trusted source text", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    Path(first["text_path"]).write_text("tampered cache", "utf-8")
    with pytest.raises(InvalidArguments, match="snapshot is damaged"):
        call(project, document_id=first["document_id"])
    refreshed = call(project, path=str(path))
    assert "trusted source text" in refreshed["text"] and not refreshed["cached"]


def test_legacy_registered_snapshot_upgrades_without_original(project, tmp_path):
    path = tmp_path / "legacy.md"
    path.write_text("# First\n\n" + "old saved text " * 1000 + "\n# Last\n\nLegacy tail", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    (Path(first["artifact_dir"]) / "manifest.json").unlink()
    path.unlink()
    continued = call(project, **first["next_call"]["arguments"])
    assert continued["cached"] and continued["snapshot_id"] == first["snapshot_id"]
    assert "Legacy text snapshot" in continued["metadata"]["note"]


def test_search_is_literal_bounded_and_paginates_matches(project, tmp_path):
    path = tmp_path / "search.txt"
    path.write_text("İstanbul e\u0301 🐱 " + "literal [x]. evidence " * 20, "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    result = call(project, "search_document", document_id=first["document_id"], query="[x].", max_matches=3, context_chars=10)
    found = []
    while True:
        found += result["matches"]
        if result["next_offset"] is None:
            break
        result = call(project, "search_document", **result["next_call"]["arguments"])
    assert len(found) == 20
    for match in found:
        assert first["text"][match["offset"]:match["end_offset"]] == "[x]."
        assert len(match["text"]) <= 24
    with pytest.raises(InvalidArguments, match="non-whitespace"):
        call(project, "search_document", document_id=first["document_id"], query="  ")


def test_blank_pdf_page_remains_in_structure_with_explicit_warning(project, tmp_path):
    path = make_pdf(tmp_path / "mixed.pdf", count=3, blank_pages=(2,))
    project.allow.allow(path)
    result = call(project, path=str(path), max_chars=12000)
    assert result["pages"] == 3 and result["coverage"]["total_units"] == 3
    assert result["unreadable_units"] == 1 and not result["extraction_complete"]
    assert result["metadata"]["pages_without_text"] == [2]
    assert "OCR" in result["structure"][1]["warning"]
    assert "## Page 2" in result["text"] and "## Page 3" in result["text"]


def test_legacy_large_max_chars_request_is_capped_not_silently_lost(project, tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("evidence " * 8000, "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=400000)
    assert 1000 <= len(first["text"]) <= 12000
    assert first["next_offset"] == len(first["text"]) and first["truncated"]
    full, _ = read_all(project, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert len(json.dumps(first, ensure_ascii=False)) < 60000


def test_long_headings_are_bounded_in_metadata_but_not_lost_from_text(project, tmp_path):
    path = tmp_path / "long-heading.md"
    heading = "Long scientific heading " * 100 + "IMPORTANT_HEADING_END"
    path.write_text("# " + heading + "\n\nBody evidence", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=1000)
    full, _ = read_all(project, first)
    assert heading in full
    assert len(first["structure"][0]["title"]) == 120


@pytest.mark.parametrize("corrupt", [[], {"units": "bad"}])
def test_corrupt_manifest_shape_returns_recoverable_error(project, tmp_path, corrupt):
    path = tmp_path / "manifest.txt"
    path.write_text("Snapshot data", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    (Path(first["artifact_dir"]) / "manifest.json").write_text(json.dumps(corrupt), "utf-8")
    with pytest.raises(InvalidArguments, match="snapshot is damaged"):
        call(project, document_id=first["document_id"])


def test_many_short_sections_fit_entire_payload_without_losing_text(project, tmp_path):
    path = tmp_path / "many-sections.md"
    path.write_text("\n\n".join(f"# Section {i}\n\nEvidence {i}." for i in range(800)), "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=12000)
    full, chunks = read_all(project, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert any(chunk["structure_truncated"] for chunk in chunks)
    for chunk in chunks:
        assert chunk["returned_chars"] > 0
        assert len(json.dumps(chunk["structure"], ensure_ascii=True)) <= documents.MAX_STRUCTURE_CHARS
        units = [unit for unit in json.loads((Path(first["artifact_dir"]) / "manifest.json").read_text())["units"]
                 if unit["start_offset"] < chunk["end_offset"] and unit["end_offset"] > chunk["offset"]]
        assert chunk["coverage"]["first_unit"] == units[0]["index"]
        assert chunk["coverage"]["last_unit"] == units[-1]["index"]


def test_escape_heavy_unicode_fits_serialized_budget_and_preserves_all_content(project, tmp_path):
    path = tmp_path / "escaped.txt"
    path.write_text(('"\\\x01\t👨‍👩‍👧‍👦 नाम 東京 e\u0301\n' * 3000) + "FINAL_ESCAPED_EVIDENCE", "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=12000)
    full, chunks = read_all(project, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert "FINAL_ESCAPED_EVIDENCE" in chunks[-1]["text"]
    assert 1 < len(chunks) < 100
    assert all(chunk["returned_chars"] > 0 for chunk in chunks)
    assert first["returned_chars"] < 12000


def test_large_pdf_metadata_is_bounded_before_text_pagination(project, tmp_path):
    from pypdf import PdfWriter

    path = make_pdf(tmp_path / "metadata.pdf", count=3)
    writer = PdfWriter(clone_from=str(path))
    writer.add_metadata({"/Title": "📚" * 5000, "/Author": '"\\' * 5000})
    writer.write(str(path))
    project.allow.allow(path)
    first = call(project, path=str(path), max_chars=12000)
    full, chunks = read_all(project, first)
    assert full == Path(first["text_path"]).read_text("utf-8")
    assert all(len(json.dumps(chunk["metadata"], ensure_ascii=True)) <= documents.MAX_METADATA_CHARS for chunk in chunks)


def test_impossible_metadata_budget_fails_without_recording_read_coverage(project, tmp_path):
    path = tmp_path / "metadata-limit.txt"
    path.write_text("Evidence " * 1000, "utf-8")
    project.allow.allow(path)
    first = call(project, path=str(path))
    project.document(first["document_id"]).name = "oversized registry name " * 1000
    before = (project.root / "audit.jsonl").read_text("utf-8")
    with pytest.raises(InvalidArguments, match="No document interval was delivered"):
        call(project, document_id=first["document_id"])
    assert (project.root / "audit.jsonl").read_text("utf-8") == before
