"""Saved formats preserve design, never source findings or executable objects."""
import hashlib
from io import StringIO
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import docx
from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordColor
from docx.enum.section import WD_SECTION
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import pytest
import pandas as pd

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, InvalidArguments, registry
from squipi_worker.security import AllowList, AccessDenied
from squipi_worker.server import Server
from squipi_worker.templates import library_action, resolve_template, template_chart_palette

SENTINEL = "OLD_PRIVATE_RESEARCH_72FA"


@pytest.fixture
def ctx(tmp_path):
    return JobContext(Project(tmp_path / "project", AllowList()), lambda *a: None, lambda: False)


def word_template(path):
    document = docx.Document()
    document.styles["Normal"].font.name = "Georgia"
    document.styles["Normal"].font.size = WordPt(12)
    document.styles["Heading 1"].font.color.rgb = WordColor.from_string("715A21")
    document.styles["Title"].font.size = WordPt(33)
    document.sections[0].header.paragraphs[0].text = "SYNTHETIC RESEARCH LAB"
    document.sections[0].footer.paragraphs[0].text = "Approved house format"
    document.add_paragraph(SENTINEL)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 2"
    table.cell(0, 0).text = SENTINEL
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = WordInches(7), WordInches(10)
    section.left_margin, section.right_margin = WordInches(.8), WordInches(.7)
    document.core_properties.title = SENTINEL
    document.core_properties.author = SENTINEL
    document.save(path)


def ppt_template(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    deck.core_properties.title = SENTINEL
    deck.core_properties.author = SENTINEL
    for layout_index in (0, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[layout_index])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F1E8")
        title = slide.shapes.title
        title.text = SENTINEL
        title.text_frame.paragraphs[0].runs[0].font.name = "Georgia"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(31)
        title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("715A21")
        slide.placeholders[1].text = SENTINEL
        slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.name = "Verdana"
        slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.size = Pt(19)
        slide.notes_slide.notes_text_frame.text = SENTINEL
    # Common editable charts carry harmless embedded XLSX workbooks.
    chart_data = CategoryChartData()
    chart_data.categories = [SENTINEL, "B"]
    chart_data.add_series(SENTINEL, [1, 2])
    deck.slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(3), Inches(2), chart_data)
    deck.save(path)


def imported(ctx, tmp_path, fmt, default=True):
    path = tmp_path / ("reference." + fmt)
    (word_template if fmt == "docx" else ppt_template)(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    ctx.project.allow.allow(path)
    result = library_action(ctx.project, {"action": "import", "path": str(path), "make_default": default})
    return result["templates"][-1], path, before


def exported(ctx, fmt, **extra):
    return registry.call("export_document", {"format": fmt, "title": "New household findings", "subtitle": "Synthetic current analysis",
        "sections": [{"heading": "Wellbeing", "text": "Only fresh results belong in the report.",
                      "table": {"columns": ["Measure", "Value"], "rows": [["Households", 60]]}},
                     {"heading": "", "text": ""}], **extra}, ctx)


def package_xml(path):
    with ZipFile(path) as package:
        return b"\n".join(package.read(name) for name in package.namelist() if name.endswith((".xml", ".rels")))


def test_word_preserves_format_branding_and_removes_old_content(ctx, tmp_path):
    item, source, digest = imported(ctx, tmp_path, "docx")
    result = exported(ctx, "docx")
    document = docx.Document(result["path"])
    assert result["template_applied"]["id"] == item["id"]
    assert document.styles["Normal"].font.name == "Georgia"
    assert document.styles["Normal"].font.size.pt == 12
    assert str(document.styles["Heading 1"].font.color.rgb) == "715A21"
    assert document.styles["Title"].font.size.pt == 33
    assert document.tables[0].style.name == "Light Shading Accent 2"
    assert len(document.sections) == 1
    assert document.sections[0].page_width == WordInches(7)
    assert document.sections[0].left_margin == WordInches(.8)
    assert document.sections[0].header.paragraphs[0].text == "SYNTHETIC RESEARCH LAB"
    assert document.sections[0].footer.paragraphs[0].text == "Approved house format"
    assert SENTINEL.encode() not in package_xml(result["path"])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert hashlib.sha256((ctx.project.root / "templates" / item["filename"]).read_bytes()).hexdigest() == digest


def test_ppt_preserves_sample_styles_master_layout_size_and_removes_all_old_data(ctx, tmp_path):
    item, source, digest = imported(ctx, tmp_path, "pptx")
    result = exported(ctx, "pptx")
    deck = Presentation(result["path"])
    assert result["template_applied"]["id"] == item["id"]
    assert deck.slide_width == Inches(10) and deck.slide_height == Inches(7.5)
    assert len(deck.slides) == 2  # title plus a real section, never empty third slide
    assert len(deck.slide_layouts) == len(Presentation(source).slide_layouts)
    assert str(deck.slides[0].background.fill.fore_color.rgb) == "F5F1E8"
    assert deck.slides[1].shapes.title.text_frame.paragraphs[0].font.name == "Georgia"
    assert deck.slides[1].shapes.title.text_frame.paragraphs[0].font.size.pt == 31
    assert str(deck.slides[1].shapes.title.text_frame.paragraphs[0].font.color.rgb) == "715A21"
    assert deck.slides[1].placeholders[1].text_frame.paragraphs[0].font.name == "Verdana"
    assert deck.slides[1].placeholders[1].text_frame.paragraphs[0].font.size.pt == 19
    assert SENTINEL.encode() not in package_xml(result["path"])
    with ZipFile(result["path"]) as package:
        assert not any("notesSlides/notesSlide" in name or "/embeddings/" in name for name in package.namelist())
        assert package.read("ppt/theme/theme1.xml") == ZipFile(source).read("ppt/theme/theme1.xml")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_none_and_explicit_style_override(ctx, tmp_path, fmt):
    item, _, _ = imported(ctx, tmp_path, fmt)
    plain = exported(ctx, fmt, template_id="none")
    assert plain["template_applied"] is None
    custom = exported(ctx, fmt, style={"font_family": "Arial", "heading_color": "#123456"})
    if fmt == "docx":
        artifact = docx.Document(custom["path"])
        assert artifact.styles["Normal"].font.name == "Arial"
        assert str(artifact.styles["Heading 1"].font.color.rgb) == "123456"
        assert artifact.styles["Normal"].font.size.pt == 12  # unspecified field retained
    else:
        artifact = Presentation(custom["path"])
        font = artifact.slides[1].shapes.title.text_frame.paragraphs[0].font
        assert font.name == "Arial" and str(font.color.rgb) == "123456"
        assert font.size.pt == 31


def test_library_persistence_rename_clear_remove_original_safe(ctx, tmp_path):
    item, source, digest = imported(ctx, tmp_path, "docx")
    restarted = Project(ctx.project.root, AllowList())
    assert library_action(restarted, {"action": "list"})["defaults"]["docx"] == item["id"]
    renamed = library_action(restarted, {"action": "update", "id": item["id"], "name": "House style", "make_default": False})
    assert renamed["templates"][0]["name"] == "House style" and renamed["defaults"] == {}
    removed = library_action(restarted, {"action": "remove", "id": item["id"]})
    assert removed["templates"] == [] and removed["defaults"] == {}
    assert not (restarted.root / "templates" / item["filename"]).exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def test_frozen_hermes_default_wins_and_missing_snapshot_is_visible(ctx, tmp_path, monkeypatch):
    item, _, _ = imported(ctx, tmp_path, "docx")
    monkeypatch.setenv("SQUIPI_TEMPLATE_DEFAULTS", json.dumps({"docx": item["id"]}))
    library_action(ctx.project, {"action": "update", "id": item["id"], "make_default": False})
    assert exported(ctx, "docx")["template_applied"]["id"] == item["id"]
    monkeypatch.setenv("SQUIPI_TEMPLATE_DEFAULTS", "{}")
    assert exported(ctx, "docx")["template_applied"] is None
    monkeypatch.setenv("SQUIPI_TEMPLATE_DEFAULTS", "not-json")
    with pytest.raises(InvalidArguments, match="snapshot"):
        exported(ctx, "docx")


@pytest.mark.parametrize("bad", ["../outside", "tpl_deadbeefdead", "", "other-name"])
def test_bad_template_selection_fails_before_artifact_write(ctx, bad):
    with pytest.raises(InvalidArguments):
        exported(ctx, "docx", template_id=bad)
    assert list((ctx.project.root / "artifacts").iterdir()) == []


def test_format_mismatch_missing_asset_and_ungranted_import(ctx, tmp_path):
    item, source, _ = imported(ctx, tmp_path, "docx")
    with pytest.raises(InvalidArguments, match="PPTX"):
        exported(ctx, "pptx", template_id=item["id"])
    (ctx.project.root / "templates" / item["filename"]).unlink()
    with pytest.raises(InvalidArguments):
        exported(ctx, "docx")
    other = tmp_path / "ungranted.docx"
    word_template(other)
    with pytest.raises(AccessDenied):
        library_action(ctx.project, {"action": "import", "path": str(other)})


@pytest.mark.parametrize("bad_entry,data", [
    ("../outside.xml", b"<a/>"), ("word/vbaProject.bin", b"x"),
    ("word/activeX/activeX1.xml", b"<a/>"), ("word/embeddings/oleObject1.bin", b"x"),
    ("word/entities.xml", b'<!DOCTYPE a [<!ENTITY x "private">]><a>&x;</a>'),
    ("word/_rels/linked.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" TargetMode="External" Target="https://example.test/private.png"/></Relationships>'),
])
def test_unsafe_packages_rejected_without_library_changes(ctx, tmp_path, bad_entry, data):
    source = tmp_path / "bad.docx"
    word_template(source)
    with ZipFile(source, "a", ZIP_DEFLATED) as package:
        package.writestr(bad_entry, data)
    ctx.project.allow.allow(source)
    with pytest.raises(InvalidArguments):
        library_action(ctx.project, {"action": "import", "path": str(source)})
    assert library_action(ctx.project, {"action": "list"})["templates"] == []


def test_manifest_and_asset_symlinks_rejected(ctx, tmp_path):
    item, _, _ = imported(ctx, tmp_path, "docx")
    asset = ctx.project.root / "templates" / item["filename"]
    asset.unlink()
    asset.symlink_to(tmp_path / "private.docx")
    with pytest.raises(AccessDenied):
        library_action(ctx.project, {"action": "list"})
    manifest = ctx.project.root / "templates.json"
    manifest.unlink()
    manifest.symlink_to(tmp_path / "private.json")
    with pytest.raises(AccessDenied):
        library_action(ctx.project, {"action": "list"})


def test_corrupt_manifest_not_overwritten(ctx):
    path = ctx.project.root / "templates.json"
    path.write_text("not json")
    with pytest.raises(InvalidArguments, match="not been overwritten"):
        library_action(ctx.project, {"action": "list"})
    assert path.read_text() == "not json"


def test_host_actions_are_private_and_read_only_catalog_has_no_source_text(ctx, tmp_path):
    item, _, _ = imported(ctx, tmp_path, "docx")
    assert "list_templates" in registry.names()
    assert not {"templates", "import_template", "remove_template", "update_template"}.intersection(registry.names())
    listed = registry.call("list_templates", {}, ctx)
    assert SENTINEL not in json.dumps(listed)
    out = StringIO()
    server = Server(ctx.project, out=out)
    server.handle({"id": "test", "method": "templates", "params": {"action": "list"}})
    assert json.loads(out.getvalue())["result"] == listed
    server._jobs.put(None)
    server._worker.join(timeout=2)


def test_blank_sections_fail_but_title_only_divider_is_kept(ctx):
    with pytest.raises(InvalidArguments, match="no section content"):
        registry.call("export_document", {"format": "pptx", "title": "Analysis", "sections": [{"heading": " "}]}, ctx)
    result = registry.call("export_document", {"format": "pptx", "title": "Analysis", "sections": [{"heading": "Discussion"}]}, ctx)
    assert len(Presentation(result["path"]).slides) == 2
    word = registry.call("export_document", {"format": "docx", "title": "Analysis", "sections": [{"heading": "", "table": {"columns": ["N"], "rows": [[30]]}}]}, ctx)
    document = docx.Document(word["path"])
    assert len(document.tables) == 1
    assert not any(p.style.name == "Heading 1" and not p.text.strip() for p in document.paragraphs)


def test_dataset_report_selects_formats_and_optout_separately(ctx, tmp_path):
    word, _, _ = imported(ctx, tmp_path, "docx")
    slides, _, _ = imported(ctx, tmp_path, "pptx")
    source = tmp_path / "households.csv"
    pd.DataFrame({"income": [1000 + index * 50 for index in range(30)],
                  "wellbeing": [20 + index * .9 for index in range(30)],
                  "area": ["North", "South", "West"] * 10}).to_csv(source, index=False)
    ctx.project.allow.allow(source)
    dataset = registry.call("import_dataset", {"path": str(source)}, ctx)["dataset_id"]
    result = registry.call("report_dataset", {"dataset_id": dataset, "format": "all",
                           "docx_template_id": word["id"], "pptx_template_id": "none"}, ctx)
    assert result["exports"]["docx"]["template_applied"]["id"] == word["id"]
    assert result["exports"]["pptx"]["template_applied"] is None
    assert docx.Document(result["docx_path"]).styles["Normal"].font.name == "Georgia"
    assert Presentation(result["pptx_path"]).slide_width != Inches(10)
    assert result["chart_template_applied"]["id"] == word["id"]
    assert result["chart_style_applied"]["chart_palette"][0] == "#715A21"
    assert result["style_applied"] == {}  # derived colors do not restyle Office text
    charts = [p for p in ctx.project.root.glob("artifacts/**/*.svg") if "_profile_" not in str(p)]
    assert charts and all("#715a21" in p.read_text().lower() for p in charts)
    explicit = registry.call("report_dataset", {"dataset_id": dataset, "format": "pptx", "style": {"chart_palette": ["#126543"]}}, ctx)
    assert explicit["chart_template_applied"] is None
    assert explicit["chart_style_applied"]["chart_palette"] == ["#126543"]
    assert explicit["exports"]["pptx"]["template_applied"]["id"] == slides["id"]


def test_chart_palette_reads_selected_powerpoint_theme_and_heading_literal(ctx, tmp_path):
    word, _, _ = imported(ctx, tmp_path, "docx")
    slides, _, _ = imported(ctx, tmp_path, "pptx")
    assert template_chart_palette(resolve_template(ctx.project, "docx"))[0] == "#715A21"
    palette = template_chart_palette(resolve_template(ctx.project, "pptx"))
    from lxml import etree
    with ZipFile(ctx.project.root / "templates" / slides["filename"]) as package:
        theme = etree.fromstring(package.read("ppt/theme/theme1.xml"))
        expected = theme.xpath("//a:clrScheme/a:accent1/a:srgbClr/@val", namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"})[0]
    assert len(palette) == 6 and palette[0] == "#" + expected.upper()


def test_narrow_word_table_headers_get_word_safe_widths_and_body_font_is_retained(ctx, tmp_path):
    imported(ctx, tmp_path, "docx")
    columns = ["Variable", "Valid N", "Missing", "Mean", "SD", "Median"]
    result = registry.call("export_document", {"format": "docx", "title": "Narrow report", "sections": [
        {"heading": "Descriptive statistics", "table": {"columns": columns, "rows": [["wellbeing", 30, 0, 61.2, 9.4, 62.1]]}},
        {"heading": "Quartiles", "table": {"columns": ["Variable", "Minimum", "25th percentile", "75th percentile", "Maximum"], "rows": [["wellbeing", 40, 50, 70, 89]]}},
    ]}, ctx)
    document = docx.Document(result["path"])
    assert document.styles["Normal"].font.name == "Georgia" and document.styles["Normal"].font.size.pt == 12
    from squipi_worker.tools.export import _docx_header_minimums
    for table in document.tables:
        headers = table.rows[0].cells
        size = headers[0].paragraphs[0].runs[0].font.size.pt
        assert size >= 9
        required = _docx_header_minimums([cell.text for cell in headers], size)
        assert all(cell.width + 635 >= minimum for cell, minimum in zip(headers, required))


def test_four_by_three_chart_uses_space_below_text_and_sources_have_one_native_bullet(ctx, tmp_path):
    imported(ctx, tmp_path, "pptx")
    from PIL import Image
    from pptx.oxml.ns import qn
    chart = ctx.project.root / "synthetic-chart.svg"
    chart.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="800" height="400"/>')
    Image.new("RGB", (800, 400), "white").save(chart.with_suffix(".png"))
    result = registry.call("export_document", {"format": "pptx", "title": "Current findings", "sections": [
        {"heading": "Household wellbeing", "text": "A short current finding with a full-width chart beneath it.", "chart_path": str(chart)}],
        "sources": ["Synthetic household survey"]}, ctx)
    deck = Presentation(result["path"])
    body = deck.slides[1].placeholders[1]
    picture = next(s for s in deck.slides[1].shapes if s.shape_type == 13)
    assert picture.width > Inches(4.5)
    assert picture.top >= body.top + body.height
    assert picture.top + picture.height <= deck.slide_height
    assert body.text_frame.paragraphs[0]._p.pPr.find(qn("a:buNone")) is not None
    source = deck.slides[2].placeholders[1].text_frame.paragraphs[0]
    assert source.text == "Synthetic household survey"
    assert source._p.pPr.find(qn("a:buChar")).get("char") == "•"


def test_old_word_annotations_and_body_images_are_not_carried_into_new_file(ctx, tmp_path):
    source = tmp_path / "annotated.docx"
    word_template(source)
    from PIL import Image
    picture = tmp_path / "private.png"
    Image.new("RGB", (5, 5), "red").save(picture)
    document = docx.Document(source)
    document.add_picture(str(picture))
    document.save(source)
    # Simulate a source comments part, which python-docx preserves unless its
    # relationship is explicitly discarded when clearing the reference body.
    with ZipFile(source) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    from lxml import etree
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    rel = etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
    rel.set("Id", "rIdSourceComments")
    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
    rel.set("Target", "comments.xml")
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    types = etree.fromstring(parts["[Content_Types].xml"])
    override = etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
    override.set("PartName", "/word/comments.xml")
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    parts["[Content_Types].xml"] = etree.tostring(types)
    parts["word/comments.xml"] = ('<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="0"><w:p><w:r><w:t>' + SENTINEL + '</w:t></w:r></w:p></w:comment></w:comments>').encode()
    with ZipFile(source, "w", ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    ctx.project.allow.allow(source)
    library_action(ctx.project, {"action": "import", "path": str(source), "make_default": True})
    result = exported(ctx, "docx")
    with ZipFile(result["path"]) as package:
        assert "word/comments.xml" not in package.namelist()
        assert not any(name.startswith("word/media/") for name in package.namelist())
    assert SENTINEL.encode() not in package_xml(result["path"])
