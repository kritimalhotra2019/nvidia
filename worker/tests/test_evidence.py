"""Evidence synthesis: extract_study_evidence (deterministic retrieval + regex extraction with quotes and offsets)
and meta_analysis (inverse-variance fixed effect, DerSimonian–Laird random effects) checked against hand computation."""
import json
import math
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.embeddings import Embedder
from squipi_worker.project import Project
from squipi_worker.registry import InvalidArguments, JobContext, registry
from squipi_worker.security import AllowList
from tests.test_embeddings import fake_embed  # deterministic stand-in for nomic


@pytest.fixture
def project(tmp_path):
    return Project(tmp_path / "project", AllowList())


@pytest.fixture
def fake_embedder(monkeypatch):
    monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": fake_embed(texts, kind))
    monkeypatch.setattr(Embedder, "available", lambda self: True)
    yield


def call(project, name, **args):
    return registry.call(name, args, JobContext(project, lambda *_: None, lambda: False))


PAPER_A = """# Cash transfers and household income in Kenya

## Abstract

We evaluate an unconditional cash transfer programme in rural Kenya. We randomly assigned 1,240 households in 62 villages to
treatment and control groups; the control group received nothing. The primary outcome was household income, measured at endline.

## Results

Household income rose by 11.6% (95% CI 4.2 to 19.0) relative to the control group, p = 0.004. Attrition was 14% at the
24-month follow-up and balanced across arms. The programme was delivered at a cost of USD 45 per household.
"""

PAPER_C = """# School grants and test scores

Journal of Development Economics, 2022. https://doi.org/10.1016/j.jdeveco.2022.102

## Design

We use a regression discontinuity design around the eligibility cutoff of the school grants programme in Ghana. The sample
comprises n = 5,112 students in 84 schools, with a comparison group of ineligible schools just above the cutoff.

## Findings

Test scores increased by 0.18 standard deviations (95% CI 0.05 to 0.31), p < 0.01, twelve months after the grants were paid.
Loss to follow-up was 6 percent.
"""

PAPER_D = """# Market access and smallholder welfare

## Summary

This cross-sectional study surveyed 820 respondents in urban Ghana about market access and food expenditure. The survey
describes patterns only; no programme was evaluated and no treatment effect is reported.
"""


def _write_docx(path: Path):
    import docx
    d = docx.Document()
    d.add_heading("Employment effects of a wage subsidy", 1)
    d.add_paragraph("We use a difference-in-differences design comparing 3,480 firms in Tanzania that received a wage subsidy "
                    "with a matched comparison group of firms that did not. The outcome is formal employment.")
    d.add_heading("Results", 2)
    d.add_paragraph("The estimated effect of the programme on employment was 2.3 percentage points (SE 0.9), p = 0.011. "
                    "Attrition was 8% between the baseline and the follow-up survey 18 months after the subsidy started.")
    d.save(str(path))


def four_papers(project, tmp_path):
    (tmp_path / "paper_a.md").write_text(PAPER_A)
    _write_docx(tmp_path / "paper_b.docx")
    (tmp_path / "paper_c.md").write_text(PAPER_C)
    (tmp_path / "paper_d.md").write_text(PAPER_D)
    ids = {}
    for name in ("paper_a.md", "paper_b.docx", "paper_c.md", "paper_d.md"):
        p = tmp_path / name
        project.allow.allow(p)
        ids[name] = call(project, "read_document", path=str(p))["document_id"]
    return ids


def _text(project, doc_id):
    out = call(project, "read_document", document_id=doc_id, max_chars=12000)
    assert out["complete"]
    return out["text"]


def _cell(out, doc_id, field):
    row = next(r for r in out["evidence_matrix"] if r["document_id"] == doc_id)
    return row["cells"][field]


def _check_quote(project, cell):
    """Every filled cell points at its own text: the offsets cut the match out of the saved snapshot and the quote contains it."""
    text = _text(project, cell["document_id"])
    match = text[cell["offset"]:cell["end_offset"]]
    assert match and " ".join(match.split()) in cell["quote"]
    assert cell["confidence"] in ("low", "medium", "high")


def test_evidence_matrix_finds_designs_sizes_effects_and_leaves_missing_cells_empty(project, tmp_path):
    ids = four_papers(project, tmp_path)
    a, b, c, d = ids["paper_a.md"], ids["paper_b.docx"], ids["paper_c.md"], ids["paper_d.md"]
    out = call(project, "extract_study_evidence")
    assert out["retrieval"] == "keyword"          # the embedding daemon is off in tests
    assert out["content_trust"] == "untrusted_document_data"
    assert {r["document_id"] for r in out["evidence_matrix"]} == {a, b, c, d}
    # Designs
    assert _cell(out, a, "design")["value"] == "randomised controlled trial"
    assert _cell(out, b, "design")["value"] == "difference-in-differences"
    assert _cell(out, c, "design")["value"] == "regression discontinuity"
    assert _cell(out, d, "design")["value"] == "cross-sectional"
    # Sample sizes with their unit from the same sentence
    assert _cell(out, a, "sample_size")["value"] == 1240 and _cell(out, a, "unit")["value"] == "households"
    assert "We randomly assigned 1,240 households" in _cell(out, a, "sample_size")["quote"]
    assert _cell(out, b, "sample_size")["value"] == 3480 and _cell(out, b, "unit")["value"] == "firms"
    assert _cell(out, c, "sample_size")["value"] == 5112
    assert _cell(out, d, "sample_size")["value"] == 820 and _cell(out, d, "unit")["value"] == "respondents"
    # Effect sizes, units, intervals, SE and p from the effect sentence
    ea = _cell(out, a, "effect_size")
    assert ea["value"] == 11.6 and ea["confidence"] == "high" and "rose by 11.6%" in ea["quote"]
    assert _cell(out, a, "effect_unit")["value"] == "percent"
    assert _cell(out, a, "ci_low")["value"] == 4.2 and _cell(out, a, "ci_high")["value"] == 19.0
    assert _cell(out, a, "p_value")["value"] == 0.004
    eb = _cell(out, b, "effect_size")
    assert eb["value"] == 2.3 and _cell(out, b, "effect_unit")["value"] == "percentage points"
    assert _cell(out, b, "standard_error")["value"] == 0.9 and _cell(out, b, "p_value")["value"] == 0.011
    assert _cell(out, b, "ci_low") is None
    ec = _cell(out, c, "effect_size")
    assert ec["value"] == 0.18 and _cell(out, c, "effect_unit")["value"] == "standard deviations"
    assert _cell(out, c, "ci_low")["value"] == 0.05 and _cell(out, c, "ci_high")["value"] == 0.31
    assert _cell(out, c, "p_value")["value"] == 0.01 and _cell(out, c, "p_value")["operator"] == "<"
    # The descriptive paper has no effect: the cells stay empty rather than guessed
    for f in ("effect_size", "effect_unit", "ci_low", "ci_high", "standard_error", "p_value", "attrition"):
        assert _cell(out, d, f) is None, f
    cov_d = next(cv for cv in out["coverage"] if cv["document_id"] == d)
    assert "effect_size" in cov_d["not_found"] and cov_d["fields_asked"] == 18
    assert any("paper_d.md: not found" in n and "effect_size" in n for n in out["notes"])
    # Attrition, follow-up, cost, country, setting, intervention, comparison, outcome
    assert _cell(out, a, "attrition")["value"] == 14 and _cell(out, b, "attrition")["value"] == 8 and _cell(out, c, "attrition")["value"] == 6
    assert _cell(out, a, "follow_up")["value"] == "24 months" and _cell(out, b, "follow_up")["value"] == "18 months"
    assert _cell(out, c, "follow_up")["value"] == "12 months" and _cell(out, c, "follow_up")["months"] == 12
    assert _cell(out, c, "citation")["value"] == "doi:10.1016/j.jdeveco.2022.102" and _cell(out, a, "citation") is None
    assert _cell(out, c, "unit")["value"] == "students" and _cell(out, c, "outcome")["value"] == "test scores"
    assert _cell(out, c, "country")["value"] == "Ghana"
    assert _cell(out, a, "cost")["value"] == "USD 45 per household"
    assert _cell(out, a, "country")["value"] == "Kenya" and _cell(out, b, "country")["value"] == "Tanzania" and _cell(out, d, "country")["value"] == "Ghana"
    assert _cell(out, a, "setting")["value"] == "rural" and _cell(out, d, "setting")["value"] == "urban"
    assert _cell(out, a, "intervention")["value"] == "cash transfer"
    assert _cell(out, b, "intervention")["value"] == "wage subsidy"
    assert _cell(out, a, "comparison")["value"] == "control group"
    assert _cell(out, b, "comparison")["value"] == "matched comparison group"
    assert _cell(out, a, "outcome")["value"] == "household income"
    assert _cell(out, b, "outcome")["value"] == "employment"
    # Every filled cell carries document id, offsets and a verbatim quote that the offsets point into
    for row in out["evidence_matrix"]:
        for field, cell in row["cells"].items():
            if cell is not None:
                assert cell["document_id"] == row["document_id"]
                _check_quote(project, cell)
    # CSV written and registered as a dataset named "evidence matrix.csv"
    assert out["dataset_name"] == "evidence matrix.csv" and Path(out["csv_path"]).is_file()
    assert any(ds.id == out["dataset_id"] for ds in project.datasets())
    table = pd.read_csv(out["csv_path"])
    assert len(table) == 4 and {"study", "effect_size", "ci_low", "ci_high", "standard_error", "sample_size", "design"} <= set(table.columns)
    assert table.set_index("document_id").loc[a, "effect_size"] == 11.6
    assert pd.isna(table.set_index("document_id").loc[d, "effect_size"])
    assert Path(out["json_path"]).is_file()
    assert "Empty cells mean no confident match" in out["summary_text"]
    json.dumps(out)


def test_evidence_matrix_subset_of_fields_and_documents_and_semantic_mode(project, tmp_path, fake_embedder):
    ids = four_papers(project, tmp_path)
    a = ids["paper_a.md"]
    out = call(project, "extract_study_evidence", document_ids=[a], fields=["design", "effect_size", "ci_low", "ci_high"], max_passages_per_field=2)
    assert out["retrieval"] == "semantic" and out["embedding_model"]
    assert [r["document_id"] for r in out["evidence_matrix"]] == [a]
    assert set(out["fields"]) == {"design", "effect_size", "ci_low", "ci_high"}
    assert _cell(out, a, "effect_size")["value"] == 11.6 and _cell(out, a, "ci_high")["value"] == 19.0
    assert out["coverage"][0]["fields_asked"] == 4 and out["coverage"][0]["fields_found"] == 4
    with pytest.raises(InvalidArguments):
        call(project, "extract_study_evidence", document_ids=["doc_missing"])


def test_evidence_matrix_refuses_without_documents(project):
    with pytest.raises(InvalidArguments):
        call(project, "extract_study_evidence")


# ---------------------------------------------------------------------------------------------
# meta_analysis
# ---------------------------------------------------------------------------------------------

STUDIES = [{"study": "Kenya 2019", "effect": 11.6, "se": 3.8, "n": 1240},
           {"study": "Tanzania 2021", "effect": 2.3, "se": 0.9, "n": 3480},
           {"study": "Uganda 2020", "effect": 6.0, "se": 2.0, "n": 900}]


def hand_dersimonian_laird(y, se):
    """DerSimonian & Laird (1986) by hand, kept in the test so the numbers can be audited line by line."""
    w = [1 / s ** 2 for s in se]
    sw = sum(w)
    theta_f = sum(wi * yi for wi, yi in zip(w, y)) / sw            # inverse-variance fixed effect
    se_f = math.sqrt(1 / sw)
    q = sum(wi * (yi - theta_f) ** 2 for wi, yi in zip(w, y))       # Cochran's Q
    df = len(y) - 1
    c = sw - sum(wi ** 2 for wi in w) / sw
    tau2 = max(0.0, (q - df) / c)                                   # method-of-moments between-study variance
    i2 = max(0.0, (q - df) / q) * 100
    w_r = [1 / (s ** 2 + tau2) for s in se]
    theta_r = sum(wi * yi for wi, yi in zip(w_r, y)) / sum(w_r)
    se_r = math.sqrt(1 / sum(w_r))
    return {"theta_f": theta_f, "se_f": se_f, "q": q, "tau2": tau2, "i2": i2, "theta_r": theta_r, "se_r": se_r,
            "weights_r": [wi / sum(w_r) * 100 for wi in w_r]}


def test_meta_analysis_matches_hand_computed_dersimonian_laird(project):
    out = call(project, "meta_analysis", studies=STUDIES, effect_type="percent")
    hand = hand_dersimonian_laird([s["effect"] for s in STUDIES], [s["se"] for s in STUDIES])
    assert hand["tau2"] > 0 and hand["i2"] > 0           # the example has real heterogeneity
    assert out["model"] == "random" and out["k"] == 3
    assert abs(out["unrounded"]["estimate"] - hand["theta_r"]) < 1e-6
    assert abs(out["unrounded"]["se"] - hand["se_r"]) < 1e-6
    assert abs(out["unrounded"]["tau2"] - hand["tau2"]) < 1e-6
    assert abs(out["unrounded"]["i2"] - hand["i2"]) < 1e-6
    assert abs(out["unrounded"]["q"] - hand["q"]) < 1e-6
    assert abs(out["unrounded"]["fixed_estimate"] - hand["theta_f"]) < 1e-6
    assert abs(out["pooled"]["ci_low"] - (hand["theta_r"] - 1.959964 * hand["se_r"])) < 1e-3
    for exact, row, w in zip(out["unrounded"]["weights_pct"], out["studies"], hand["weights_r"]):
        assert abs(exact - w) < 1e-6 and abs(row["weight_random_pct"] - w) < 1e-2    # table weights are rounded to 4 significant figures
    assert abs(sum(r["weight_pct"] for r in out["studies"]) - 100) < 1e-6
    # Leave-one-out: three rows, each recomputed without that study
    assert [r["omitted"] for r in out["leave_one_out"]] == [s["study"] for s in STUDIES]
    without_kenya = hand_dersimonian_laird([2.3, 6.0], [0.9, 2.0])
    assert abs(out["leave_one_out"][0]["estimate"] - without_kenya["theta_r"]) < 1e-3
    # Prediction interval (k = 3 → t with 1 df)
    from scipy import stats as st
    half = st.t.ppf(0.975, 1) * math.sqrt(hand["tau2"] + hand["se_r"] ** 2)
    assert abs(out["prediction_interval"]["low"] - (hand["theta_r"] - half)) < 1e-3
    # Forest plot written and parses as XML; CSVs exist
    svg_path = Path(out["forest_svg"])
    assert svg_path.is_file()
    root = ET.parse(svg_path).getroot()
    assert root.tag.endswith("svg")
    assert "polygon" in svg_path.read_text()           # the pooled diamond
    assert Path(out["csv_path"]).is_file() and Path(out["leave_one_out_csv"]).is_file()
    assert len(pd.read_csv(out["leave_one_out_csv"])) == 3
    # Wording and warnings
    assert out["heterogeneity"]["i2_band"] in ("low", "moderate", "substantial", "considerable")
    assert "Heterogeneity is " + out["heterogeneity"]["i2_band"] in out["interpretation"]
    assert any("Only 3 studies" in n for n in out["notes"])
    json.dumps(out)


def test_meta_analysis_fixed_effect_is_inverse_variance_weighted(project):
    out = call(project, "meta_analysis", studies=STUDIES, model="fixed", effect_type="percent")
    hand = hand_dersimonian_laird([s["effect"] for s in STUDIES], [s["se"] for s in STUDIES])
    assert abs(out["unrounded"]["estimate"] - hand["theta_f"]) < 1e-6
    assert abs(out["unrounded"]["se"] - hand["se_f"]) < 1e-6
    w = [1 / s["se"] ** 2 for s in STUDIES]
    for exact, row, wi in zip(out["unrounded"]["weights_pct"], out["studies"], w):
        assert abs(exact - wi / sum(w) * 100) < 1e-6 and abs(row["weight_pct"] - wi / sum(w) * 100) < 1e-2
    assert out["prediction_interval"] is None
    assert "fixed-effect" in out["method"]


def test_meta_analysis_derives_se_from_ci_and_handles_subgroups(project):
    studies = [{"study": "A", "effect": 11.6, "ci_low": 4.2, "ci_high": 19.0, "subgroup": "cash"},
               {"study": "B", "effect": 2.3, "se": 0.9, "subgroup": "training"},
               {"study": "C", "effect": 6.0, "ci_low": 2.08, "ci_high": 9.92, "subgroup": "cash"}]
    out = call(project, "meta_analysis", studies=studies, subgroup="subgroup", effect_type="percent", order="effect")
    se_a = (19.0 - 4.2) / (2 * 1.959963984540054)
    assert abs(out["studies"][0]["se"] - round(float(f"{se_a:.4g}"), 6)) < 1e-6
    assert any("SE derived from the 95% CI" in n and "A" in n for n in out["notes"])
    groups = {g["subgroup"]: g for g in out["subgroups"]}
    assert groups["cash"]["k"] == 2 and groups["training"]["k"] == 1 and groups["training"]["estimate"] == 2.3
    hand_cash = hand_dersimonian_laird([11.6, 6.0], [se_a, (9.92 - 2.08) / (2 * 1.959963984540054)])
    assert abs(groups["cash"]["estimate"] - float(f"{hand_cash['theta_r']:.4g}")) < 1e-3
    assert out["subgroup_test"]["df"] == 1 and 0 <= out["subgroup_test"]["p"] <= 1
    assert "Subgroups:" in out["interpretation"]


def test_meta_analysis_from_the_evidence_matrix_dataset(project, tmp_path):
    four_papers(project, tmp_path)
    matrix = call(project, "extract_study_evidence")
    # Papers A and C carry a CI, B carries an SE, D has no effect: three studies pooled, one row skipped with a note.
    out = call(project, "meta_analysis", dataset_id=matrix["dataset_id"], effect_type="as reported")
    assert out["k"] == 3 and sorted(r["study"] for r in out["studies"]) == ["paper_a", "paper_b", "paper_c"]
    assert any("without an effect value skipped" in n and "paper_d" in n for n in out["notes"])
    assert any("effect units differ" in n for n in out["notes"])      # percent, percentage points, standard deviations
    assert next(r for r in out["studies"] if r["study"] == "paper_b")["se"] == 0.9


def test_meta_analysis_refuses_bad_inputs(project):
    with pytest.raises(InvalidArguments, match="at least 2"):
        call(project, "meta_analysis", studies=[{"study": "only", "effect": 1.0, "se": 0.5}])
    with pytest.raises(InvalidArguments, match="positive"):
        call(project, "meta_analysis", studies=[{"study": "a", "effect": 1.0, "se": 0.0}, {"study": "b", "effect": 2.0, "se": 0.5}])
    with pytest.raises(InvalidArguments, match="ci_low"):
        call(project, "meta_analysis", studies=[{"study": "a", "effect": 1.0, "ci_low": 3.0, "ci_high": 2.0}, {"study": "b", "effect": 2.0, "se": 0.5}])
    with pytest.raises(InvalidArguments, match="implies SE"):
        call(project, "meta_analysis", studies=[{"study": "a", "effect": 1.0, "se": 0.1, "ci_low": 0.0, "ci_high": 2.0}, {"study": "b", "effect": 2.0, "se": 0.5}])
    with pytest.raises(InvalidArguments, match="give se or a 95% interval"):
        call(project, "meta_analysis", studies=[{"study": "a", "effect": 1.0}, {"study": "b", "effect": 2.0, "se": 0.5}])
    with pytest.raises(InvalidArguments, match="exactly one"):
        call(project, "meta_analysis", studies=STUDIES, dataset_id="ds_x")
    with pytest.raises(InvalidArguments):
        call(project, "meta_analysis", dataset_id="ds_missing")
