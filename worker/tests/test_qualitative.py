"""Qualitative coding: themes from transcripts (keyword and embedding modes), codebook assignment, anonymised quotes."""
import json

import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.embeddings import Embedder
from squipi_worker.project import Project
from squipi_worker.registry import InvalidArguments, JobContext, registry
from squipi_worker.security import AllowList
from squipi_worker.tools.qualitative import anonymise_text

# A deterministic stand-in for nomic over the three interview topics (same pattern as tests/test_embeddings.py).
TOPICS = {
    "water": 0, "borehole": 0, "jerrycan": 0, "river": 0, "well": 0, "thirst": 0,
    "school": 1, "fees": 1, "tuition": 1, "uniform": 1, "teacher": 1, "pupil": 1,
    "clinic": 2, "health": 2, "nurse": 2, "dispensary": 2, "medicine": 2, "sick": 2,
}


def fake_embed(texts, kind="document"):
    out = []
    for text in texts:
        vector = [0.01] * 8
        folded = text.lower()
        for word, slot in TOPICS.items():
            if word in folded:
                vector[slot] += 1.0
        out.append(vector)
    return out


@pytest.fixture
def project(tmp_path):
    return Project(tmp_path / "project", AllowList())


@pytest.fixture
def fake_embedder(monkeypatch):
    monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": fake_embed(texts, kind))
    monkeypatch.setattr(Embedder, "available", lambda self: True)
    yield


def call(project, name, **args):
    return registry.call(name, args, JobContext(project, lambda *_: None, lambda: False))


NAME = "Grace Wanjiru"
PHONE = "0712 345 678"
PRE_CONSENT = "we have lived here for ten years already"
CLINIC_QUOTE = "The health centre is nine kilometres away and the nurse only comes on Tuesdays, so when a child is sick we hire a motorbike."

T1 = f"""Interview 1, village A
[00:00:05] Interviewer: Thank you for meeting me today, before we start can you tell me a little about yourself?
[00:00:12] Respondent: My name is {NAME} and {PRE_CONSENT}, near the river.
[00:00:30] Interviewer: Do you agree to take part and to be recorded?
[00:00:35] Respondent: Yes, I consent to be recorded for this study.
[00:01:00] Interviewer: Let us start with water. Where do you fetch water?
[00:01:10] Respondent: The borehole dried up in March so we walk to the river with jerrycans, two hours each way for water.
[00:02:00] Interviewer: And the children, are they in school?
[00:02:08] Respondent: Two are in school but the school fees doubled this year and the uniform alone costs more than a goat, so the youngest stays home.
[00:03:00] Interviewer: What about health care?
[00:03:10] Respondent: {CLINIC_QUOTE}
[00:04:00] Interviewer: Anything else about water?
[00:04:05] Respondent: The water from the river makes the children sick, so we boil it, but firewood for boiling water is also far.
[00:05:00] Interviewer: Thank you.
"""

T2 = f"""# Interview 2
[00:00:04] Interviewer: Do you agree to take part and to be recorded?
[00:00:09] Participant: Yes, I consent to be recorded, you can use what I say.
[00:00:40] Interviewer: Tell me about water in this area.
[00:00:50] Participant: My neighbour {NAME} gave me her number {PHONE} so we share the cost of a water cart from the borehole in the next village; without that cart we would have no water at all.
[00:01:40] Interviewer: How are the school fees?
[00:01:50] Participant: The school fees are the biggest worry; the teacher sends pupils home when tuition is late and the uniform must be bought in town.
[00:02:40] Interviewer: Is the clinic far?
[00:02:50] Participant: The dispensary closed, so the nearest clinic is in the district town and the nurse there charges for medicine we cannot afford.
[00:03:30] Interviewer: What else about school?
[00:03:40] Participant: If the fees were lower every child here would be in school; the teacher is good but the fees push families out.
[00:04:30] Interviewer: And water again?
[00:04:40] Participant: Water is life; without the borehole women spend the whole morning on water and the jerrycans are heavy.
"""

T3_LINES = [
    "Interviewer: Do you agree to take part and to be recorded?",
    "Respondent: Yes, I consent to be recorded.",
    "Interviewer: Where do you get water?",
    "Respondent: We queue at the borehole from five in the morning and the water is rationed to two jerrycans per household each day.",
    "Interviewer: Tell me about the school.",
    "Respondent: The school fees and the uniform are too much; my son missed a whole term because the tuition was not paid.",
    "Interviewer: And the clinic?",
    "Respondent: The clinic is far, nearly ten kilometres, and the health workers have no medicine, so sick people stay at home.",
    "Interviewer: Is there anything else you would like to add?",
    "Respondent: The football match last weekend was exciting and the whole village came to watch the final game together.",
]


def three_transcripts(project, tmp_path):
    a = tmp_path / "interview_1.txt"
    a.write_text(T1)
    b = tmp_path / "interview_2.md"
    b.write_text(T2)
    import docx
    d = docx.Document()
    for line in T3_LINES:
        d.add_paragraph(line)
    c = tmp_path / "interview_3.docx"
    d.save(str(c))
    ids = []
    for p in (a, b, c):
        project.allow.allow(p)
        ids.append(call(project, "read_document", path=str(p))["document_id"])
    return ids


def theme_for(themes, word):
    return next((t for t in themes if any(word in term for term in t["top_terms"])), None)


def assert_no_personal_data(blob: str):
    assert NAME not in blob and "Wanjiru" not in blob and "Grace" not in blob
    assert PHONE not in blob and "0712" not in blob and "345 678" not in blob


def test_anonymise_text_rules():
    out = anonymise_text(f"My neighbour {NAME} gave me {PHONE} and grace@example.com; Mr Otieno agreed. The Water Committee met Last Tuesday in 2024 for 1,500 shillings.")
    assert_no_personal_data(out)
    assert "[name]" in out and "[phone]" in out and "[email]" in out and "Otieno" not in out
    assert "Water Committee" in out and "2024" in out and "1,500" in out  # not names, not phones


def test_keyword_mode_discovers_the_three_topics_and_protects_people(project, tmp_path):
    ids = three_transcripts(project, tmp_path)
    out = call(project, "code_transcripts", document_ids=ids, n_themes=3, consent_marker="consent to be recorded")
    assert out["mode"] == "keyword" and out["n_documents"] == 3
    themes = out["themes"]
    assert len(themes) == 3
    water, school, clinic = theme_for(themes, "water"), theme_for(themes, "fee"), theme_for(themes, "clinic")
    assert water and school and clinic and len({water["theme_id"], school["theme_id"], clinic["theme_id"]}) == 3
    assert water["n_segments"] >= 4 and school["n_segments"] >= 3 and clinic["n_segments"] >= 3
    assert all(1 <= len(t["quotes"]) <= 3 for t in themes)
    assert sum(t["share_pct"] for t in themes) == pytest.approx(100, abs=0.3)
    assert all(t["n_documents"] >= 2 for t in (water, school, clinic))
    # Interviewer turns are excluded; consent cut-offs are counted; quotes carry their provenance.
    notes = out["notes"]
    assert notes["excluded_speaker_turns"] >= 15
    assert notes["segments_before_consent"][ids[0]] >= 1 and notes["mode"] == "keyword"
    for t in themes:
        for q in t["quotes"]:
            assert not q["speaker"].lower().startswith("interviewer")
            assert q["document_id"] in ids and q["end_offset"] > q["offset"] >= 0 and len(q["text"].split()) <= 61
    # Personal data never appears, in the result or in the CSV; placeholders do.
    blob = json.dumps(out, default=str)
    assert_no_personal_data(blob)
    assert PRE_CONSENT not in blob
    import pandas as pd
    seg = pd.read_csv(out["segments_csv_path"])
    csv_text = seg.to_csv(index=False)
    assert_no_personal_data(csv_text)
    assert PRE_CONSENT not in csv_text
    assert "[name]" in csv_text and "[phone]" in csv_text
    assert not seg["speaker"].str.lower().str.startswith("interviewer").any()
    assert {"document_id", "speaker", "timestamp", "offset", "end_offset", "theme_id", "theme", "text"} <= set(seg.columns)
    assert out["segments_dataset_id"] in {ds.id for ds in project.datasets()}
    # Offsets point into the saved snapshot: reading the document there returns the (un-anonymised) original passage.
    row = seg[seg["text"].str.contains("borehole dried")].iloc[0]
    window = call(project, "read_document", document_id=row["document_id"], offset=int(row["offset"]), max_chars=int(row["end_offset"] - row["offset"]))
    assert window["text"].startswith("The borehole dried up")
    assert row["timestamp"] == "00:01:10"
    # Artifacts and templated wording.
    from pathlib import Path
    assert Path(out["chart_path"]).exists() and Path(out["matrix_csv_path"]).exists()
    matrix = pd.read_csv(out["matrix_csv_path"])
    assert len(matrix) == 3 and "interview_1.txt" in matrix.columns
    assert out["interpretation"][0].startswith("3 themes across 3 transcripts")
    assert "largest theme" in out["summary_text"]
    assert out["content_trust"] == "untrusted_document_data"
    assert len(out["codebook"]) == 3 and all(c["top_terms"] for c in out["codebook"])


def test_auto_k_by_silhouette_still_covers_each_topic(project, tmp_path):
    ids = three_transcripts(project, tmp_path)
    out = call(project, "code_transcripts", document_ids=ids)
    assert 4 <= len(out["themes"]) <= 12
    assert out["notes"]["k_candidates"][0]["k"] == 4 and out["notes"]["silhouette"] is not None
    for word in ("water", "fee", "clinic"):
        assert theme_for(out["themes"], word), word


def test_semantic_mode_with_fake_embedder(project, tmp_path, fake_embedder):
    ids = three_transcripts(project, tmp_path)
    out = call(project, "code_transcripts", document_ids=ids, n_themes=3, consent_marker="consent to be recorded")
    assert out["mode"] == "semantic" and out["embedding_model"]
    themes = out["themes"]
    water, school, clinic = theme_for(themes, "water"), theme_for(themes, "fee"), theme_for(themes, "clinic")
    assert water and school and clinic and len({water["theme_id"], school["theme_id"], clinic["theme_id"]}) == 3
    assert_no_personal_data(json.dumps(out, default=str))


@pytest.mark.parametrize("semantic", [False, True])
def test_codebook_mode_assigns_codes_and_leaves_off_topic_uncoded(project, tmp_path, monkeypatch, semantic):
    if semantic:
        monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": fake_embed(texts, kind))
    ids = three_transcripts(project, tmp_path)
    codebook = [
        {"code": "Water access", "description": "getting water: boreholes, rivers, jerrycans, distance and time to fetch water",
         "examples": ["we walk far to fetch water", "the borehole is dry"]},
        {"code": "School fees", "description": "cost of schooling: school fees, tuition, uniform, children sent home",
         "examples": ["the fees are too high", "the teacher sent the pupils home"]},
        {"code": "Clinic distance", "description": "distance to the clinic or health centre, nurse and medicine availability",
         "examples": ["the clinic is far", "the dispensary has no medicine"]},
    ]
    out = call(project, "code_transcripts", document_ids=ids, codebook=codebook, consent_marker="consent to be recorded")
    assert out["mode"] == ("semantic" if semantic else "keyword")
    by_id = {t["theme_id"]: t for t in out["themes"]}
    labels = [t["label"] for t in out["themes"]]
    assert labels == ["Water access", "School fees", "Clinic distance", "uncoded"]
    import pandas as pd
    seg = pd.read_csv(out["segments_csv_path"])
    assert (seg[seg["text"].str.contains("borehole|jerrycan", regex=True)]["theme"] == "Water access").all()
    assert (seg[seg["text"].str.contains("fees")]["theme"] == "School fees").all()
    assert (seg[seg["text"].str.contains("clinic|health centre|dispensary", regex=True)]["theme"] == "Clinic distance").all()
    football = seg[seg["text"].str.contains("football")]
    assert len(football) == 1 and football["theme"].iloc[0] == "uncoded"
    assert by_id["uncoded"]["n_segments"] >= 1
    assert all(t["n_segments"] >= 3 for t in out["themes"][:3])
    assert "uncoded" in out["interpretation"][2]
    assert [c["code"] for c in out["codebook"]] == ["Water access", "School fees", "Clinic distance"]
    assert_no_personal_data(json.dumps(out, default=str))


def test_search_quotes_finds_the_clinic_distance_quote(project, tmp_path, monkeypatch):
    ids = three_transcripts(project, tmp_path)
    out = call(project, "search_quotes", document_ids=ids, query="how far is the health centre", max_quotes=3)
    assert out["mode"] == "keyword" and "unavailable" in out["fallback"]
    top = out["quotes"][0]
    assert top["text"] == anonymise_text(CLINIC_QUOTE) == CLINIC_QUOTE
    assert top["document_id"] == ids[0] and top["timestamp"] == "00:03:10" and top["speaker"] == "Respondent"
    window = call(project, "read_document", document_id=ids[0], offset=top["offset"], max_chars=top["end_offset"] - top["offset"])
    assert window["text"] == CLINIC_QUOTE
    assert all(not q["speaker"].lower().startswith("interviewer") for q in out["quotes"])
    assert_no_personal_data(json.dumps(out, default=str))
    # By meaning: the three clinic quotes come first even when the words differ.
    monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": fake_embed(texts, kind))
    sem = call(project, "search_quotes", document_ids=ids, query="how far is the health centre", max_quotes=3)
    assert sem["mode"] == "semantic"
    assert all(any(w in q["text"].lower() for w in ("clinic", "health", "dispensary")) for q in sem["quotes"])


def test_invalid_arguments_are_refused(project, tmp_path):
    ids = three_transcripts(project, tmp_path)
    auto = call(project, "code_transcripts", document_ids=ids, n_themes=1)    # under 2 means automatic, not an error
    assert len(auto["themes"]) >= 2
    with pytest.raises(InvalidArguments):
        call(project, "code_transcripts", document_ids=ids, n_themes=21)
    with pytest.raises(InvalidArguments):
        call(project, "code_transcripts", document_ids=ids, n_themes=19)   # more themes than segments
    with pytest.raises(InvalidArguments):
        call(project, "code_transcripts", document_ids=ids, n_themes=3, codebook=[{"code": "Water"}])
    with pytest.raises(InvalidArguments):
        call(project, "code_transcripts", document_ids=["doc_missing"])
    with pytest.raises(InvalidArguments):
        call(project, "code_transcripts", document_ids=ids + [ids[0]])
    with pytest.raises(InvalidArguments):
        call(project, "search_quotes", document_ids=ids, query="   ")
