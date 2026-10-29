"""Presentation preferences reach real Office files and charts, without becoming code."""
from pathlib import Path
from zipfile import ZipFile

import docx
import pandas as pd
import pytest
from pptx import Presentation
from docx.oxml.ns import qn

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, InvalidArguments, registry
from squipi_worker.security import AllowList
from squipi_worker.tools import charts
from squipi_worker.tools.document_style import normalize_style


STYLE = {"accent_color": "#136F63", "heading_color": "#123456", "body_color": "#243746",
         "background_color": "#FEFCF4", "font_family": "Arial", "font_size": 12,
         "chart_palette": ["#136F63", "#F4A261", "#357EDD"]}
SECTIONS = [{"heading": "Household wellbeing", "text": "Descriptive statistics for the synthetic pilot.\n\n- Report missing observations\n- Avoid causal claims",
             "table": {"columns": ["Measure", "Value"], "rows": [["Households", 60], ["Mean wellbeing", 61.25]]}}]


@pytest.fixture
def ctx(tmp_path):
    return JobContext(Project(tmp_path / "project", AllowList()), lambda *a: None, lambda: False)


def export_file(ctx, fmt, style=None):
    args = {"format": fmt, "title": "Household research", "subtitle": "Synthetic style fixture", "sections": SECTIONS}
    if style is not None:
        args["style"] = style
    return registry.call("export_document", args, ctx)


def test_word_style_is_present_in_editable_document_and_table(ctx):
    result = export_file(ctx, "docx", STYLE)
    document = docx.Document(result["path"])
    assert result["style_applied"] == STYLE
    assert document.styles["Normal"].font.name == "Arial"
    assert document.styles["Normal"].font.size.pt == 12
    assert str(document.styles["Normal"].font.color.rgb) == "243746"
    for name in ("Title", "Heading 1", "Subtitle"):
        assert document.styles[name].font.name == "Arial"
        assert str(document.styles[name].font.color.rgb) == "123456"
        color = document.styles[name].element.xpath("./w:rPr/w:color")[0]
        assert color.get(qn("w:themeColor")) is None
    assert document._element.xpath("./w:background")[0].get(qn("w:color")) == "FEFCF4"
    assert document.settings.element.xpath("./w:displayBackgroundShape")
    header = document.tables[0].cell(0, 0)
    assert header._tc.xpath("./w:tcPr/w:shd")[0].get(qn("w:fill")) == "136F63"
    assert str(header.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert str(document.tables[0].cell(1, 0).paragraphs[0].runs[0].font.color.rgb) == "243746"


def test_powerpoint_style_reaches_titles_body_tables_and_background(ctx):
    result = export_file(ctx, "pptx", STYLE)
    presentation = Presentation(result["path"])
    assert result["style_applied"] == STYLE
    for slide in presentation.slides:
        assert str(slide.background.fill.fore_color.rgb) == "FEFCF4"
        title_font = slide.shapes.title.text_frame.paragraphs[0].runs[0].font
        assert title_font.name == "Arial" and str(title_font.color.rgb) == "123456"
        assert title_font.size.pt == pytest.approx(19.2)
    slide = presentation.slides[1]
    body = next(s for s in slide.shapes if s.has_text_frame and s != slide.shapes.title)
    assert body.text_frame.paragraphs[0].font.size.pt == 12
    assert str(body.text_frame.paragraphs[0].runs[0].font.color.rgb) == "243746"
    table = next(s.table for s in slide.shapes if s.has_table)
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "136F63"
    assert str(table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert str(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "243746"


@pytest.mark.parametrize("fmt,parts", [("docx", ["word/document.xml", "word/styles.xml"]),
                                      ("pptx", ["ppt/slides/slide1.xml", "ppt/slides/slide2.xml"])])
def test_omitted_and_empty_style_leave_existing_defaults_identical(ctx, fmt, parts):
    plain = export_file(ctx, fmt)
    empty = export_file(ctx, fmt, {})
    with ZipFile(plain["path"]) as before, ZipFile(empty["path"]) as after:
        assert all(before.read(part) == after.read(part) for part in parts)


@pytest.mark.parametrize("bad", [
    [], "blue", {"accent_color": "#GG0000"}, {"accent_color": "red"}, {"body_color": 123},
    {"font_size": True}, {"font_size": float("nan")}, {"font_size": float("inf")}, {"font_size": 8}, {"font_size": 33},
    {"chart_palette": []}, {"chart_palette": ["#FF0000"] * 11}, {"chart_palette": ["#000000", "#BADHEX"]},
    {"font_family": "Arial; color:red"}, {"font_family": "../../font"}, {"font_family": "https://font.test"},
    {"font_family": "<script>"}, {"font_family": " "}, {"template_path": "/etc/passwd"},
])
def test_invalid_style_is_rejected_before_writing_artifacts(ctx, bad):
    with pytest.raises(InvalidArguments):
        export_file(ctx, "docx", bad)
    assert not list(ctx.project.root.glob("artifacts/*export*"))


def test_chart_palette_changes_picture_without_changing_numbers_and_does_not_leak(tmp_path):
    frame = pd.DataFrame({"group": ["A", "B", "A", "B", "A", "C"]})
    before = charts.draw(tmp_path, "bar", frame, "group", stem="before")
    styled = charts.draw(tmp_path, "bar", frame, "group", stem="styled", style=STYLE)
    after = charts.draw(tmp_path, "bar", frame, "group", stem="after")
    assert before["data"] == styled["data"] == after["data"]
    svg = Path(styled["chart_path"]).read_text().lower()
    assert "#136f63" in svg and "#f4a261" in svg and "#123456" in svg and "#243746" in svg
    assert "#5b4ca3" not in svg and "#b6a6ee" not in svg
    default = Path(after["chart_path"]).read_text().lower()
    assert "#5b4ca3" in default and "#b6a6ee" in default and "#136f63" not in default
    assert Path(styled["png_path"]).stat().st_size > 1000


def test_report_forwards_style_to_office_exports_and_all_generated_charts(ctx, tmp_path):
    source = tmp_path / "sample.csv"
    pd.DataFrame({"income": [1000 + i * 123.45 for i in range(60)],
                  "wellbeing": [30 + i * .81 for i in range(60)],
                  "region": ["North", "South", "West"] * 20}).to_csv(source, index=False)
    ctx.project.allow.allow(source)
    dataset = registry.call("import_dataset", {"path": str(source)}, ctx)["dataset_id"]
    result = registry.call("report_dataset", {"dataset_id": dataset, "format": "all", "style": STYLE}, ctx)
    assert result["style_applied"] == STYLE and result["charts"] >= 3
    assert result["warnings"] == []
    assert docx.Document(result["docx_path"]).styles["Normal"].font.name == "Arial"
    assert str(Presentation(result["pptx_path"]).slides[0].background.fill.fore_color.rgb) == "FEFCF4"
    all_charts = [path for path in ctx.project.root.glob("artifacts/**/*.svg") if "_profile_" not in str(path)]
    assert len(all_charts) == result["charts"]
    for chart in all_charts:
        svg = chart.read_text().lower()
        assert "#136f63" in svg and "#5b4ca3" not in svg and "#b6a6ee" not in svg
    html = Path(result["report_path"]).read_text()
    assert "--bg:#FEFCF4" in html and 'font-family:"Arial"' in html and "h1,h2 {color:#123456}" in html
    assert ".chart svg text { fill:#cfc6ef }" not in html


def test_normalized_style_does_not_change_user_preferences():
    original = {"accent_color": "#aabbcc", "font_family": " Arial ", "chart_palette": ["#1122aa"]}
    assert normalize_style(original) == {"accent_color": "#AABBCC", "font_family": "Arial", "chart_palette": ["#1122AA"]}
    assert original["accent_color"] == "#aabbcc" and original["chart_palette"] == ["#1122aa"]
