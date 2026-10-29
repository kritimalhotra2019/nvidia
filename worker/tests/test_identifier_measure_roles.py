"""Keep genuine numeric measurements usable after cleaning changes uniqueness."""
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, registry
from squipi_worker.security import AllowList
from squipi_worker.tools.datasets import analyse_column


@pytest.mark.parametrize("name,values", [
    ("monthly_income_inr", [18000 + 137 * i for i in range(48)]),
    ("annual_revenue_usd", [100000 + 1700 * i for i in range(48)]),
    ("age_years", list(range(18, 66))),
    ("weight_kg", list(range(40, 88))),
    ("height_cm", list(range(130, 178))),
    ("travel_time_minutes", list(range(30, 78))),
    ("weekly_energy_kwh", [100.25 + 1.17 * i for i in range(48)]),
    ("visits_count", list(range(1, 49))),
    ("number_of_children", [1, 2, 3, 4] * 12),
    ("Number of visits", list(range(1, 49))),
    ("number-of-patients", [10 + i for i in range(48)]),
])
def test_measurements_are_not_identifiers_because_values_are_distinct(name, values):
    series = pd.Series(values)
    info = analyse_column(name, series, len(series))
    assert info["role"] == "numeric"
    assert info["stats"]["mean"] == pytest.approx(float(series.mean()))
    assert "histogram" in info


@pytest.mark.parametrize("name", [
    "household_id", "deviceid", "simid", "subscriberid", "device_phone_num",
    "phone_number", "account_number", "record_code", "case_id",
    "number_of_phone_ids", "Number of household_id",
])
def test_explicit_identifier_evidence_stays_protected(name):
    # Repeated and entirely distinct values must both retain the semantic role.
    for values in [[100000000000 + i for i in range(48)], [100000000000 + i % 4 for i in range(48)]]:
        info = analyse_column(name, pd.Series(values), len(values))
        assert info["role"] == "id"
        assert "stats" not in info and "histogram" not in info


def test_explicit_id_override_and_nonnumeric_key_evidence_remain_available():
    values = pd.Series(range(30, 78))
    assert analyse_column("opaque_field", values, len(values), override_role="id")["role"] == "id"
    codes = pd.Series([f"HX-{i:04d}" for i in range(48)])
    assert analyse_column("record", codes, len(codes))["role"] == "id"
    assert analyse_column("household_id", codes.astype("string"), len(codes))["role"] == "id"


def test_deduplication_does_not_remove_measurement_from_requested_report(tmp_path):
    frame = pd.DataFrame({
        "household_id": [f"H{i:03d}" for i in range(48)],
        "district": ["North"] * 24 + ["South"] * 24,
        "annual_revenue_usd": [10000.0 + 137 * i for i in range(47)] + [np.nan],
    })
    duplicated = pd.concat([frame, frame.iloc[[0]]], ignore_index=True)
    source = tmp_path / "businesses.csv"
    duplicated.to_csv(source, index=False)
    original = source.read_bytes()
    project = Project(tmp_path / "project", AllowList())
    project.allow.allow(source)
    ctx = JobContext(project, lambda *_: None, lambda: False)
    imported = registry.call("import_dataset", {"path": str(source)}, ctx)
    before = analyse_column("annual_revenue_usd", duplicated.annual_revenue_usd, len(duplicated))
    cleaned = registry.call("clean_dataset", {"dataset_id": imported["dataset_id"], "pii": "keep", "steps": ["drop_duplicates"]}, ctx)
    profile = registry.call("profile_dataset", {"dataset_id": cleaned["dataset_id"]}, ctx)
    measure = next(c for c in profile["columns"] if c["name"] == "annual_revenue_usd")
    assert before["role"] == measure["role"] == "numeric"
    assert profile["rows"] == 48 and measure["missing"] == 1 and measure["unique"] == 47
    assert measure["stats"]["mean"] == pytest.approx(float(frame.annual_revenue_usd.mean()))
    report = registry.call("report_dataset", {"dataset_id": cleaned["dataset_id"], "outcome": "annual_revenue_usd", "format": "docx", "title": "Business research"}, ctx)
    assert source.read_bytes() == original
    docx = Path(report["docx_path"])
    doc = Document(docx)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert docx.is_file() and "annual_revenue_usd" in text
    assert "household_id" in text  # column guide may name an ID without averaging it
