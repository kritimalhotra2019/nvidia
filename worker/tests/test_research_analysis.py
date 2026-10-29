import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document

from squipi_worker import tools
from squipi_worker.project import Project
from squipi_worker.security import AllowList
from squipi_worker.registry import JobContext,registry,InvalidArguments
from squipi_worker.execution_cache import ExecutionCache
from squipi_worker.tools import research_analysis


def setup(tmp_path,frame,fmt='csv',name='experiment'):
    p=Project(tmp_path/'project',AllowList());ctx=JobContext(p,lambda *a:None,lambda:False)
    path=tmp_path/(name+'.'+fmt)
    if fmt=='csv':frame.to_csv(path,index=False)
    elif fmt=='xlsx':frame.to_excel(path,index=False)
    else:frame.to_parquet(path,index=False)
    p.allow.allow(path);ds=registry.call('import_dataset',{'path':str(path)},ctx)['dataset_id']
    return ctx,ds,path


@pytest.mark.parametrize('fmt',['csv','xlsx','parquet'])
def test_unrelated_formats_and_schema_have_same_numbers(tmp_path,fmt):
    ctx,ds,path=setup(tmp_path,pd.DataFrame({'item_id':['K1','K2','K3','K4','K5','K6'],
       'branch':['North']*3+['South']*3,'sales':[10,20,None,30,40,50],'quality':[1,2,3,4,5,6]}),fmt)
    before=path.read_bytes()
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['sales','quality'],'group_by':['branch'],'units':{'sales':'INR','quality':'score'},'charts':False},ctx)
    sales=next(x for x in r['statistics'] if x['scope']=='overall' and x['variable']=='sales')
    assert sales['n']==5 and sales['missing_or_invalid']==1
    assert sales['mean']==30 and sales['sd']==pytest.approx(np.sqrt(250)) and sales['median']==30
    north=next(x for x in r['statistics'] if x['group']=='North' and x['variable']=='sales')
    assert north['mean']==15 and north['n']==2 and north['missing_or_invalid']==1
    assert path.read_bytes()==before


def test_mixed_numeric_values_and_identifiers_are_not_silently_averaged(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'device_id':[11,12,13,14],'vibration':['2','bad','4',None]}))
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['vibration'],'charts':False},ctx)
    assert r['statistics'][0]['mean']==3 and r['statistics'][0]['missing_or_invalid']==2
    with pytest.raises(InvalidArguments,match='Identifier'):
        registry.call('summarize_research_data',{'dataset_id':ds,'columns':['device_id']},ctx)


def test_counts_are_measures_but_explicit_identifier_roles_are_respected(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'number_of_children':[1,2,3],
        'account_number':[1001,1002,1003],'serial':[11,12,13]}))
    ctx.project.dataset(ds).overrides={'serial':{'role':'id'}}
    r=registry.call('summarize_research_data',{'dataset_id':ds,'charts':False},ctx)
    assert [row['variable'] for row in r['statistics']]==['number_of_children']
    assert r['statistics'][0]['mean']==2
    with pytest.raises(InvalidArguments,match='identifiers'):
        registry.call('summarize_research_data',{'dataset_id':ds,'columns':['number_of_children'],
          'paired_changes':[{'baseline':'number_of_children','followup':'serial','label':'change','unit':'count'}]},ctx)


@pytest.mark.parametrize('grouping',[{}, {'group_by':[]}])
def test_overall_aggregation_does_not_require_a_dummy_group(tmp_path,grouping):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'shipment_id':range(4),'cost':[10,20,None,30]}))
    r=registry.call('aggregate',{'dataset_id':ds,**grouping,'metrics':[
        {'column':'cost','agg':'count'},{'column':'cost','agg':'mean'},{'column':'cost','agg':'std'}]},ctx)
    assert r['table']=={'columns':['count_cost','mean_cost','std_cost'],'rows':[[3,20,10]]}
    assert r['n_groups']==1 and r['group_by']==[] and r['chart_path'] is None


def test_linkage_exclusions_reach_the_actual_brief_after_cleaning(tmp_path):
    ctx,left,_=setup(tmp_path,pd.DataFrame({'unit_id':[1,2,3,4],'before':[4,5,6,7]}),name='baseline')
    path=tmp_path/'endline.csv';pd.DataFrame({'unit_id':[2,3,4,5,6],'after':[7,8,9,10,11]}).to_csv(path,index=False)
    ctx.project.allow.allow(path);right=registry.call('import_dataset',{'path':str(path)},ctx)['dataset_id']
    joined=registry.call('combine_datasets',{'dataset_ids':[left,right],'mode':'join','how':'inner','keys':['unit_id']},ctx)
    clean=registry.call('clean_dataset',{'dataset_id':joined['dataset_id'],'pii':'keep','drop_columns':['unit_id'],'steps':[]},ctx)
    analysis=registry.call('summarize_research_data',{'dataset_id':clean['dataset_id'],'columns':['before','after'],'charts':False},ctx)
    brief=registry.call('write_research_brief',{'analysis_id':analysis['analysis_id'],'title':'Matched sample'},ctx)
    doc=Document(brief['path'])
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==[
        ['Source','Input rows','Matched rows','Unmatched rows'],['baseline.csv','4','3','1'],['endline.csv','5','3','2']]
    assert brief['verification']['passed']
    Path(joined['provenance_path']).write_text('{}')
    with pytest.raises(InvalidArguments,match='chart changed'):
        registry.call('write_research_brief',{'analysis_id':analysis['analysis_id'],'title':'Stale audit'},ctx)


def test_paired_change_uses_complete_pairs_and_explicit_units(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'station':['A']*4,'before':[10,20,None,40],'after':[13,18,30,None]}))
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['before'],'group_by':['station'],
        'paired_changes':[{'baseline':'before','followup':'after','label':'Change','unit':'percentage points'}],'charts':False},ctx)
    change=next(x for x in r['statistics'] if x['variable']=='Change' and x['scope']=='overall')
    assert change['n']==2 and change['missing_or_invalid']==2 and change['mean']==.5 and change['unit']=='percentage points'
    assert change['sd']==pytest.approx(np.sqrt(12.5))


def test_targeted_distributions_and_changes_do_not_add_unrequested_measures(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'age':[30,40,50],'answer':['Yes','No','Yes'],'before':[10,20,30],'after':[12,23,34]}))
    categories=registry.call('summarize_research_data',{'dataset_id':ds,'categorical_columns':['answer'],'charts':False},ctx)
    assert categories['statistics']==[] and len(categories['frequencies'])==2
    changes=registry.call('summarize_research_data',{'dataset_id':ds,'columns':[],
        'paired_changes':[{'baseline':'before','followup':'after','label':'Change','unit':'days'}],'charts':False},ctx)
    assert [r['variable'] for r in changes['statistics']]==['Change']
    assert changes['statistics'][0]['mean']==3


def test_weighted_descriptives_are_explicit_about_exclusions_and_inference(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'income':[10,20,100,None],'weight':[1,3,0,1]}))
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['income'],'weight_column':'weight','charts':False},ctx)
    row=r['statistics'][0]
    assert row['weighted_mean']==17.5 and row['weighted_n']==2 and row['weighted_excluded']==2
    data=research_analysis.load_analysis(ctx.project,r['analysis_id'])
    assert any('No survey-design-adjusted' in x for x in data['limitations'])


def test_categorical_only_unicode_and_group_named_overall(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'区':['Overall','Overall','B','B'],'回答':['हाँ','हाँ','नहीं',None]}),name='service-feedback')
    r=registry.call('summarize_research_data',{'dataset_id':ds,'categorical_columns':['回答'],'group_by':['区']},ctx)
    total=[x for x in r['frequencies'] if x['scope']=='overall']
    yes=next(x for x in total if x['category']=='हाँ')
    assert yes['n']==2 and yes['nonmissing_denominator']==3 and yes['percent_nonmissing']==pytest.approx(200/3)
    assert len(r['chart_paths'])==1
    brief=registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Service feedback','filename':'Feedback.docx'},ctx)
    assert brief['verification']['passed'] and Path(brief['path']).name=='Feedback.docx'
    assert len(Document(brief['path']).inline_shapes)==1


def test_brief_retains_all_requested_values_and_detects_changes(tmp_path):
    ctx,ds,source=setup(tmp_path,pd.DataFrame({'region':['West']*6+['East']*6,'delivery_days':range(12),'cost':range(100,112)}))
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['delivery_days','cost'],'group_by':['region'],'units':{'delivery_days':'days','cost':'USD'}},ctx)
    brief=registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Logistics findings'},ctx)
    assert brief['verification']['all_requested_table_cells_match'] and brief['charts_embedded']==2
    d=Document(brief['path']);assert len(d.tables)==2 and len(d.inline_shapes)==2
    assert all(len(t.rows)==4 for t in d.tables)
    # The generic exporter must not round 105.5 to 106 merely because it is >=100.
    assert d.tables[1].cell(1,3).text=='105.5'
    assert d.tables[1].cell(1,0).text=='Overall (all rows)'
    source.write_text('changed source')
    with pytest.raises(InvalidArguments,match='Source data changed'):
        registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Stale'},ctx)


def test_modified_chart_and_analysis_are_rejected(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'measure':range(12)}))
    r=registry.call('summarize_research_data',{'dataset_id':ds,'columns':['measure']},ctx)
    Path(r['chart_paths'][0]).with_suffix('.png').write_bytes(b'changed')
    with pytest.raises(InvalidArguments,match='chart changed'):
        registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Bad asset'},ctx)
    Path(r['results_path']).write_text('{}')
    with pytest.raises(InvalidArguments,match='Saved analysis changed'):
        registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Bad manifest'},ctx)


@pytest.mark.parametrize('how,rows,left_missing,right_missing',[('inner',2,1,1),('left',3,1,1),('right',3,1,1),('outer',4,1,1)])
def test_join_semantics_work_for_general_keys(tmp_path,how,rows,left_missing,right_missing):
    ctx,a,_=setup(tmp_path,pd.DataFrame({'sku':['A','B','C'],'cost':[1,2,3]}),name='costs')
    path=tmp_path/'stock.parquet';pd.DataFrame({'sku':['B','C','D'],'stock':[10,20,30]}).to_parquet(path,index=False)
    ctx.project.allow.allow(path);b=registry.call('import_dataset',{'path':str(path)},ctx)['dataset_id']
    args={'dataset_ids':[a,b],'mode':'join','keys':['sku'],'how':how}
    if how=='left':args['require_all_left']=False
    r=registry.call('combine_datasets',args,ctx)
    assert r['rows']==rows and r['join_audit']['unmatched_left_rows']==left_missing and r['join_audit']['unmatched_right_rows']==right_missing
    assert r['join_audit']['matched_left_rows']==2
    with pytest.raises(InvalidArguments,match='matched cohort'):
        registry.call('combine_datasets',{'dataset_ids':[a,b],'mode':'join','keys':['sku']},ctx)


def test_cardinality_and_missing_keys_still_block_unsafe_joins(tmp_path):
    ctx,a,_=setup(tmp_path,pd.DataFrame({'key':['A','A'],'x':[1,2]}))
    path=tmp_path/'lookup.csv';pd.DataFrame({'key':['A','A'],'y':[3,4]}).to_csv(path,index=False)
    ctx.project.allow.allow(path);b=registry.call('import_dataset',{'path':str(path)},ctx)['dataset_id']
    with pytest.raises(InvalidArguments,match='cardinality'):
        registry.call('combine_datasets',{'dataset_ids':[a,b],'mode':'join','keys':['key'],'how':'inner'},ctx)


def test_appraisal_time_zero_and_power_have_independent_known_answers(tmp_path):
    ctx=JobContext(Project(tmp_path/'p',AllowList()),lambda *a:None,lambda:False)
    r=registry.call('economic_appraisal',{'costs':[100,0],'benefits':[0,110],'discount_rate':.1,'currency':'INR','price_basis':'constant 2026 prices','first_period':0},ctx)
    assert r['present_value_costs']==100 and r['present_value_benefits']==pytest.approx(100)
    assert r['net_present_value']==pytest.approx(0) and r['benefit_cost_ratio']==pytest.approx(1)
    p=registry.call('plan_research_power',{'difference':.5,'sd':1,'alpha':.05,'power':.8},ctx)
    assert p['per_arm']==63 and p['total']==126
    p=registry.call('plan_research_power',{'difference':.5,'sd':1,'alpha':.05,'power':.8,'design_effect':2,'attrition':.2},ctx)
    assert p['per_arm']==157


def test_new_analysis_and_brief_reuse_survive_process_recreation(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'outcome':range(8)}))
    args={'dataset_id':ds,'columns':['outcome'],'charts':False}
    cache=ExecutionCache(ctx.project);scope='test_objective_123'
    r=cache.call('summarize_research_data',args,ctx,scope,lambda:registry.call('summarize_research_data',args,ctx))
    reopened=ExecutionCache(ctx.project)
    reused=reopened.call('summarize_research_data',args,ctx,scope,lambda:pytest.fail('analysis was rerun'))
    assert reused['reused_saved_result'] and reused['analysis_id']==r['analysis_id']
    a={'analysis_id':r['analysis_id'],'title':'Persisted results'}
    brief=cache.call('write_research_brief',a,ctx,scope,lambda:registry.call('write_research_brief',a,ctx))
    reused=reopened.call('write_research_brief',a,ctx,scope,lambda:pytest.fail('report was rerun'))
    assert reused['path']==brief['path'] and reused['reused_saved_result']


def test_modified_office_chart_invalidates_the_analysis_cache(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'outcome':range(8)}))
    args={'dataset_id':ds,'columns':['outcome']};scope='chart_cache_123'
    cache=ExecutionCache(ctx.project)
    saved=cache.call('summarize_research_data',args,ctx,scope,lambda:registry.call('summarize_research_data',args,ctx))
    Path(saved['chart_png_paths'][0]).write_bytes(b'altered image')
    assert cache.lookup('summarize_research_data',args,scope) is None
    fresh=cache.call('summarize_research_data',args,ctx,scope,lambda:registry.call('summarize_research_data',args,ctx))
    assert fresh['analysis_id']!=saved['analysis_id']


def test_explicit_defaults_reuse_the_same_analysis_without_merging_different_measures(tmp_path):
    ctx,ds,_=setup(tmp_path,pd.DataFrame({'value':range(8),'other':range(10,18)}))
    args={'dataset_id':ds,'columns':['value']};scope='equivalent_analysis_123';cache=ExecutionCache(ctx.project)
    saved=cache.call('summarize_research_data',args,ctx,scope,lambda:registry.call('summarize_research_data',args,ctx))
    equivalent={**args,'charts':True,'allow_sample':False,'units':{},'group_by':[],'paired_changes':[],'categorical_columns':[]}
    hit=cache.call('summarize_research_data',equivalent,ctx,scope,lambda:pytest.fail('Equivalent analysis was repeated'))
    assert hit['analysis_id']==saved['analysis_id'] and hit['reused_saved_result']
    assert cache.lookup('summarize_research_data',{**args,'columns':['other']},scope) is None


def test_changed_ancestor_invalidates_downstream_saved_work(tmp_path):
    ctx,ds,source=setup(tmp_path,pd.DataFrame({'name':['A','B'],'measure':[10,20]}))
    clean=registry.call('clean_dataset',{'dataset_id':ds,'pii':'keep','drop_columns':['name'],'steps':[]},ctx)
    args={'dataset_id':clean['dataset_id'],'columns':['measure'],'charts':False};cache=ExecutionCache(ctx.project)
    r=cache.call('summarize_research_data',args,ctx,'objective_ancestor',lambda:registry.call('summarize_research_data',args,ctx))
    assert cache.lookup('summarize_research_data',args,'objective_ancestor') is not None
    source.write_text('name,measure\nA,1000\nB,2000\n')
    assert cache.lookup('summarize_research_data',args,'objective_ancestor') is None
    with pytest.raises(InvalidArguments,match='upstream source'):
        registry.call('write_research_brief',{'analysis_id':r['analysis_id'],'title':'Stale descendant'},ctx)


def test_invalid_numbers_and_oversized_tables_do_not_silently_export(tmp_path):
    ctx=JobContext(Project(tmp_path/'p',AllowList()),lambda *a:None,lambda:False)
    with pytest.raises(InvalidArguments,match='finite'):
        registry.call('plan_research_power',{'difference':float('nan'),'sd':1,'alpha':.05,'power':.8},ctx)
    with pytest.raises(InvalidArguments):
        registry.call('export_document',{'format':'docx','title':'Too much','sections':[{'heading':'Table','table':{'columns':['x'],'rows':[[i] for i in range(201)]}}]},ctx)


def test_skills_are_read_only_guidance_and_reference_existing_tools(tmp_path):
    ctx=JobContext(Project(tmp_path/'p',AllowList()),lambda *a:None,lambda:False)
    skills=registry.call('read_research_skill',{},ctx)['skills']
    assert {s['name'] for s in skills} >= {'data-integrity','descriptive-research','survey-design','private-text','web-evidence'}
    for item in skills:
        skill=registry.call('read_research_skill',{'name':item['name']},ctx)
        assert skill['references'] and skill['instructions']
        assert all(registry.get(name) for name in skill['tools'])
