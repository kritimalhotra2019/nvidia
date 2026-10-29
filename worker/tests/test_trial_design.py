"""Trial-design tools: seeded randomisation on a synthetic frame (units in schools, two strata) and the pre-analysis plan document."""
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from squipi_worker import tools  # noqa: F401 - registers the tools
from squipi_worker.project import Project
from squipi_worker.security import AllowList
from squipi_worker.registry import JobContext, registry, InvalidArguments

N = 400
SCHOOLS = 40


def _frame(rng):
    school = np.repeat(np.arange(SCHOOLS), N // SCHOOLS)
    region = np.where(school < SCHOOLS // 2, "north", "south")
    urban = (school % 2 == 0).astype(int)
    return pd.DataFrame({"pupil_id": [f"p{i:04d}" for i in range(N)], "school": school, "region": region, "urban": urban,
                         "age": rng.integers(6, 12, N), "baseline_score": rng.normal(50, 10, N), "female": rng.integers(0, 2, N),
                         "email": [f"u{i}@example.org" for i in range(N)]})


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    rng = np.random.default_rng(3)
    base = tmp_path_factory.mktemp("trial")
    p = Project(base / "project", AllowList())
    ctx = JobContext(p, lambda *a: None, lambda: False)
    frame = _frame(rng)
    path = base / "pupils.csv"
    frame.to_csv(path, index=False)
    p.allow.allow(path)
    ds = registry.call("import_dataset", {"path": str(path)}, ctx)["dataset_id"]
    dup = pd.concat([frame.head(5), frame.head(5)], ignore_index=True)
    dpath = base / "dup.csv"
    dup.to_csv(dpath, index=False)
    p.allow.allow(dpath)
    dup_id = registry.call("import_dataset", {"path": str(dpath)}, ctx)["dataset_id"]
    small = pd.DataFrame({"uid": range(9), "g": list("aaaaaabbc"), "x": rng.normal(size=9)})
    spath = base / "small.csv"
    small.to_csv(spath, index=False)
    p.allow.allow(spath)
    small_id = registry.call("import_dataset", {"path": str(spath)}, ctx)["dataset_id"]
    return {"ctx": ctx, "ds": ds, "dup": dup_id, "small": small_id, "frame": frame, "base": base}


def _assignment(r):
    return pd.read_csv(r["assignment_path"])


def test_same_seed_reproduces_and_different_seed_differs(world):
    ctx, ds = world["ctx"], world["ds"]
    args = {"dataset_id": ds, "id_column": "pupil_id", "seed": 7, "strata": ["region"], "balance_columns": ["baseline_score", "female"]}
    a = registry.call("randomize_assignment", args, ctx)
    b = registry.call("randomize_assignment", args, ctx)
    c = registry.call("randomize_assignment", {**args, "seed": 8}, ctx)
    assert _assignment(a)["arm"].tolist() == _assignment(b)["arm"].tolist()
    assert _assignment(a)["arm"].tolist() != _assignment(c)["arm"].tolist()
    assert a["seed"] == 7 and a["dataset_id"] != b["dataset_id"] and a["dataset_name"] == "pupils assignment.csv"
    assert ctx.project.dataset(a["dataset_id"]).derived_from == ds
    assert list(_assignment(a).columns) == ["pupil_id", "region", "arm"]
    assert "seed 7" in a["procedure"] and "blocks of 2" in a["procedure"] and a["rerandomisation"] is None
    assert Path(a["chart_path"]).exists() and Path(a["chart_path"]).with_suffix(".png").exists() and Path(a["csv_path"]).exists()
    json.loads(json.dumps(a))


def test_stratified_counts_are_exact_for_equal_allocation(world):
    ctx, ds = world["ctx"], world["ds"]
    r = registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 11, "strata": ["region", "urban"]}, ctx)
    assert r["counts_per_arm"] == {"treatment": N // 2, "control": N // 2}
    assert len(r["counts_per_stratum_arm"]) == 4
    for row in r["counts_per_stratum_arm"]:
        assert row["treatment"] == row["control"] == row["total"] // 2 == 50
    a = _assignment(r)
    assert a.groupby(["region", "urban"])["arm"].value_counts().eq(50).all()
    # unequal allocation in blocks of 5 (0.6 / 0.4): each stratum within one block of target
    u = registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 11, "arms": ["A", "B"], "allocation": [0.6, 0.4], "strata": ["region"]}, ctx)
    assert u["block_size"] == 5 and u["counts_per_arm"] == {"A": 240, "B": 160}


def test_cluster_randomisation_keeps_clusters_intact(world):
    ctx, ds = world["ctx"], world["ds"]
    r = registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 5, "cluster_column": "school", "strata": ["region"],
                                               "balance_columns": ["baseline_score", "age"]}, ctx)
    a = _assignment(r)
    assert list(a.columns) == ["pupil_id", "region", "school", "arm"]
    assert (a.groupby("school")["arm"].nunique() == 1).all()
    assert r["clusters_per_arm"] == {"treatment": SCHOOLS // 2, "control": SCHOOLS // 2} and r["n_units"] == SCHOOLS
    assert all(row["clusters"] == SCHOOLS // 2 for row in r["counts_per_stratum_arm"])
    assert all(j["se_kind"].startswith("cluster-robust") for j in r["joint_test"])
    assert "All rows of a cluster share its arm" in r["procedure"]
    with pytest.raises(InvalidArguments, match="constant within a cluster"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 5, "cluster_column": "school", "strata": ["female"]}, ctx)


def test_balance_table_columns_and_joint_test(world):
    ctx, ds = world["ctx"], world["ds"]
    r = registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 2, "balance_columns": ["baseline_score", "female", "age"]}, ctx)
    assert {row["column"] for row in r["balance"]} == {"baseline_score", "female", "age"}
    for row in r["balance"]:
        assert set(row) >= {"column", "comparison", "mean_treatment", "n_treatment", "mean_control", "n_control", "difference", "std_diff", "p", "flag"}
        assert row["comparison"] == "treatment vs control" and row["n_treatment"] + row["n_control"] == N
        assert row["flag"] in ("balanced", "imbalanced") and 0 <= row["p"] <= 1
    assert len(r["joint_test"]) == 1 and r["joint_test"][0]["df"] == 3 and 0 <= r["joint_test"][0]["p"] <= 1 and r["joint_test"][0]["se_kind"] == "HC1 robust"
    assert any("Balance:" in s for s in r["interpretation"])
    bal = pd.read_csv(r["csv_path"])
    assert len(bal) == 3


def test_rerandomisation_reports_draws_and_improves_balance(world):
    ctx, ds = world["ctx"], world["ds"]
    args = {"dataset_id": ds, "id_column": "pupil_id", "seed": 21, "balance_columns": ["baseline_score", "female", "age"]}
    one = registry.call("randomize_assignment", args, ctx)
    first_d = max(abs(row["std_diff"]) for row in one["balance"])
    r = registry.call("randomize_assignment", {**args, "rerandomize_max": 30}, ctx)
    rr = r["rerandomisation"]
    assert rr["draws"] == 30 and 1 <= rr["chosen_draw"] <= 30
    assert rr["first_draw_max_abs_std_diff"] == pytest.approx(first_d, abs=1e-3)      # the first draw is the plain seeded draw
    assert rr["chosen_max_abs_std_diff"] <= rr["first_draw_max_abs_std_diff"]
    assert max(abs(row["std_diff"]) for row in r["balance"]) == pytest.approx(rr["chosen_max_abs_std_diff"], abs=1e-3)
    assert any("Re-randomisation" in n for n in r["notes"]) and "re-randomised" in " ".join(r["interpretation"])
    again = registry.call("randomize_assignment", {**args, "rerandomize_max": 30}, ctx)
    assert _assignment(r)["arm"].tolist() == _assignment(again)["arm"].tolist()
    with pytest.raises(InvalidArguments, match="balance_columns"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 21, "rerandomize_max": 5}, ctx)


def test_refusals(world):
    ctx, ds = world["ctx"], world["ds"]
    with pytest.raises(InvalidArguments, match="seed"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id"}, ctx)
    with pytest.raises(InvalidArguments, match="not unique"):
        registry.call("randomize_assignment", {"dataset_id": world["dup"], "id_column": "pupil_id", "seed": 1}, ctx)
    with pytest.raises(InvalidArguments, match="sum to 1"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 1, "allocation": [0.7, 0.4]}, ctx)
    with pytest.raises(InvalidArguments, match="personal data"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 1, "strata": ["email"]}, ctx)
    with pytest.raises(InvalidArguments, match="personal data"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 1, "balance_columns": ["email"]}, ctx)
    # strata cells smaller than the number of arms: refused, or warned with allow_small_strata
    with pytest.raises(InvalidArguments, match="2 strata have fewer units than arms"):
        registry.call("randomize_assignment", {"dataset_id": world["small"], "id_column": "uid", "seed": 1, "strata": ["g"], "arms": ["a", "b", "c"]}, ctx)
    w = registry.call("randomize_assignment", {"dataset_id": world["small"], "id_column": "uid", "seed": 1, "strata": ["g"], "arms": ["a", "b", "c"], "allow_small_strata": True}, ctx)
    assert any("Warning" in n and "fewer units than arms" in n for n in w["notes"]) and sum(w["counts_per_arm"].values()) == 9
    with pytest.raises(InvalidArguments, match="differ from the id and the strata"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 1, "cluster_column": "school", "strata": ["school"]}, ctx)
    with pytest.raises(InvalidArguments, match="block"):
        registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 1, "allocation": [0.123, 0.877]}, ctx)


def test_reproduce_snippet_recreates_the_arm_column(world):
    ctx, ds = world["ctx"], world["ds"]
    for extra in ({}, {"strata": ["region", "urban"], "cluster_column": "school"}, {"balance_columns": ["baseline_score"], "rerandomize_max": 12, "strata": ["region"]}):
        r = registry.call("randomize_assignment", {"dataset_id": ds, "id_column": "pupil_id", "seed": 99, **extra}, ctx)
        snippet = r["reproduce_python"]
        assert "import numpy as np" in snippet and "squipi" not in snippet
        assert Path(r["reproduce_path"]).read_text() == snippet
        ns = {"__name__": "reproduce_check", "__builtins__": __builtins__}
        exec(compile(snippet, "reproduce_assignment.py", "exec"), ns)  # the snippet only imports numpy and pandas and reads the frame
        expected = _assignment(r).set_index("pupil_id")["arm"]
        got = ns["frame"].set_index("pupil_id")["arm"]
        assert got.loc[expected.index].tolist() == expected.tolist()


# ---------- pre_analysis_plan

PLAN = {"title": "Reading tutors pilot", "background": "Tutoring in grade 3 classrooms in two regions.", "treatment": "Twice-weekly small-group tutoring for one term.",
        "primary_outcomes": [{"name": "Reading score", "column": "endline_score", "direction": "higher is better", "unit": "points"}],
        "secondary_outcomes": [{"name": "Attendance", "column": "attendance_pct", "direction": "higher is better", "unit": "percent"}],
        "randomisation": {"level": "school", "strata": ["region"], "arms": ["treatment", "control"], "allocation": [0.5, 0.5], "seed": 5, "method": "permuted blocks within strata"},
        "sample": {"n": 400, "clusters": 40, "expected_attrition_pct": 10, "eligibility": "pupils enrolled at baseline"},
        "hypotheses": ["Tutoring raises reading scores."],
        "estimation": {"estimator": "ANCOVA", "covariates": ["baseline_score"], "clustering": "school", "multiple_testing": "BH", "heterogeneity": ["female"]},
        "power": {"mde": 5, "sd": 10, "alpha": 0.05, "power": 0.8, "icc": 0.1, "cluster_size": 10},
        "analysis_data_rules": ["Exclude pupils who left the school before the intervention started.", "Winsorise scores at the 1st and 99th percentiles."],
        "timeline": {"baseline": "2027-01", "intervention": "2027-02 to 2027-06", "endline": "2027-07", "analysis": "2027-09", "stopping_rules": "none"}}


def test_pre_analysis_plan_docx_headings_power_and_sidecar(world):
    from docx import Document
    ctx = world["ctx"]
    r = registry.call("pre_analysis_plan", PLAN, ctx)
    assert Path(r["docx_path"]).name == "Reading tutors pilot pre-analysis plan.docx" and Path(r["docx_path"]).exists()
    doc = Document(r["docx_path"])
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    expected = ["1. Background", "2. Hypotheses", "3. Outcomes", "4. Randomisation", "5. Sample and power", "6. Estimation", "7. Multiple testing",
                "8. Heterogeneity", "9. Data rules", "10. Timeline and stopping rules", "11. Deviations log", "12. Registration checklist"]
    assert headings == expected == r["sections"]
    assert any(p.text == "Reading tutors pilot" for p in doc.paragraphs)
    # power numbers are plan_research_power's for the same inputs (design effect 1 + (10 - 1) × 0.1 = 1.9, attrition 0.1)
    ref = registry.call("plan_research_power", {"difference": 5, "sd": 10, "alpha": 0.05, "power": 0.8, "design_effect": 1.9, "attrition": 0.1}, ctx)
    assert r["power"]["per_arm"] == ref["per_arm"] and r["power"]["total"] == ref["total"]
    assert r["power"]["independent_complete_cases_per_arm"] == ref["independent_complete_cases_per_arm"] == 63
    assert r["power"]["inputs"]["design_effect"] == pytest.approx(1.9)
    cells = {c.text for t in doc.tables for row in t.rows for c in row.cells}
    assert f"{ref['per_arm']:,}" in cells and f"{ref['total']:,}" in cells and "endline_score" in cells and "AEA" in " ".join(p.text for p in doc.paragraphs)
    assert r["checklist"] == [] and "0 item(s) still missing" in r["summary_text"]
    side = json.loads(Path(r["json_path"]).read_text())
    assert side["plan"] == PLAN == r["plan"] and side["power"] == r["power"] and side["checklist"] == [] and side["sections"] == expected
    json.loads(json.dumps(r))


def test_pre_analysis_plan_checklist_lists_missing_items(world):
    ctx = world["ctx"]
    r = registry.call("pre_analysis_plan", {"title": "Bare plan", "primary_outcomes": [{"name": "Income"}, {"name": "Savings"}]}, ctx)
    text = "\n".join(r["checklist"])
    for needle in ("secondary outcomes: none specified", "no attrition allowance", "no seed recorded", "no data column named", "no power calculation",
                   "multiple testing: no rule chosen", "2 primary outcomes", "data rules", "hypotheses: none stated", "timeline"):
        assert needle in text, needle
    assert r["power"] is None and "no power calculation" in r["summary_text"]
    from docx import Document
    doc = Document(r["docx_path"])
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "no attrition allowance" in body and "To do: Registry" in body
    # the document holds no invented numbers: the power table is absent, and the sidecar says so
    assert not any("Recruit per arm" in c.text for t in doc.tables for row in t.rows for c in row.cells)
    side = json.loads(Path(r["json_path"]).read_text())
    assert side["checklist"] == r["checklist"] and side["power"] is None
    with pytest.raises(InvalidArguments, match="distinct"):
        registry.call("pre_analysis_plan", {"title": "Dup", "primary_outcomes": [{"name": "A"}], "secondary_outcomes": [{"name": "A"}]}, ctx)
    with pytest.raises(InvalidArguments, match="sum to 1"):
        registry.call("pre_analysis_plan", {"title": "Alloc", "primary_outcomes": [{"name": "A"}], "randomisation": {"arms": ["t", "c"], "allocation": [0.7, 0.5]}}, ctx)


def test_tools_are_deterministic_and_registered():
    for name in ("randomize_assignment", "pre_analysis_plan"):
        t = registry.get(name)
        assert t is not None and t.needs_model is False
    assert "seed" in registry.get("randomize_assignment").params["required"]
