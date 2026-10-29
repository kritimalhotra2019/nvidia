from pathlib import Path
from zipfile import ZipFile
import hashlib
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import pytest
from squipi_worker import tools
from squipi_worker.project import Project
from squipi_worker.security import AllowList
from squipi_worker.registry import JobContext, registry, InvalidArguments
from squipi_worker.templates import library_action


def call(ctx, name, **args):
    return registry.call(name, args, ctx)


@pytest.fixture
def draft(tmp_path):
    ctx = JobContext(Project(tmp_path/'project', AllowList()), lambda *a: None, lambda: False)
    source = tmp_path/'template.pptx'
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(8.5), Inches(11)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    label = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(7.5), Inches(.5))
    label.text = '[Finding]'
    run = label.text_frame.paragraphs[0].runs[0]
    run.font.name, run.font.size = 'Calibri', Pt(18)
    run.font.color.rgb = RGBColor.from_string('24594D')
    panel = slide.shapes.add_shape(1, Inches(.5), Inches(2), Inches(7.5), Inches(5))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string('F2F4F6')
    slide.notes_slide.notes_text_frame.text = 'OLD_PRIVATE_NOTES'
    deck.save(source)
    ctx.project.allow.allow(source)
    template = library_action(ctx.project, {'action':'import', 'path':str(source)})['templates'][-1]
    result = call(ctx, 'begin_template_presentation', template_id=template['id'])
    call(ctx, 'inspect_presentation_draft', draft_id=result['draft_id'], revision=0, slide_number=1)
    return ctx, source, result, label.shape_id, panel.shape_id


def test_preserves_real_objects_and_style_without_mutating_source(draft):
    ctx, source, result, label, panel = draft
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    original = Presentation(result['draft_path'])
    edited = call(ctx,'edit_presentation_slide', draft_id=result['draft_id'], revision=0, slide_number=1,
                  text_edits=[{'shape_id':label,'text':'Mean income: 3,950'}])
    output = call(ctx,'finalize_template_presentation',draft_id=result['draft_id'],revision=1,title='Test')
    final = Presentation(output['path'])
    assert final.slide_width == Inches(8.5) and final.slide_height == Inches(11)
    assert len(final.slides[0].shapes) == len(original.slides[0].shapes)
    for a,b in zip(final.slides[0].shapes,original.slides[0].shapes):
        assert (a.left,a.top,a.width,a.height)==(b.left,b.top,b.width,b.height)
    run=final.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size==Pt(18) and str(run.font.color.rgb)=='24594D'
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    assert Presentation(result['draft_path']).slides[0].shapes[0].text=='[Finding]'
    with ZipFile(output['path']) as archive:
        assert not any(b'OLD_PRIVATE_NOTES' in archive.read(n) for n in archive.namelist() if n.endswith('.xml'))
    assert output['visual_review_required'] is True


def test_rejects_placeholders_stale_edits_and_invalid_target(draft):
    ctx,_,result,label,_=draft
    with pytest.raises(InvalidArguments,match='Unfinished'):
        call(ctx,'finalize_template_presentation',draft_id=result['draft_id'],revision=0,title='Test')
    with pytest.raises(InvalidArguments,match='target'):
        call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=0,slide_number=1,
             text_edits=[{'shape_id':999,'text':'Bad'}])
    call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=0,slide_number=1,
         text_edits=[{'shape_id':label,'text':'Done'}])
    with pytest.raises(InvalidArguments,match='Stale'):
        call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=0,slide_number=1,
             text_edits=[{'shape_id':label,'text':'Overwrite'}])
    with pytest.raises(InvalidArguments,match='Invalid presentation'):
        call(ctx,'inspect_presentation_draft',draft_id='../escape',revision=0,slide_number=1)


def test_native_chart_computes_counts_and_rejects_pii(draft):
    ctx,_,result,label,panel=draft
    data=ctx.project.root/'input.csv'
    pd.DataFrame({'gender':['Female','Male']*20,'category':['=1+1','ordinary']*20,
                  'email':[f'person{i}@example.test' for i in range(40)]}).to_csv(data,index=False)
    ds=call(ctx,'import_dataset',path=str(data))['dataset_id']
    with pytest.raises(InvalidArguments,match='non-PII'):
        call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=0,slide_number=1,
             chart={'shape_id':panel,'dataset_id':ds,'column':'email','kind':'counts'})
    edited=call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=0,slide_number=1,
                text_edits=[{'shape_id':label,'text':'Equal representation'}],
                chart={'shape_id':panel,'dataset_id':ds,'column':'gender','kind':'counts'})
    assert edited['chart_evidence'][0]['counts']==[20,20]
    deck=Presentation(edited['draft_path'])
    chart=next(s.chart for s in deck.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values)==[20,20]
    assert all(int(e.get('val')) >= 0 for e in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'))
    with ZipFile(edited['draft_path']) as archive:
        assert any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in archive.namelist())
        assert not any(b'@example.test' in archive.read(n) for n in archive.namelist() if n.endswith('.xml'))
    replacement=call(ctx,'edit_presentation_slide',draft_id=result['draft_id'],revision=1,slide_number=1,
                     chart={'shape_id':panel,'dataset_id':ds,'column':'category','kind':'counts'})
    deck=Presentation(replacement['draft_path'])
    assert sum(s.has_chart for s in deck.slides[0].shapes)==1
    chart=next(s.chart for s in deck.slides[0].shapes if s.has_chart)
    from openpyxl import load_workbook
    from io import BytesIO
    workbook=load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob))
    labels=[cell for row in workbook.active for cell in row if cell.value=='=1+1']
    assert labels and all(cell.data_type=='s' for cell in labels)


def test_completion_progress_and_final_export_evidence(draft):
    ctx, source, initial, label, panel = draft
    progress = initial['completion']
    assert progress['untouched_slides'] == [1]
    assert progress['placeholder_shapes'] == [{'slide': 1, 'shape_id': label}]
    assert not progress['ready_to_finalize']
    with pytest.raises(InvalidArguments, match='Unfinished template'):
        call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=0, title='Synthetic report')
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
                  slide_number=1, text_edits=[{'shape_id': label, 'text': 'Synthetic report'}])
    assert edited['completion']['ready_to_finalize']
    final = call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=1, title='Synthetic report')
    assert Path(final['path']).is_file() and final['completion']['ready_to_finalize']
    assert final['charts_embedded'] == 0 and final['charts_referenced_only'] == []
    assert final['visual_review_required'] is True
    assert final['review_slides'][0]['slide_number'] == 1
    assert any(e.get('text') == 'Synthetic report' for e in final['review_slides'][0]['text_elements'])


def test_new_draft_requires_inspection_before_any_edit(draft):
    ctx, source, initial, label, panel = draft
    template_id = ctx.project.root.joinpath('templates.json')
    import json
    template_id = json.loads(template_id.read_text())['templates'][0]['id']
    fresh = call(ctx, 'begin_template_presentation', template_id=template_id)
    before = Path(fresh['draft_path']).read_bytes()
    with pytest.raises(InvalidArguments, match='Inspect slide 1 first'):
        call(ctx, 'edit_presentation_slide', draft_id=fresh['draft_id'], revision=0,
             slide_number=1, text_edits=[{'shape_id': label, 'text': 'New finding'}])
    assert Path(fresh['draft_path']).read_bytes() == before
    call(ctx, 'inspect_presentation_draft', draft_id=fresh['draft_id'], revision=0, slide_number=1)
    edited = call(ctx, 'edit_presentation_slide', draft_id=fresh['draft_id'], revision=0,
                  slide_number=1, text_edits=[{'shape_id': label, 'text': 'New finding'}])
    assert edited['completion']['ready_to_finalize']


def test_inspection_can_recover_current_revision_without_guessing(draft):
    ctx, source, initial, label, panel = draft
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
                  slide_number=1, text_edits=[{'shape_id': label, 'text':'A verified update'}])
    restored = call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    assert restored['revision'] == edited['revision'] == 1
    assert restored['completion']['ready_to_finalize']
    with pytest.raises(InvalidArguments, match='Stale revision'):
        call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], revision=0, slide_number=1)


def test_gross_text_overflow_is_rejected_without_a_new_revision(draft):
    ctx, source, initial, label, panel = draft
    before = Path(initial['draft_path']).read_bytes()
    with pytest.raises(InvalidArguments, match='does not fit'):
        call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
             slide_number=1, text_edits=[{'shape_id':label, 'text':'Long paragraph ' * 150}])
    assert Path(initial['draft_path']).read_bytes() == before
    inspected = call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    assert inspected['revision'] == 0
    assert any(shape['region'] == 'body' for shape in inspected['shapes'])


def test_decorative_shape_cannot_receive_report_text(draft):
    ctx, source, initial, label, panel = draft
    with pytest.raises(InvalidArguments, match='decorative'):
        call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
             slide_number=1, text_edits=[{'shape_id':panel, 'text':'Important finding'}])
    assert call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)['revision'] == 0


def test_repeated_identical_edit_does_not_advance_revision(draft):
    ctx, source, initial, label, panel = draft
    call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
         slide_number=1, text_edits=[{'shape_id':label, 'text':'Verified finding'}])
    with pytest.raises(InvalidArguments, match='No content changed'):
        call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=1,
             slide_number=1, text_edits=[{'shape_id':label, 'text':'Verified finding'}])
    assert call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)['revision'] == 1


def test_legacy_draft_recovers_fields_from_original_template(draft):
    import json
    ctx, source, initial, label, panel = draft
    record = ctx.project.root/'presentation-drafts'/initial['draft_id']/'draft.json'
    manifest = json.loads(record.read_text())
    manifest.pop('text_fields')
    record.write_text(json.dumps(manifest))
    restored = call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    fields = {s['shape_id']: s for s in restored['shapes']}
    assert fields[label]['text_editable'] is True
    assert fields[panel]['text_editable'] is False
    assert 'text_fields' in json.loads(record.read_text())


def test_empty_standard_placeholders_are_editable_text_fields(tmp_path):
    ctx = JobContext(Project(tmp_path/'project', AllowList()), lambda *a: None, lambda: False)
    source = tmp_path/'simple.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    title, body = slide.shapes.title.shape_id, slide.placeholders[1].shape_id
    assert not slide.shapes.title.text and not slide.placeholders[1].text
    deck.save(source)
    ctx.project.allow.allow(source)
    template = library_action(ctx.project, {'action':'import', 'path':str(source)})['templates'][-1]
    result = call(ctx, 'begin_template_presentation', template_id=template['id'])
    state = call(ctx, 'inspect_presentation_draft', draft_id=result['draft_id'], slide_number=1)
    shapes = {s['shape_id']:s for s in state['shapes']}
    assert shapes[title]['text_editable'] and shapes[body]['text_editable']
    edited = call(ctx, 'edit_presentation_slide', draft_id=result['draft_id'], revision=0, slide_number=1,
                  text_edits=[{'shape_id':title,'text':'Household profile'}, {'shape_id':body,'text':'Synthetic sample: 500 households'}])
    final = call(ctx, 'finalize_template_presentation', draft_id=result['draft_id'], revision=edited['revision'], title='Profile')
    exported = Presentation(final['path'])
    assert exported.slides[0].shapes.title.text == 'Household profile'
    assert exported.slides[0].placeholders[1].text == 'Synthetic sample: 500 households'


@pytest.mark.parametrize('text', ['Research finding [1]', 'Estimate [95% CI]', 'Name [REDACTED]'])
def test_new_bracketed_research_content_is_not_a_source_prompt(draft, text):
    ctx, source, initial, label, panel = draft
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
                  slide_number=1, text_edits=[{'shape_id': label, 'text': text}])
    assert edited['completion']['placeholder_shape_count'] == 0
    state = call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    assert not next(s for s in state['shapes'] if s['shape_id'] == label)['contains_placeholder']
    final = call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=1, title='Research')
    assert Presentation(final['path']).slides[0].shapes[0].text == text


def test_original_source_prompt_still_blocks_after_slide_is_edited(draft):
    ctx, source, initial, label, panel = draft
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
                  slide_number=1, text_edits=[{'shape_id': label, 'text': 'Updated [finding]'}])
    assert edited['completion']['untouched_slides'] == []
    assert edited['completion']['placeholder_shapes'] == [{'slide': 1, 'shape_id': label}]
    with pytest.raises(InvalidArguments, match='Unfinished template'):
        call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=1, title='Research')


def test_table_source_prompts_block_but_research_brackets_are_allowed(tmp_path):
    ctx = JobContext(Project(tmp_path/'project', AllowList()), lambda *a: None, lambda: False)
    source = tmp_path/'table-template.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(.5), Inches(1), Inches(8), Inches(2))
    table.table.cell(0, 0).text = '[Finding]'
    table.table.cell(0, 1).text = 'References [1]'
    deck.save(source)
    ctx.project.allow.allow(source)
    template = library_action(ctx.project, {'action':'import', 'path':str(source)})['templates'][-1]
    initial = call(ctx, 'begin_template_presentation', template_id=template['id'])
    call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0, slide_number=1,
                  table_edits=[{'shape_id': table.shape_id, 'cells': [['[Finding]', 'References [1]']]}])
    assert edited['completion']['placeholder_shape_count'] == 1
    with pytest.raises(InvalidArguments, match='Unfinished template'):
        call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=1, title='Research')
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=1, slide_number=1,
                  table_edits=[{'shape_id': table.shape_id, 'cells': [['Estimate [95% CI]', 'References [1]']]}])
    assert edited['completion']['ready_to_finalize']
    final = call(ctx, 'finalize_template_presentation', draft_id=initial['draft_id'], revision=2, title='Research')
    assert Presentation(final['path']).slides[0].shapes[0].table.cell(0, 0).text == 'Estimate [95% CI]'


def test_legacy_draft_restores_original_prompt_tokens_without_rejecting_new_brackets(draft):
    import json
    ctx, source, initial, label, panel = draft
    call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=0,
         slide_number=1, text_edits=[{'shape_id': label, 'text': 'New result [1]'}])
    record = ctx.project.root/'presentation-drafts'/initial['draft_id']/'draft.json'
    manifest = json.loads(record.read_text())
    manifest.pop('source_placeholders')
    record.write_text(json.dumps(manifest))
    restored = call(ctx, 'inspect_presentation_draft', draft_id=initial['draft_id'], slide_number=1)
    assert restored['revision'] == 1 and restored['completion']['ready_to_finalize']
    assert json.loads(record.read_text())['source_placeholders'] == ['[finding]']
    edited = call(ctx, 'edit_presentation_slide', draft_id=initial['draft_id'], revision=1,
                  slide_number=1, text_edits=[{'shape_id': label, 'text': '[Finding]'}])
    assert edited['completion']['placeholder_shape_count'] == 1
