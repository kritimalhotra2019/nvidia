"""export_replication_package: a journaled import → clean → aggregate → treatment effect → Excel run becomes a
folder with README, steps.json, reproduce.py, reproduce.R, data, outputs, manifest and zip; reproduce.py replays
the run under the research-env python and regenerates the Excel output."""
import hashlib
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, registry
from squipi_worker.research import ResearchStore
from squipi_worker.objectives import ObjectiveStore
from squipi_worker.security import AllowList

RESEARCH_PYTHON = Path(os.environ.get("SQUIPI_RESEARCH_PYTHON") or (Path.home() / "Library/Application Support/Squipi/research-env/bin/python"))


def _ctx(tmp_path):
    allow = AllowList()
    project = Project(tmp_path / "project", allow)
    return project, JobContext(project, report=lambda p, n: None, cancelled=lambda: False)


def _trial(path: Path, n=60):
    rng = np.random.default_rng(7)
    treated = rng.integers(0, 2, n)
    frame = pd.DataFrame({"household_id": [f"HH{i:03d}" for i in range(n)], "respondent_name": [f"Person {i}" for i in range(n)],
                          "district": np.where(np.arange(n) % 3 == 0, "North", "South"), "treated": treated,
                          "baseline_income": rng.normal(1000, 100, n).round(1),
                          "income": (rng.normal(1000, 100, n) + 80 * treated).round(1)})
    frame.to_csv(path, index=False)
    return path


def _journaled_run(project, ctx, src):
    """Run the tools through the registry and record each call in the journal the way the host does."""
    store = ResearchStore(project.root)
    task = store.start_research_task("Clean the trial data, drop respondent names, summarise income by district, estimate the effect of the programme on income and save the clean data as an Excel workbook.", [])["task_id"]
    calls = [("import_dataset", lambda ids: {"path": str(src)}),
             ("clean_dataset", lambda ids: {"dataset_id": ids["raw"], "drop_columns": ["respondent_name"], "pii": "keep", "steps": []}),
             ("aggregate", lambda ids: {"dataset_id": ids["clean"], "group_by": ["district"], "metrics": [{"column": "income", "agg": "mean"}, {"column": "income", "agg": "count"}]}),
             ("estimate_treatment_effect", lambda ids: {"dataset_id": ids["clean"], "outcome": "income", "treatment": "treated", "covariates": ["baseline_income"], "design": "randomized"}),
             ("export_excel", lambda ids: {"dataset_id": ids["clean"], "filename": "Clean trial.xlsx"})]
    ids, results = {}, {}
    for name, make in calls:
        args = make(ids)
        result = registry.call(name, args, ctx)
        store.record_research_step(task, name, args, result, True)
        results[name] = result
        if name == "import_dataset":
            ids["raw"] = result["dataset_id"]
        if name == "clean_dataset":
            ids["clean"] = result["dataset_id"]
    # a failed call is journaled too and must be kept for the record but not replayed
    store.record_research_step(task, "aggregate", {"dataset_id": ids["clean"], "metrics": [{"column": "nope", "agg": "mean"}]}, {"error": "column nope not found"}, False)
    store.finish_research_task(task, "completed", "Done.")
    return task, results


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def test_replication_package_is_complete_and_reproduce_py_regenerates_the_excel(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _trial(tmp_path / "trial.csv")
    project.allow.allow(src)
    task, results = _journaled_run(project, ctx, src)

    out = registry.call("export_replication_package", {"task_id": task, "name": "trial-package"}, ctx)
    pkg = Path(out["package_dir"])
    for name in ("README.md", "steps.json", "reproduce.py", "reproduce.R", "manifest.json"):
        assert (pkg / name).is_file(), name
    assert Path(out["zip_path"]).is_file()
    with zipfile.ZipFile(out["zip_path"]) as z:
        names = z.namelist()
    assert any(n.endswith("README.md") for n in names) and any(n.endswith("reproduce.py") for n in names)
    assert (pkg / "code" / "squipi_worker" / "registry.py").is_file()

    # manifest hashes match the files on disk, and every file is listed
    manifest = json.loads((pkg / "manifest.json").read_text("utf-8"))
    listed = set(manifest["files"])
    on_disk = {str(p.relative_to(pkg)) for p in pkg.rglob("*") if p.is_file() and p.name != "manifest.json"}
    assert listed == on_disk
    for rel, meta in manifest["files"].items():
        assert _sha(pkg / rel) == meta["sha256"], rel
    assert out["file_count"] == len(listed)

    # derived data copied, original source referenced by path + sha256 only
    clean_id = results["clean_dataset"]["dataset_id"]
    raw_id = results["import_dataset"]["dataset_id"]
    assert manifest["datasets"][clean_id]["package_path"].startswith("data/")
    assert manifest["datasets"][raw_id]["package_path"] is None
    assert manifest["datasets"][raw_id]["sha256"] == _sha(src)
    assert manifest["datasets"][raw_id]["rows"] == 60
    assert manifest["original_sources"][0]["path"] == str(src)
    assert out["sources_referenced_only"] == 1 and out["datasets_copied"] == 1
    assert not list((pkg / "sources").iterdir())
    # outputs copied: the Excel workbook, the aggregate table/chart, the treatment-effect chart/balance table
    outputs = list((pkg / "outputs").iterdir())
    assert any(p.suffix == ".xlsx" for p in outputs) and any(p.suffix == ".svg" for p in outputs) and any(p.suffix == ".csv" for p in outputs)
    assert out["outputs_missing"] == []

    # steps.json is the journal slice with copied paths rewritten
    steps = json.loads((pkg / "steps.json").read_text("utf-8"))["steps"]
    assert [s["tool"] for s in steps] == ["import_dataset", "clean_dataset", "aggregate", "estimate_treatment_effect", "export_excel", "aggregate"]
    assert steps[-1]["ok"] is False and steps[-1]["result"]["error"]
    assert steps[1]["result"]["path"].startswith("data/")
    assert steps[4]["result"]["path"].startswith("outputs/")
    assert steps[1]["arguments"]["drop_columns"] == ["respondent_name"]   # column names stay as they are
    assert out["step_count"] == 6 and out["reproducible_step_count"] == 5

    # README: title, instruction, inputs with hashes and rows, steps, outputs, software, limitations, original sources
    readme = (pkg / "README.md").read_text("utf-8")
    for heading in ("## What was asked", "## Inputs", "## Steps", "## Outputs", "## Software", "## Limitations recorded in results", "## Original sources"):
        assert heading in readme
    assert "estimate the effect of the programme" in readme
    assert manifest["datasets"][raw_id]["sha256"] in readme and "| 60 |" in readme
    assert "pandas " in readme and "`estimate_treatment_effect`" in readme
    assert '"design":"randomized"' in readme
    home = str(Path.home())
    body, _, originals = readme.partition("## Original sources")
    assert home not in body and str(src) not in body, "no absolute home/source path outside the original-sources section"
    assert str(src) in originals and f"<original source {raw_id}: trial.csv>" in body
    assert home not in (pkg / "reproduce.py").read_text("utf-8")
    assert home not in (pkg / "reproduce.R").read_text("utf-8")

    # reproduce.R covers cleaning, aggregate, treatment effect and the export; the import reads the packaged data or the source
    r = (pkg / "reproduce.R").read_text("utf-8")
    assert "setdiff(names(d2), c(\"respondent_name\"))" in r
    assert "aggregate(" in r and "lm(`income` ~ treated4 + `baseline_income`" in r
    assert "write.csv(d2" in r
    assert sorted(out["r_coverage"]["covered_steps"]) == [1, 2, 3, 4, 5]

    # reproduce.py replays the run from the package folder under the research-env python and recreates the Excel output
    python = str(RESEARCH_PYTHON) if RESEARCH_PYTHON.is_file() else sys.executable
    proc = subprocess.run([python, "reproduce.py"], cwd=str(pkg), capture_output=True, text=True, timeout=600,
                          env={**os.environ, "SQUIPI_EMBEDDINGS": "off"})
    assert proc.returncode == 0, proc.stdout + proc.stderr
    log = json.loads((pkg / "run" / "results.json").read_text("utf-8"))
    assert log["failures"] == 0 and [s["tool"] for s in log["steps"] if "sequence" in s] == ["import_dataset", "clean_dataset", "aggregate", "estimate_treatment_effect", "export_excel"]
    regenerated = [Path(p) for p in log["steps"][-1]["outputs"] if p.endswith(".xlsx")]
    assert regenerated and regenerated[0].name == "Clean trial.xlsx" and str(regenerated[0]).startswith(str(pkg / "run"))
    original = openpyxl.load_workbook(results["export_excel"]["path"], read_only=True)["Data"]
    fresh = openpyxl.load_workbook(regenerated[0], read_only=True)["Data"]
    assert (fresh.max_row, fresh.max_column) == (original.max_row, original.max_column) == (61, 5)
    assert "respondent_name" not in [c.value for c in next(fresh.iter_rows(min_row=1, max_row=1))]


def test_replication_package_from_objective_receipts_and_missing_outputs_are_noted(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _trial(tmp_path / "trial.csv")
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    store = ObjectiveStore(project)
    store.open("objective-rep-1", "Summarise income by district and save it as 'Income.xlsx'.", [ds])
    receipt = store.begin("objective-rep-1", "aggregate", {"dataset_id": ds, "group_by": ["district"], "metrics": [{"column": "income", "agg": "mean"}]})
    agg = registry.call("aggregate", {"dataset_id": ds, "group_by": ["district"], "metrics": [{"column": "income", "agg": "mean"}]}, ctx)
    store.finish_step(receipt, agg, True)
    receipt = store.begin("objective-rep-1", "export_excel", {"dataset_id": ds, "filename": "Income.xlsx"})
    exported = registry.call("export_excel", {"dataset_id": ds, "filename": "Income.xlsx"}, ctx)
    store.finish_step(receipt, exported, True)
    Path(exported["path"]).unlink()   # the user deleted the workbook afterwards
    out = registry.call("export_replication_package", {"objective_id": "objective-rep-1", "copy_sources": True}, ctx)
    pkg = Path(out["package_dir"])
    assert out["source"] == {"kind": "objective", "id": "objective-rep-1"} and out["step_count"] == 2
    assert out["outputs_missing"] == [exported["path"]] and any("Income.xlsx" in n for n in out["notes"])
    manifest = json.loads((pkg / "manifest.json").read_text("utf-8"))
    assert manifest["datasets"][ds]["package_path"].startswith("sources/") and (pkg / manifest["datasets"][ds]["package_path"]).is_file()
    assert "Income.xlsx" in (pkg / "README.md").read_text("utf-8")
    with pytest.raises(Exception):
        registry.call("export_replication_package", {}, ctx)


def test_r_script_marks_steps_without_an_r_equivalent(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _trial(tmp_path / "trial.csv")
    project.allow.allow(src)
    store = ResearchStore(project.root)
    task = store.start_research_task("Profile the data", [])["task_id"]
    imported = registry.call("import_dataset", {"path": str(src)}, ctx)
    store.record_research_step(task, "import_dataset", {"path": str(src)}, imported, True)
    profiled = registry.call("profile_dataset", {"dataset_id": imported["dataset_id"]}, ctx)
    store.record_research_step(task, "profile_dataset", {"dataset_id": imported["dataset_id"]}, profiled, True)
    out = registry.call("export_replication_package", {"task_id": task, "include_data": False}, ctx)
    r = Path(out["reproduce_r_path"]).read_text("utf-8")
    assert "run reproduce.py (tool: profile_dataset)" in r
    assert out["r_coverage"]["not_in_r"] == [{"sequence": 2, "tool": "profile_dataset"}]
    assert any("not reproduced in R" in n for n in out["notes"])
    assert "This source file is not included in the package" in r
    json.dumps(out)
