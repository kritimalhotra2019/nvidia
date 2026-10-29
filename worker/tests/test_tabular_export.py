import pandas as pd
from openpyxl import load_workbook
import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, registry, Cancelled
from squipi_worker.security import AllowList


def context(tmp_path):
    source = tmp_path / "pending.csv"
    pd.DataFrame({"household_id": ["HH1", "HH2"], "name": ["Synthetic A", "Synthetic B"],
                  "notes": ["=1+1", "+literal"], "income": [100, None]}).to_csv(source, index=False)
    allow = AllowList()
    allow.allow(source)
    project = Project(tmp_path / "project", allow)
    ctx = JobContext(project, lambda *args: None, lambda: False)
    imported = registry.call("import_dataset", {"path": str(source)}, ctx)
    return ctx, imported["dataset_id"], source


def test_shared_excel_exports_cleaned_data_and_literal_text(tmp_path):
    ctx, dataset, source = context(tmp_path)
    original = source.read_bytes()
    cleaned = registry.call("clean_dataset", {"dataset_id": dataset, "pii": "keep",
                            "drop_columns": ["name"], "steps": []}, ctx)
    result = registry.call("export_excel", {"dataset_id": cleaned["dataset_id"]}, ctx)
    book = load_workbook(result["path"])
    sheet = book["Data"]
    assert list(sheet.values) == [("household_id", "notes", "income"), ("HH1", "=1+1", 100), ("HH2", "+literal", None)]
    assert sheet["B2"].data_type == "s"
    assert sheet.column_dimensions["A"].width >= len("household_id") + 2
    assert sheet["A1"].alignment.wrap_text is True
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:C3"
    assert source.read_bytes() == original
    assert result["rows"] == 2 and result["format"] == "xlsx"


def test_excel_layout_handles_long_multiline_and_unicode_without_changing_data(tmp_path):
    long_header = "Household wellbeing index measured using the complete survey instrument and reported on its original scale"
    wide_header = "家庭收入与生活状况调查指标" * 4
    frame = pd.DataFrame({
        long_header: [1.25, None],
        "=HEADER_IS_LITERAL": ["=1+1", '=HYPERLINK("https://example.invalid","literal")'],
        "Interview notes": ["A long synthetic interview response. " * 18, "First line\nSecond line\nThird line"],
        wide_header: ["印度家庭", "中文访谈"],
        "Survey date": pd.to_datetime(["2026-10-01", "2026-10-02"]),
    })
    source = tmp_path / "typed-households.parquet"
    frame.to_parquet(source, index=False)
    original = source.read_bytes()
    allow = AllowList()
    allow.allow(source)
    ctx = JobContext(Project(tmp_path / "project", allow), lambda *args: None, lambda: False)
    dataset = registry.call("import_dataset", {"path": str(source)}, ctx)["dataset_id"]
    result = registry.call("export_excel", {"dataset_id": dataset, "filename": "Readable households.xlsx"}, ctx)
    book = load_workbook(result["path"], data_only=False)
    sheet = book["Data"]
    rows = list(sheet.values)
    assert rows[0] == tuple(frame.columns)
    assert rows[1] == (1.25, "=1+1", frame.iloc[0, 2], "印度家庭", frame.iloc[0, 4].to_pydatetime())
    assert rows[2] == (None, frame.iloc[1, 1], frame.iloc[1, 2], "中文访谈", frame.iloc[1, 4].to_pydatetime())
    assert sheet["B1"].data_type == sheet["B2"].data_type == sheet["B3"].data_type == "s"
    assert sheet["A2"].data_type == "n" and sheet["E2"].data_type == "d"
    assert all(12 <= sheet.column_dimensions[column].width <= 48 for column in ["A", "B", "C", "D", "E"])
    assert sheet.column_dimensions["A"].width == sheet.column_dimensions["D"].width == 48
    assert all(cell.alignment.wrap_text is True for cell in sheet[1])
    assert sheet["C2"].alignment.wrap_text is True and sheet["C3"].alignment.wrap_text is True
    assert 30 < sheet.row_dimensions[1].height <= 90
    assert 30 < sheet.row_dimensions[2].height <= 90
    assert 30 < sheet.row_dimensions[3].height <= 90
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:E3"
    assert source.read_bytes() == original
    assert result["rows"] == 2 and result["columns"] == list(frame.columns)


def test_export_cancellation_leaves_no_workbook(tmp_path):
    ctx, dataset, source = context(tmp_path)
    ctx._cancelled = lambda: True
    with pytest.raises(Cancelled):
        registry.call("export_excel", {"dataset_id": dataset}, ctx)
    assert not list(ctx.project.root.glob("artifacts/**/*.xlsx"))
