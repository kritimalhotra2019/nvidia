"""Day-1/2 harness changes: interrupted-task sweep, small-group guards, PII-absent requirement,
quoted filenames with dots, learned-hint relevance, layout overlap detection, docx chart verification."""
import json
import sqlite3
from pathlib import Path

import pandas as pd
import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.project import Project
from squipi_worker.registry import JobContext, registry
from squipi_worker.security import AllowList
from squipi_worker.research import ResearchStore
from squipi_worker.objectives import ObjectiveStore
from squipi_worker.learning import LearningStore
from squipi_worker.tools import presentation_edit


def _ctx(tmp_path):
    allow = AllowList()
    project = Project(tmp_path / "project", allow)
    return project, JobContext(project, report=lambda p, n: None, cancelled=lambda: False)


def _households(path: Path, n=40, tiny=2):
    rows = []
    for i in range(n):
        rows.append({"household_id": f"HH{i:03d}", "respondent_name": f"Person {i}", "phone": f"+91 98{i:08d}",
                     "email": f"p{i}@example.org", "district": "North" if i < n // 2 else "South",
                     "programme_group": "pilot" if i < tiny else ("treatment" if i % 2 else "comparison"),
                     "monthly_income_inr": 12000 + 517 * i, "wellbeing_index": 40 + (i % 7) * 5.5})
    pd.DataFrame(rows).to_csv(path, index=False)
    return path


def test_running_tasks_are_swept_once_per_process(tmp_path):
    store = ResearchStore(tmp_path)
    task = store.start_research_task("Clean this", [])["task_id"]
    assert store.read_research_task(task)["task"]["status"] == "running"
    assert ResearchStore(tmp_path).sweep_interrupted() == 1
    record = store.read_research_task(task)["task"]
    assert record["status"] == "partial" and "Interrupted" in record["summary"]
    assert ResearchStore(tmp_path).sweep_interrupted() == 0


def test_tiny_groups_are_not_compared(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _households(tmp_path / "h.csv")
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    out = registry.call("compare_groups", {"dataset_id": ds, "outcome": "wellbeing_index", "group": "programme_group"}, ctx)
    assert any("pilot" in s for s in out["excluded_small_groups"])
    assert all(r["group"] != "pilot" for r in out["groups"])
    assert "fewer than" in out["interpretation"][0]
    agg = registry.call("aggregate", {"dataset_id": ds, "group_by": ["programme_group"], "metrics": [{"column": "wellbeing_index", "agg": "mean"}]}, ctx)
    assert any(g["group"] == "pilot" and g["n"] == 2 for g in agg["small_groups"])
    assert "Caution" in agg["summary_text"] and "n = " in agg["summary_text"]
    insights = registry.call("find_insights", {"dataset_id": ds}, ctx)
    for finding in insights["findings"]:
        if finding.get("kind") == "group_difference":
            assert all(g["n"] >= 5 for g in finding["numbers"]["groups"])
    assert any("fewer than" in limit for limit in insights["limits"])


def test_pii_absent_requirement_fails_on_raw_export_and_passes_after_cleaning(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _households(tmp_path / "h.csv")
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    store = ObjectiveStore(project)
    instruction = "Clean this dataset, remove personal information, and save it as 'Clean households.xlsx'."
    snap = store.open("objective-pii-1", instruction, [ds])
    # The quoted filename with a dot now registers as a host output requirement.
    assert any(r["id"] == "host-output-xlsx" and r.get("filename") == "Clean households.xlsx" for r in snap["requirements"])
    store.requirements("objective-pii-1", [{"id": "contract-pii", "kind": "pii_absent", "quote": "remove personal information",
                                            "description": "No personal-data columns in the delivered data", "source_ids": [ds]}])
    ctx.objective_id = "objective-pii-1"
    receipt = store.begin("objective-pii-1", "export_excel", {"dataset_id": ds, "filename": "Clean households.xlsx"})
    raw = registry.call("export_excel", {"dataset_id": ds, "filename": "Clean households.xlsx"}, ctx)
    store.finish_step(receipt, raw, True)
    review = store.review("objective-pii-1")
    assert not review["passed"]
    assert any(c["id"] == "contract-pii" and not c["passed"] for c in review["checks"])
    # After a real cleaning with PII dropped the same requirement passes.
    receipt = store.begin("objective-pii-1", "clean_dataset", {"dataset_id": ds, "pii": "drop"})
    cleaned = registry.call("clean_dataset", {"dataset_id": ds, "pii": "drop"}, ctx)
    store.finish_step(receipt, cleaned, True)
    receipt = store.begin("objective-pii-1", "export_excel", {"dataset_id": cleaned["dataset_id"], "filename": "Clean households.xlsx"})
    exported = registry.call("export_excel", {"dataset_id": cleaned["dataset_id"], "filename": "Clean households.xlsx"}, ctx)
    store.finish_step(receipt, exported, True)
    review = store.review("objective-pii-1")
    assert any(c["id"] == "contract-pii" and c["passed"] for c in review["checks"])
    assert any(c["id"] == "host-output-xlsx" and c["passed"] for c in review["checks"])


def test_output_requirement_regex_accepts_give_me_and_dotted_names(tmp_path):
    project, ctx = _ctx(tmp_path)
    store = ObjectiveStore(project)
    snap = store.open("objective-word-1", "Using the cleaned data, give me descriptive statistics by district, with labelled graphs, as a Word report named 'Income brief.docx'.", [])
    assert any(r["id"] == "host-output-docx" and r.get("filename") == "Income brief.docx" for r in snap["requirements"])


def test_learned_hints_require_relevance(tmp_path):
    store = LearningStore(tmp_path)
    with store.connect() as c:
        c.execute("INSERT INTO lessons VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  ("learn_x", "fp", "procedure", "Clean and export dataset", "do it", "cleaning a dataset and exporting excel",
                   json.dumps(["clean_dataset"]), json.dumps(["check"]), json.dumps(["s1"]), "task_a", "trial", 1, 0, "2026-01-01"))
    assert store.context("clean this dataset and export excel")["lessons"]
    assert not store.context("summarise the fieldbook recommendation")["lessons"]


def test_overlap_detector_flags_chart_over_text():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)); box.text_frame.text = "Narrative paragraph that must stay readable"
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    data = CategoryChartData(); data.categories = ["a", "b"]; data.add_series("s", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(2.5), Inches(4), Inches(3), data)
    issues = presentation_edit._overlap_issues(deck)
    assert issues and issues[0]["slide"] == 1 and issues[0]["overlap_share"] > 0.3
    clean = Presentation(); s2 = clean.slides.add_slide(clean.slide_layouts[5])
    t = s2.shapes.add_textbox(Inches(0.5), Inches(1), Inches(3), Inches(1)); t.text_frame.text = "left"
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(3), data)
    assert presentation_edit._overlap_issues(clean) == []


def test_docx_export_reports_verified_inline_charts(tmp_path):
    project, ctx = _ctx(tmp_path)
    src = _households(tmp_path / "h.csv")
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    chart = registry.call("create_chart", {"dataset_id": ds, "kind": "histogram", "x": "monthly_income_inr"}, ctx)
    out = registry.call("export_document", {"format": "docx", "title": "Brief", "sections": [{"heading": "Income", "text": "Distribution.", "chart_path": chart["chart_path"]}]}, ctx)
    assert out["charts_verified_in_file"] == out["charts_embedded"] >= 1


def test_analysis_requirement_is_satisfied_by_a_deterministic_impact_tool(tmp_path):
    """An 'analysis' contract row is met by estimate_treatment_effect on the required columns (26 Oct x1 case)."""
    import numpy as np, pandas as pd
    project, ctx = _ctx(tmp_path)
    rng = np.random.default_rng(3)
    src = tmp_path / "trial.csv"
    pd.DataFrame({"treated": rng.integers(0, 2, 200), "attendance_rate": rng.normal(80, 5, 200).round(1), "year": 2025}).to_csv(src, index=False)
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    store = ObjectiveStore(project)
    snap = store.open("objective-did-1", "Estimate the effect of the programme on attendance_rate", [ds])
    store.requirements("objective-did-1", [{"id": "contract-analysis", "kind": "analysis", "quote": "effect of the programme on attendance",
                                            "description": "Estimate the effect on attendance", "columns": ["attendance_rate", "treated"], "source_ids": [ds]}])
    ctx.objective_id = "objective-did-1"
    assert not store.review("objective-did-1")["passed"]
    receipt = store.begin("objective-did-1", "estimate_treatment_effect", {"dataset_id": ds, "outcome": "attendance_rate", "treatment": "treated"})
    result = registry.call("estimate_treatment_effect", {"dataset_id": ds, "outcome": "attendance_rate", "treatment": "treated"}, ctx)
    store.finish_step(receipt, result, True)
    review = store.review("objective-did-1")
    assert any(c["id"] == "contract-analysis" and c["passed"] for c in review["checks"]), review


def test_pii_absent_respects_columns_the_user_chose_to_keep(tmp_path):
    """'…but keep the village GPS' → the contract lists village_gps as an exception; the review passes with it present."""
    import pandas as pd
    project, ctx = _ctx(tmp_path)
    src = tmp_path / "ct.csv"
    pd.DataFrame({"beneficiary_name": ["A B", "C D", "E F", "G H"], "mobile": ["+1 555 010 1234", "+1 555 010 2345", "+1 555 010 3456", "+1 555 010 4567"],
                  "village_gps": [1.34, 0.43, 1.2, 0.9], "consumption": [100, 120, 90, 110]}).to_csv(src, index=False)
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    store = ObjectiveStore(project)
    store.open("objective-keep-1", "Strip out anything that identifies beneficiaries but keep the village GPS column, and save it as 'Kept.xlsx'.", [ds])
    store.requirements("objective-keep-1", [{"id": "contract-pii", "kind": "pii_absent", "quote": "strip out anything that identifies beneficiaries",
                                             "description": "No personal-data columns in the delivered data", "columns": ["village_gps"], "source_ids": [ds]}])
    ctx.objective_id = "objective-keep-1"
    receipt = store.begin("objective-keep-1", "clean_dataset", {"dataset_id": ds, "drop_columns": ["beneficiary_name", "mobile"]})
    cleaned = registry.call("clean_dataset", {"dataset_id": ds, "drop_columns": ["beneficiary_name", "mobile"]}, ctx)
    store.finish_step(receipt, cleaned, True)
    review = store.review("objective-keep-1")
    check = next(c for c in review["checks"] if c["id"] == "contract-pii")
    assert check["passed"], check
    assert "kept by request: village_gps" in check["description"]
    # Without the exception the same data fails, because coordinates are personal data by default.
    store.requirements("objective-keep-1", [{"id": "contract-pii", "kind": "pii_absent", "quote": "strip out anything that identifies beneficiaries",
                                             "description": "No personal-data columns in the delivered data", "source_ids": [ds]}])
    strict = next(c for c in store.review("objective-keep-1")["checks"] if c["id"] == "contract-pii")
    assert not strict["passed"] and "village_gps" in strict["description"]


def test_analysis_requirement_is_satisfied_by_find_insights_on_the_source(tmp_path):
    """'analyze this dataset and produce key insights': the contract bound schema columns the user never named;
    find_insights examines every column, so it meets the requirement (27 Oct, household dataset)."""
    import numpy as np, pandas as pd
    project, ctx = _ctx(tmp_path)
    rng = np.random.default_rng(5)
    src = tmp_path / "households.csv"
    pd.DataFrame({"district": rng.choice(["North", "South"], 200), "head_age": rng.integers(20, 70, 200),
                  "food_insecure_days": rng.integers(0, 12, 200), "monthly_income": rng.normal(9000, 2000, 200).round()}).to_csv(src, index=False)
    project.allow.allow(src)
    ds = registry.call("import_dataset", {"path": str(src)}, ctx)["dataset_id"]
    store = ObjectiveStore(project)
    store.open("objective-insights-1", "analyze this dataset and produce key insights", [ds])
    store.requirements("objective-insights-1", [{"id": "contract-analysis-1", "kind": "analysis", "quote": "analyze this dataset and produce key insights",
                                                 "description": "Perform statistical analysis", "columns": ["head_age", "food_insecure_days", "monthly_income"], "source_ids": [ds]}])
    ctx.objective_id = "objective-insights-1"
    assert not store.review("objective-insights-1")["passed"]
    receipt = store.begin("objective-insights-1", "find_insights", {"dataset_id": ds, "charts": False})
    result = registry.call("find_insights", {"dataset_id": ds, "charts": False}, ctx)
    store.finish_step(receipt, result, True)
    review = store.review("objective-insights-1")
    assert any(c["id"] == "contract-analysis-1" and c["passed"] for c in review["checks"]), review
