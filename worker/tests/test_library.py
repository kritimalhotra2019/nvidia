"""Project library: cross-document meaning-based search and table extraction into datasets and the index."""
import json

import pytest

from squipi_worker import tools  # noqa: F401
from squipi_worker.embeddings import Embedder
from squipi_worker.project import Project
from squipi_worker.registry import InvalidArguments, JobContext, registry
from squipi_worker.security import AllowList
from tests.test_embeddings import fake_embed  # deterministic stand-in for nomic


@pytest.fixture
def project(tmp_path):
    return Project(tmp_path / "project", AllowList())


@pytest.fixture
def fake_embedder(monkeypatch):
    monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": fake_embed(texts, kind))
    monkeypatch.setattr(Embedder, "available", lambda self: True)
    yield


def call(project, name, **args):
    return registry.call(name, args, JobContext(project, lambda *_: None, lambda: False))


def two_documents(project, tmp_path):
    roads = tmp_path / "roads_report.md"
    roads.write_text("# Roads\n" + ("Rural roads were rehabilitated; transport time fell. " * 60) +
                     "\n# Attrition\n" + ("Loss to follow-up reached 14 percent; retention was lowest in the poorest quintile. " * 25))
    cash = tmp_path / "cash_report.md"
    cash.write_text("# Cash\n" + ("Treatment households received transfers; consumption and spending rose at endline. " * 60) +
                    "\n# Dropout\n" + ("Dropout between rounds was 9 percent and balanced across arms. " * 25))
    project.allow.allow(roads); project.allow.allow(cash)
    a = call(project, "read_document", path=str(roads))["document_id"]
    b = call(project, "read_document", path=str(cash))["document_id"]
    return a, b


def test_search_library_ranks_passages_across_documents(project, tmp_path, fake_embedder):
    roads, cash = two_documents(project, tmp_path)
    out = call(project, "search_library", query="what was the attrition between survey rounds", max_matches=4)
    assert out["mode"] == "semantic"
    assert {d["document_id"] for d in out["documents_searched"]} == {roads, cash}
    assert out["matches"], "expected ranked passages"
    top = out["matches"][0]
    assert top["document_id"] in (roads, cash) and top["kind"] == "text"
    assert out["content_trust"] == "untrusted_document_data"
    # A second search is served from the index: nothing new is embedded.
    again = call(project, "search_library", query="attrition", max_matches=2)
    assert again["indexed_passages"] == 0
    # Restricting to one document keeps results inside it.
    only = call(project, "search_library", query="attrition", document_ids=[cash])
    assert all(m["document_id"] == cash for m in only["matches"])


def test_search_library_literal_fallback_and_errors(project, tmp_path, monkeypatch):
    with pytest.raises(InvalidArguments):
        call(project, "search_library", query="anything")
    roads, cash = two_documents(project, tmp_path)
    monkeypatch.setattr(Embedder, "embed", lambda self, texts, kind="document": None)
    out = call(project, "search_library", query="poorest quintile")
    assert out["mode"] == "literal" and "unavailable" in out["fallback"]
    assert out["matches"] and all(m["document_id"] == roads for m in out["matches"])


def test_extract_docx_tables_registers_datasets_and_indexes_them(project, tmp_path, fake_embedder):
    import docx
    d = docx.Document()
    d.add_heading("Results", 1)
    d.add_paragraph("Table 1 reports attrition by arm.")
    t = d.add_table(rows=4, cols=3)
    for j, h in enumerate(["arm", "attrition_rate", "n"]):
        t.cell(0, j).text = h
    for i, row in enumerate([["treatment", "0.18", "300"], ["control", "0.10", "300"], ["total", "0.14", "600"]], start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    path = tmp_path / "attrition_tables.docx"
    d.save(str(path))
    project.allow.allow(path)
    doc_id = call(project, "read_document", path=str(path))["document_id"]
    out = call(project, "extract_document_tables", document_id=doc_id)
    assert out["tables_kept"] == 1
    table = out["tables"][0]
    assert table["columns"] == ["arm", "attrition_rate", "n"] and table["rows"] == 3
    datasets = {ds.id for ds in project.datasets()}
    assert table["dataset_id"] in datasets
    assert out["indexed_passages"] == 1
    hit = call(project, "search_library", query="attrition rate by arm", max_matches=3)
    kinds = {m["kind"] for m in hit["matches"]}
    assert "table" in kinds
    tbl = next(m for m in hit["matches"] if m["kind"] == "table")
    assert "attrition_rate" in tbl["text"] and tbl["document_id"] == doc_id
    json.dumps(out, default=str)


def test_extract_tables_rejects_non_document_kinds(project, tmp_path):
    p = tmp_path / "notes.md"
    p.write_text("# Notes\nplain text")
    project.allow.allow(p)
    doc_id = call(project, "read_document", path=str(p))["document_id"]
    with pytest.raises(InvalidArguments):
        call(project, "extract_document_tables", document_id=doc_id)
