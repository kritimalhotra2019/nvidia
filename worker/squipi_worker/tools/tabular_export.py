"""Export an existing registered dataset without transforming its values."""
from pathlib import Path
from math import ceil
from unicodedata import east_asian_width
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..registry import tool, InvalidArguments
from .datasets import read_dataset

EXCEL_SCHEMA = {"type": "object", "properties": {"dataset_id": {"type": "string"}, "filename": {"type":"string","minLength":1,"maxLength":120,"description":"Exact requested output name; .xlsx is added if omitted. No directory path."}},
                "required": ["dataset_id"], "additionalProperties": False}


def _line_widths(value):
    """Approximate Excel character widths, including wide scripts and explicit newlines."""
    return [sum(2 if east_asian_width(char) in {"W", "F"} else 1 for char in line)
            for line in str(value).split("\n")]


def _format_data_sheet(sheet):
    # Bound widths so identifiers and long interview responses do not create a giant canvas.
    # Wrapping and row heights preserve readability without changing the underlying values.
    minimum, maximum = 12, 48
    widths = [minimum] * sheet.max_column
    wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for row in sheet:
        for cell in row:
            if cell.value is None:
                continue
            width = max(_line_widths(cell.value), default=0)
            widths[cell.column - 1] = min(maximum, max(widths[cell.column - 1], width + 2))
            if isinstance(cell.value, str):
                # Headers and text are data, including strings beginning with '='.
                cell.data_type = "s"
                cell.alignment = wrap
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet:
        lines = 1
        for cell in row:
            if isinstance(cell.value, str):
                capacity = max(1, widths[cell.column - 1] - 2)
                lines = max(lines, sum(max(1, ceil(width / capacity)) for width in _line_widths(cell.value)))
        # Extremely long notes remain available in the cell; cap height to keep browsing practical.
        if row[0].row == 1:
            sheet.row_dimensions[1].height = min(90, max(30, lines * 15 + 10))
        elif lines > 1:
            sheet.row_dimensions[row[0].row].height = min(90, lines * 15 + 4)
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


@tool("export_excel",
      "Save the selected registered dataset as a new Excel workbook, preserving every row and column. "
      "After cleaning, use the cleaned dataset ID. Does not clean or remove fields. Returns the actual XLSX path.",
      EXCEL_SCHEMA)
def export_excel(ctx, args):
    ds = ctx.project.dataset(args["dataset_id"])
    ctx.project.allow.check_readable(ds.path)
    frame = read_dataset(ds)
    ctx.progress(25, "writing Excel workbook")
    out = ctx.project.new_artifact_dir("excel", ds.name)
    filename = args.get('filename') or (Path(ds.name).stem + '.xlsx')
    if Path(filename).name != filename or '/' in filename or '\\' in filename or filename in {'.','..'} or any(ord(c)<32 for c in filename):
        raise InvalidArguments('Use a filename without a directory path')
    if not filename.lower().endswith('.xlsx'): filename += '.xlsx'
    target = out / filename
    temporary = target.with_suffix(".pending.xlsx")
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name="Data", index=False)
            _format_data_sheet(writer.book["Data"])
        ctx.check()
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    ctx.project.audit({"tool": "export_excel", "dataset": ds.id, "artifact": str(out)})
    ctx.progress(100, "done")
    return {"dataset_id": ds.id, "path": str(target), "rows": len(frame),
            "columns": list(frame.columns), "format": "xlsx",
            "summary_text": f"Saved {len(frame):,} rows and {len(frame.columns)} columns to a new Excel workbook. The source is unchanged."}
