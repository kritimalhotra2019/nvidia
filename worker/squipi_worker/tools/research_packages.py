"""Small, typed adapters around reviewed research libraries. Imports are lazy.

No model-generated Python, R, formula expressions, shell commands or installation
is executed. Original inputs are read-only. Package availability is explicit.
"""
import importlib.metadata
import json
import math
import os
import re
import subprocess
import uuid
from pathlib import Path

import pandas as pd
from ..registry import tool, InvalidArguments
from .datasets import read_dataset

PACKAGES={'pandera':'pandera','presidio-analyzer':'presidio_analyzer','presidio-anonymizer':'presidio_anonymizer',
          'docling':'docling','pyfixest':'pyfixest','dowhy':'dowhy'}

def schema(props,required=()):return {'type':'object','properties':props,'required':list(required),'additionalProperties':False}
STRING={'type':'string','minLength':1,'maxLength':200}
COLUMNS={'type':'array','maxItems':30,'items':STRING}
DATASET={'dataset_id':STRING}

def require(*names):
    for name in names:
        try:importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:raise InvalidArguments('Research pack is not installed: '+name+'. Set up the research runtime in Me; no data was sent anywhere.')

def frame(ctx,args,columns=()):
    ds=ctx.project.dataset(args['dataset_id']);ctx.project.allow.check_readable(ds.path)
    if Path(ds.path).stat().st_size>400*1024*1024:raise InvalidArguments('This research adapter supports inputs up to 400 MB; select a smaller input explicitly')
    if ds.rows and ds.rows>250000:raise InvalidArguments('This research adapter supports at most 250,000 rows; select a dataset explicitly instead of silently sampling')
    df=read_dataset(ds)
    if len(df)>250000:raise InvalidArguments('This research adapter supports at most 250,000 rows')
    if not set(columns)<=set(df.columns):raise InvalidArguments('Requested columns are absent: '+', '.join(sorted(set(columns)-set(df.columns))))
    ctx.check();return df,ds

def names(values):
    if len(set(values))!=len(values) or any(not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*',v) for v in values):
        raise InvalidArguments('Estimator columns must have distinct simple names; rename unusual headers explicitly first. Formula/code expressions are not accepted.')

def clean(value):
    if isinstance(value,dict):return {str(k):clean(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [clean(v) for v in value]
    if hasattr(value,'item'):return clean(value.item())
    if isinstance(value,float) and not math.isfinite(value):return None
    return value

def result(ctx,kind,ds,payload):
    from ..objectives import digest
    out=ctx.project.new_artifact_dir(kind,ds.name)
    analysis_id='research_'+uuid.uuid4().hex[:12]
    data=clean({'kind':kind,'analysis_id':analysis_id,'dataset_id':ds.id,'source_sha256':digest(ds.path),'source_lineage':{d.id: digest(d.path) for d in ctx.project.dataset_lineage(ds.id)},**payload})
    path=ctx.project.write_json(out/'results.json',data)
    index_path=ctx.project.root/'package-analyses.json'
    index=json.loads(index_path.read_text()) if index_path.exists() else {}
    index[analysis_id]={'path':str(path),'sha256':digest(path)}
    ctx.project.write_json(index_path,index)
    return {**data,'results_path':str(path),'summary_text':kind.replace('_',' ')+': results saved; assumptions and limitations are included.'}

def load_package_analysis(project,identifier):
    from ..objectives import digest
    from ..security import ensure_within
    index_path=project.root/'package-analyses.json'
    index=json.loads(index_path.read_text()) if index_path.exists() else {}
    if identifier not in index:raise InvalidArguments('Unknown package analysis id')
    item=index[identifier];path=ensure_within(Path(item['path']),project.root)
    if digest(path)!=item['sha256']:raise InvalidArguments('Saved research result changed')
    data=json.loads(path.read_text())
    for source,signature in data['source_lineage'].items():
        ds=project.dataset(source);project.allow.check_readable(ds.path)
        if digest(ds.path)!=signature:raise InvalidArguments('An analysis source changed; recompute before reporting')
    return data

@tool('write_research_results_brief','Build a checked Word brief from saved Pandera/PyFixest/DoWhy/R-survey/Kumo analysis ids. Reuses exact saved results, units, sample exclusions and limitations. Verifies exported table cells. Narrative context must be grounded in the user or sources; no automatic causal certification.',
 schema({'analysis_ids':{'type':'array','minItems':1,'maxItems':8,'items':STRING},'title':STRING,'filename':STRING,'context':{'type':'string','maxLength':3000}},['analysis_ids','title','filename']))
def write_research_results_brief(ctx,args):
    from .export import export_document,_cell
    from docx import Document
    records=[load_package_analysis(ctx.project,i) for i in args['analysis_ids']]
    if Path(args['filename']).name!=args['filename'] or not args['filename'].endswith('.docx'):raise InvalidArguments('Use a plain .docx filename')
    sections=[]
    if args.get('context'):sections.append({'heading':'Research context','text':args['context']})
    for r in records:
        sections.append({'heading':r['kind'].replace('_',' ').title(),'text':'Source dataset: '+r['dataset_id']+'\nSpecification: '+json.dumps(r.get('specification',r.get('rules',{})),ensure_ascii=False)+'\nComplete observations: '+str(r.get('n',r.get('rows','not applicable')))+'\nExcluded for missing values: '+str(r.get('excluded_missing','see results'))})
        tables=[]
        if r['kind']=='fixed_effects_model':tables=[('Model coefficients',r['coefficients'])]
        elif r['kind']=='causal_effect':tables=[('Conditional effect estimate',[{'Effect':r['effect'],'Unit':r['unit'],'N':r['n'],'Placebo effect':r['placebo_effect']}])]
        elif r['kind']=='survey_summary':
            for measure,value in r['results'].items():tables.append((measure+' ('+r['specification']['units'][measure]+')',value if isinstance(value,list) else [value]))
        elif r['kind']=='data_validation':tables=[('Validation violations',r['violations'] or [{'Validation passed':r['passed']}])]
        elif r['kind']=='kumo_forecast':
            tables=[('Forecast summary — point predictions',r['series_summary'])]
            for chart in r.get('chart_paths',[]):
                sections.append({'heading':'Observed series and Kumo point forecast','chart_path':chart})
        elif r['kind']=='kumo_prediction':
            tables=[('Prediction scope — experimental',r['prediction_summary'])]
            if r.get('prediction_statistics'):tables.append(('Prediction summary — not validated accuracy',r['prediction_statistics']))
            if r.get('class_mapping'):tables.append(('Class probability column mapping',r['class_mapping']))
            sections.append({'heading':'Per-entity predictions','text':'The full per-query prediction table is saved separately: '+r['csv_path']+'\nScores are uncalibrated and were not evaluated on a holdout by this call.'})
        for title,rows in tables:
            if not rows:continue
            columns=list(rows[0]);values=[[format(v,'.6g') if isinstance(v,float) else v for v in [row.get(c) for c in columns]] for row in rows]
            if len(values)>200:raise InvalidArguments('Report table exceeds 200 rows; refine the requested scope explicitly')
            sections.append({'heading':title,'table':{'columns':columns,'rows':values}})
        sections.append({'heading':'Assumptions and limitations','text':'\n\n'.join(r.get('limitations',[])) or 'Validation checks only the declared rules; it does not establish scientific validity.'})
    exported=export_document(ctx,{'format':'docx','title':args['title'],'sections':sections,'sources':args['analysis_ids']})
    path=Path(exported['path']);path=path.rename(path.with_name(args['filename']))
    document=Document(path)
    intended=[[[str(c) for c in s['table']['columns']]]+[[_cell(v) for v in row] for row in s['table']['rows']] for s in sections if s.get('table')]
    actual=[[[c.text for c in row.cells] for row in t.rows] for t in document.tables]
    if intended!=actual:raise InvalidArguments('The Word brief failed its actual-table-cell check')
    return {**exported,'path':str(path),'docx_path':str(path),'analysis_ids':args['analysis_ids'],
        'source_dataset_ids':list(dict.fromkeys(i for r in records for i in r['source_lineage'])),
        'review_evidence':{'title':args['title'],'filename':path.name,'specifications':[r.get('specification',r.get('rules')) for r in records],
            'sections':[{**{k:v for k,v in section.items() if k!='table'},**({'table_columns':section['table']['columns'],'table_row_count':len(section['table']['rows'])} if section.get('table') else {})} for section in sections]},
        'verification':{'all_saved_result_cells_match':True,'independent_scientific_review':False},
        'summary_text':'Created a Word brief from the verified saved research results. Assumptions and limitations are included; scientific and visual review still matter.'}

@tool('research_capabilities','List installed research-library versions, document models and the R survey adapter. This never installs packages, downloads models, or contacts an external service.',schema({}))
def research_capabilities(ctx,args):
    installed={}
    for package in PACKAGES:
        try:installed[package]=importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:installed[package]=None
    rscript=os.environ.get('SQUIPI_RSCRIPT','')
    return {'packages':installed,'rscript_available':bool(rscript and Path(rscript).is_file()),
            'document_models_ready':Path(os.environ.get('SQUIPI_DOCUMENT_MODELS','/nonexistent')).joinpath('models-ready.json').is_file(),
            'network_services':['OpenAlex','Crossref','Bright Data (requires configured key and zones)'],
            'summary_text':'Research capabilities checked locally. Missing dependencies are not silently replaced.'}

@tool('validate_research_data','Validate declared column types, missingness, ranges, uniqueness and allowed values using Pandera. Returns violations without changing data. Rules must reflect the user/data dictionary, not invented assumptions.',
 schema({**DATASET,'rules':{'type':'array','minItems':1,'maxItems':50,'items':schema({
   'column':STRING,'dtype':{'type':'string','enum':['number','integer','string','boolean','datetime']},
   'nullable':{'type':'boolean'},'unique':{'type':'boolean'},'minimum':{'type':'number'},'maximum':{'type':'number'},
   'allowed':{'type':'array','maxItems':100,'items':{'type':['string','number','boolean','null']}}},['column','dtype'])}},['dataset_id','rules']))
def validate_research_data(ctx,args):
    require('pandera')
    import pandera.pandas as pa
    df,ds=frame(ctx,args,[r['column'] for r in args['rules']]);columns={}
    mapping={'integer':int,'number':float,'string':str,'boolean':bool,'datetime':'datetime64[ns]'}
    for rule in args['rules']:
        checks=[]
        if 'minimum' in rule:checks.append(pa.Check.ge(rule['minimum']))
        if 'maximum' in rule:checks.append(pa.Check.le(rule['maximum']))
        if 'allowed' in rule:checks.append(pa.Check.isin(rule['allowed']))
        dtype=None if rule['dtype']=='number' else mapping[rule['dtype']]
        if rule['dtype']=='number':checks.insert(0,pa.Check(lambda s:pd.api.types.is_numeric_dtype(s),error='must be numeric'))
        columns[rule['column']]=pa.Column(dtype,checks,nullable=rule.get('nullable',True),unique=rule.get('unique',False),coerce=False)
    if len(columns)!=len(args['rules']):raise InvalidArguments('Use one validation rule per column')
    failures=[]
    try:pa.DataFrameSchema(columns,strict=False).validate(df,lazy=True)
    except pa.errors.SchemaErrors as exc:
        # Avoid copying raw PII values into the planner or validation log.
        failures=[{k:clean(r.get(k)) for k in ['column','check','index']} for r in exc.failure_cases.to_dict('records')[:200]]
    return result(ctx,'data_validation',ds,{'passed':not failures,'rules':args['rules'],'violations':failures,'violations_are_bounded':len(failures)==200,'rows':len(df),'original_unchanged':True})

@tool('redact_text_pii','Redact detected PII inside explicitly selected text columns using local Presidio and its English NLP model. Preserves other columns and row count. Detection is imperfect and does not guarantee anonymity; inspect the audit. Never deletes interview notes as whole columns.',
 schema({**DATASET,'columns':COLUMNS,'entities':{'type':'array','maxItems':20,'items':STRING},'filename':STRING},['dataset_id','columns']))
def redact_text_pii(ctx,args):
    require('presidio-analyzer','presidio-anonymizer','en-core-web-sm')
    if not args['columns']:raise InvalidArguments('Choose at least one text column')
    from presidio_analyzer import AnalyzerEngine
    from presidio_analyzer.nlp_engine import NlpEngineProvider
    from presidio_anonymizer import AnonymizerEngine
    engine=NlpEngineProvider(nlp_configuration={'nlp_engine_name':'spacy','models':[{'lang_code':'en','model_name':'en_core_web_sm'}]}).create_engine()
    analyzer=AnalyzerEngine(nlp_engine=engine,supported_languages=['en']);anonymizer=AnonymizerEngine()
    df,ds=frame(ctx,args,args['columns']);df=df.copy();counts={}
    for col in args['columns']:
        if pd.api.types.is_numeric_dtype(df[col]):raise InvalidArguments('PII text redaction expects text columns; original was not modified')
        values=[]
        for value in df[col]:
            ctx.check()
            if pd.isna(value):values.append(value);continue
            found=analyzer.analyze(text=str(value),language='en',entities=args.get('entities'),score_threshold=.5)
            for item in found:counts[item.entity_type]=counts.get(item.entity_type,0)+1
            values.append(anonymizer.anonymize(text=str(value),analyzer_results=found).text)
        df[col]=values
    out=ctx.project.new_artifact_dir('text-pii',ds.name);name=args.get('filename','redacted-text.csv')
    if Path(name).name!=name or not name.lower().endswith('.csv'):raise InvalidArguments('Use a plain .csv filename')
    path=out/name;df.to_csv(path,index=False)
    derived=ctx.project.register_derived(path,path.name,len(df),len(df.columns),ds.id,list(df.columns))
    audit=ctx.project.write_json(out/'redaction-audit.json',{'source_id':ds.id,'columns':args['columns'],'rows':len(df),'detections_by_type':counts,'language':'en','guarantees_anonymity':False,'other_columns_preserved':True})
    return {'dataset_id':derived.id,'from_dataset_id':ds.id,'csv_path':str(path),'audit_path':str(audit),'rows':len(df),'columns':list(df.columns),'detections_by_type':counts,'summary_text':'Redacted detected PII within selected text columns. Other columns and row count retained. Review remaining re-identification risk.'}

@tool('fit_fixed_effects_model','Fit explicit OLS/fixed-effects or IV models using PyFixest, with heteroskedastic or one-way clustered inference. Simple column names only; no formula expressions. A fitted model alone does not establish causality.',
 schema({**DATASET,'outcome':STRING,'predictors':COLUMNS,'fixed_effects':COLUMNS,'cluster':STRING,'endogenous':COLUMNS,'instruments':COLUMNS,'outcome_unit':STRING},['dataset_id','outcome','predictors','outcome_unit']))
def fit_fixed_effects_model(ctx,args):
    require('pyfixest');import pyfixest as pf
    terms=[args['outcome']]+args['predictors']+args.get('fixed_effects',[])+([args['cluster']] if args.get('cluster') else [])+args.get('endogenous',[])+args.get('instruments',[])
    names(list(dict.fromkeys(terms)));df,ds=frame(ctx,args,terms)
    if not args['predictors'] and not args.get('endogenous'):raise InvalidArguments('Select predictors or an endogenous variable')
    if bool(args.get('endogenous'))!=bool(args.get('instruments')):raise InvalidArguments('IV needs explicit endogenous columns and instruments')
    sample=df.dropna(subset=list(set(terms)))
    if len(sample)<max(10,len(set(terms))+3):raise InvalidArguments('Too few complete observations for this specification')
    formula=args['outcome']+' ~ '+(' + '.join(args['predictors']) or '1')
    fixed=' + '.join(args.get('fixed_effects',[])) or '0'
    if args.get('fixed_effects') or args.get('endogenous'):formula+=' | '+fixed
    if args.get('endogenous'):formula+=' | '+' + '.join(args['endogenous'])+' ~ '+' + '.join(args['instruments'])
    vcov={'CRV1':args['cluster']} if args.get('cluster') else 'hetero'
    fit=pf.feols(formula,data=sample,vcov=vcov)
    table=fit.tidy().reset_index().to_dict('records')
    return result(ctx,'fixed_effects_model',ds,{'specification':args,'formula':formula,'n':len(sample),'excluded_missing':len(df)-len(sample),'coefficients':table,'inference':vcov,'cluster_count':int(sample[args['cluster']].nunique()) if args.get('cluster') else None,'limitations':['Inference depends on the stated specification. Fixed effects or instruments alone do not validate identification. Few clusters require additional care.']})

@tool('estimate_causal_effect','Use DoWhy for an explicitly specified binary-treatment backdoor linear-regression estimate and a placebo refutation. Requires stated confounders and identifying assumptions; never infer these from correlations. Reports complete-case exclusions. Refutation is a diagnostic, not proof of causality.',
 schema({**DATASET,'treatment':STRING,'outcome':STRING,'confounders':COLUMNS,'identifying_assumptions':{'type':'string','minLength':20,'maxLength':3000},'outcome_unit':STRING,'placebo_simulations':{'type':'integer','minimum':5,'maximum':30}},['dataset_id','treatment','outcome','confounders','identifying_assumptions','outcome_unit']))
def estimate_causal_effect(ctx,args):
    require('dowhy');from dowhy import CausalModel
    cols=[args['treatment'],args['outcome']]+args['confounders'];names(cols)
    df,ds=frame(ctx,args,cols);sample=df[cols].dropna()
    if not set(sample[args['treatment']].unique())<={0,1} or sample[args['treatment']].nunique()!=2:raise InvalidArguments('This adapter requires a binary 0/1 treatment with both groups present')
    if len(sample)<20:raise InvalidArguments('At least 20 complete observations are required')
    model=CausalModel(data=sample,treatment=args['treatment'],outcome=args['outcome'],common_causes=args['confounders'])
    identified=model.identify_effect(proceed_when_unidentifiable=False)
    estimate=model.estimate_effect(identified,method_name='backdoor.linear_regression')
    ctx.check()
    refutation=model.refute_estimate(identified,estimate,method_name='placebo_treatment_refuter',placebo_type='permute',num_simulations=args.get('placebo_simulations',10),random_state=17)
    return result(ctx,'causal_effect',ds,{'specification':args,'n':len(sample),'excluded_missing':len(df)-len(sample),'effect':float(estimate.value),'unit':args['outcome_unit'],'identified_estimand':str(identified),'placebo_effect':float(refutation.new_effect),'placebo_test':clean(refutation.refutation_result),'limitations':['The declared no-unmeasured-confounding assumption is not verified by this tool. Placebo refutation cannot prove identification. This adapter supports backdoor linear regression, not every DoWhy estimator.']})

@tool('survey_summary','Calculate survey-design-adjusted means, standard errors and confidence intervals through R survey. Requires a positive weight column; optional clusters, strata and group. Singleton-PSU errors are reported rather than silently adjusted. No silent imputation.',
 schema({**DATASET,'columns':COLUMNS,'weight_column':STRING,'cluster':STRING,'strata':STRING,'group_by':STRING,'units':{'type':'object','additionalProperties':STRING}},['dataset_id','columns','weight_column','units']))
def survey_summary(ctx,args):
    columns=args['columns']
    if not columns:raise InvalidArguments('Select at least one outcome')
    all_cols=columns+[args['weight_column']]+[args[k] for k in ['cluster','strata','group_by'] if args.get(k)]
    names(list(dict.fromkeys(all_cols)));df,ds=frame(ctx,args,all_cols)
    weight=pd.to_numeric(df[args['weight_column']],errors='coerce')
    if weight.isna().any() or not weight.gt(0).all():raise InvalidArguments('Survey weights must be observed and positive')
    if set(columns)-set(args['units']):raise InvalidArguments('State units for every selected outcome')
    executable=os.environ.get('SQUIPI_RSCRIPT','')
    if not executable or not Path(executable).is_file():raise InvalidArguments('R survey runtime is not configured; set it up in Me')
    out=ctx.project.new_artifact_dir('survey',ds.name);data=out/'input.csv';df[all_cols].to_csv(data,index=False)
    request=ctx.project.write_json(out/'request.json',{**args,'path':str(data)})
    script=Path(__file__).parent.parent/'survey_adapter.R'
    proc=subprocess.run([executable,'--vanilla',str(script),str(request)],capture_output=True,text=True,timeout=120,env={**os.environ,'R_DEFAULT_PACKAGES':'utils,stats,methods'})
    if proc.returncode:raise InvalidArguments('R survey could not estimate this design: '+proc.stderr[-1500:])
    response=json.loads(proc.stdout)
    return result(ctx,'survey_summary',ds,{'specification':args,**response,'limitations':['Variance estimates use the supplied survey design. These are descriptive population estimates, not programme-impact evidence.']})

@tool('extract_research_document','Extract an explicitly supplied document with local Docling, including structure and tables. Optional macOS OCR for scanned PDF. Writes full text, structured JSON and page-coverage evidence. Does not download models during research or send document contents to a cloud service.',
 schema({'path':{'type':'string','maxLength':4096},'ocr':{'type':'boolean'},'ocr_engine':{'type':'string','enum':['rapidocr','macOS']}},['path']))
def extract_research_document(ctx,args):
    require('docling')
    path=ctx.project.allow.check_readable(args['path'])
    if path.stat().st_size>50*1024*1024:raise InvalidArguments('Docling adapter limit is 50 MB')
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions,OcrMacOptions,RapidOcrOptions,OcrMode,AcceleratorOptions,AcceleratorDevice
    models=Path(os.environ.get('SQUIPI_DOCUMENT_MODELS','/nonexistent'))
    if path.suffix.lower()=='.pdf' and not (models/'models-ready.json').is_file():raise InvalidArguments('Document models need explicit setup before PDF extraction; no download was attempted')
    options=PdfPipelineOptions(artifacts_path=models if models.is_dir() else None,do_ocr=args.get('ocr',False),enable_remote_services=False,
        accelerator_options=AcceleratorOptions(num_threads=2,device=AcceleratorDevice.CPU))
    if args.get('ocr'):
        if args.get('ocr_engine','rapidocr')=='macOS':options.ocr_options=OcrMacOptions(lang=['en-US'],mode=OcrMode.FULL_PAGE)
        else:
            if not (models/'ocr-ready.json').is_file():raise InvalidArguments('Local OCR models need setup first; no model download was attempted')
            import onnxruntime
            onnxruntime.disable_telemetry_events()
            options.ocr_options=RapidOcrOptions(lang=['en'],backend='onnxruntime',mode=OcrMode.FULL_PAGE)
    converter=DocumentConverter(format_options={InputFormat.PDF:PdfFormatOption(pipeline_options=options)})
    converted=converter.convert(path,max_num_pages=100)
    ctx.check();document=converted.document;out=ctx.project.new_artifact_dir('docling',path.name)
    text=document.export_to_markdown();text_path=ctx.project.write_text(out/'text.md',text)
    structure=ctx.project.write_json(out/'structure.json',document.export_to_dict())
    expected=converted.input.page_count or len(document.pages);observed=len(document.pages)
    status=str(converted.status.value)
    text_pages={p.page_no for item in document.texts if item.text.strip() for p in item.prov}
    text_pages.update(p.page_no for item in document.tables if item.data.table_cells for p in item.prov)
    missing_pages=sorted(set(range(1,expected+1))-text_pages) if path.suffix.lower()=='.pdf' and expected else []
    complete=status=='success' and (not expected or observed==expected) and bool(text.strip()) and not missing_pages
    tables=[]
    for i,table in enumerate(document.tables):
        dest=out/('table-'+str(i+1)+'.csv');table.export_to_dataframe(doc=document).to_csv(dest,index=False);tables.append(str(dest))
    from ..objectives import digest
    coverage=ctx.project.write_json(out/'coverage.json',{'source_path':str(path),'source_sha256':digest(path),'expected_pages':expected,'extracted_pages':observed,'pages_without_extracted_text_or_tables':missing_pages,'complete':complete,'status':status,'ocr':args.get('ocr',False)})
    registered=ctx.project.register_document(Path(text_path),'text',observed,len(text.split()),Path(text_path))
    return {'document_id':registered.id,'source_path':str(path),'source_sha256':digest(path),'text_path':str(text_path),'structure_path':str(structure),'coverage_path':str(coverage),'table_paths':tables,'expected_pages':expected,'extracted_pages':observed,'pages_without_extracted_text_or_tables':missing_pages,'complete':complete,'ocr_engine':args.get('ocr_engine','rapidocr') if args.get('ocr') else None,'summary_text':f'Docling extraction {status}: {observed}/{expected} pages and {len(tables)} tables. Full text saved; extraction coverage does not certify that the model has read it.'}
