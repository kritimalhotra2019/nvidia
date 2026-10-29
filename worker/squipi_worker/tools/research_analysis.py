"""Schema-independent research summaries and evidence-backed briefs.

The model selects variables, grouping and units. Calculations and report tables
come from the same saved specification, not a second transcription by the model.
"""
import hashlib
import json
import math
import uuid
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

from ..registry import tool, InvalidArguments
from .datasets import ID_NAME
import re
from .stats import _load, _col, _numeric
from .document_style import STYLE_SCHEMA
from ..templates import TEMPLATE_ID_SCHEMA


def fingerprint(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()


def scalar(value):
    if value is None or pd.isna(value):return None
    return value.item() if hasattr(value,'item') else value


def identifier_column(name, dataset):
    if (dataset.overrides or {}).get(name,{}).get('role')=='id':return True
    normalized=re.sub(r'[\s-]+','_',str(name).strip().lower())
    # Counts such as number_of_children are quantities; account_number and
    # number_of_phone_ids still contain an identifier token and stay protected.
    if normalized.startswith('number_of_'):normalized=normalized[len('number_of_'):]
    return bool(ID_NAME.search(normalized))


def numeric_summary(values):
    values=pd.to_numeric(values,errors='coerce').replace([np.inf,-np.inf],np.nan)
    good=values.dropna();n=len(good)
    return {'n':n,'missing_or_invalid':len(values)-n,'mean':float(good.mean()) if n else None,
            'sd':float(good.std(ddof=1)) if n>1 else None,'median':float(good.median()) if n else None,
            'min':float(good.min()) if n else None,'max':float(good.max()) if n else None}


def _entity_labels(df,key):
    """A text column that maps one-to-one onto a numeric or code-like grouping key (country_id → country_iso3_code):
    its values label the groups, so a ranking reads UGA, IND, TZA rather than 800, 356, 834 (29 Oct 2026)."""
    try:
        k=df[key].dropna()
        if k.empty or k.nunique()<2:return None
        if not (pd.api.types.is_numeric_dtype(k) or k.astype(str).str.fullmatch(r'[A-Za-z]{0,3}\d+').mean()>0.9):return None
        best=None
        for c in df.columns:
            if c==key or not (pd.api.types.is_object_dtype(df[c]) or pd.api.types.is_string_dtype(df[c])):continue
            pair=df[[key,c]].dropna()
            if pair.empty:continue
            per_key=pair.groupby(key)[c].nunique()
            if (per_key<=1).all() and pair[c].nunique()>=0.9*pair[key].nunique() and pair[c].astype(str).str.len().median()<=40:
                score=pair[c].nunique()
                if best is None or score>best[0]:best=(score,c,pair.drop_duplicates(key).set_index(key)[c].astype(str).to_dict())
        return {'column':best[1],'map':best[2]} if best else None
    except Exception:
        return None


def normalize_summarize(args):
    """Deterministic repairs of a legal request written loosely (27 Oct 2026: a 9B model listed the same
    eight text columns as numeric and categorical, and asked for eight charts). Nothing is guessed: every
    change is reported in argument_notes."""
    notes=[]
    # Argument names are not column names (28 Oct 2026: a host repair put 'chart_palette' into columns).
    argument_names={'dataset_id','columns','categorical_columns','group_by','units','weight_column','paired_changes','chart_columns','charts','allow_sample','style','chart_palette','accent_color','heading_color','font_family','font_size','body_color','background_color','title','rank_by'}
    for key in ('columns','categorical_columns','group_by','chart_columns'):
        value=args.get(key)
        if isinstance(value,list):
            unique=list(dict.fromkeys(v for v in value if isinstance(v,str)))
            if len(unique)<len(value):notes.append(f'Repeated names in {key} were merged.')
            dropped=[v for v in unique if v in argument_names]
            if dropped:
                unique=[v for v in unique if v not in argument_names];notes.append(f'{", ".join(dropped)} in {key} are argument names, not columns; they were dropped.')
            args[key]=unique
    charts=args.get('chart_columns')
    if isinstance(charts,list) and len(charts)>4:
        args['chart_columns']=charts[:4];notes.append(f'Charts are limited to four measures; the first four were drawn ({", ".join(charts[:4])}). The others are in the tables.')
    return args,notes


@tool('summarize_research_data',
 'Calculate selected numeric/categorical summaries and paired changes overall and by group; save all results for write_research_brief. '
 'Works with registered CSV, Excel and Parquet datasets. Excludes numeric identifiers, reports per-variable valid/missing N, '
 'and optionally computes within-row paired changes or descriptive weighted means. Saves reusable tables and labelled charts '
 'for write_research_brief. No causal inference or automatic cleaning.',
 {'type':'object','properties':{
  'dataset_id':{'type':'string'},
  'columns':{'type':'array','items':{'type':'string'},'maxItems':40,'description':'Numeric levels only; [] for just changes or categories.'},
  'categorical_columns':{'type':'array','items':{'type':'string'},'maxItems':8},
  'group_by':{'type':'array','items':{'type':'string'},'maxItems':3},
  'units':{'type':'object','additionalProperties':{'type':'string'},'description':"Units by numeric column; use the user's supplied units."},
  'weight_column':{'type':'string'},
  'paired_changes':{'type':'array','maxItems':4,'description':'Add followup-minus-baseline measures; grouping applies.', 'items':{'type':'object','properties':{
    'baseline':{'type':'string'},'followup':{'type':'string'},'label':{'type':'string'},'unit':{'type':'string'}},
    'required':['baseline','followup','label','unit'],'additionalProperties':False}},
  'chart_columns':{'type':'array','maxItems':4,'items':{'type':'string'},'description':'Exact numeric/change/category measures to graph; must be selected in this analysis.'},'charts':{'type':'boolean'},'allow_sample':{'type':'boolean'},'style':STYLE_SCHEMA,
  'rank_by':{'type':'string','description':'With more than 50 groups: the measure whose mean ranks them (15 highest and 15 lowest shown).'}},
 'required':['dataset_id'],'additionalProperties':False},normalize=normalize_summarize)
def summarize_research_data(ctx,args):
    source=ctx.project.dataset(args['dataset_id']);ctx.project.allow.check_readable(source.path)
    source_hash=fingerprint(source.path)
    ancestors={d.id:{'sha256':fingerprint(ctx.project.allow.check_readable(d.path)),'overrides':d.overrides} for d in ctx.project.dataset_lineage(source.id)}
    df,ds,sampled=_load(ctx,args['dataset_id']);ctx.check()
    if sampled and not args.get('allow_sample',False):
        raise InvalidArguments('This dataset exceeds the full-analysis limit. Select/filter the data or explicitly allow_sample=true; no silent sampling was performed.')
    keys=[_col(df,c) for c in args.get('group_by',[])]
    if len(set(keys))!=len(keys):raise InvalidArguments('Grouping columns must be distinct')
    defaults=[] if args.get('categorical_columns') or args.get('paired_changes') else df.select_dtypes(include='number').columns.tolist()
    cols=[_col(df,c) for c in args.get('columns',defaults)]
    if 'columns' not in args:cols=[c for c in cols if not identifier_column(c,ds)]
    omitted=[]
    if len(cols)>12:
        # 28 Oct: a broad first pass over 21 columns was refused and cost two rounds. The first 12 are summarised with the
        # columns that read like outcomes (wellbeing, food security, income…) first, demographics last; the rest are named.
        from .structure import outcome_candidates
        from .datasets import analyse_column
        info={str(c):analyse_column(str(c),df[c],len(df)) for c in cols}
        preferred=[c for c in outcome_candidates(df,info,limit=12) if c in cols]
        ordered=preferred+[c for c in cols if c not in preferred]
        omitted=ordered[12:];cols=ordered[:12]
    # Identifier columns among the requested measures are set aside with a note (28 Oct: a whole call was refused
    # for one id column among four outcomes); only a request made of nothing but identifiers is refused.
    dropped_ids=[c for c in cols if identifier_column(c,ds)]
    if dropped_ids:
        cols=[c for c in cols if c not in dropped_ids]
        if not cols:raise InvalidArguments('Identifier columns are not quantities. Select outcomes/measures instead of averaging IDs.')

    categories=[_col(df,c) for c in args.get('categorical_columns',[])]
    if len(set(cols))!=len(cols) or len(set(categories))!=len(categories):raise InvalidArguments('Select distinct columns')
    notes=[]
    if omitted:notes.append(f"12 of {len(cols)+len(omitted)} requested columns were summarised; not included: {', '.join(omitted[:10])}{', …' if len(omitted)>10 else ''}. Ask for them by name for a second table.")
    if dropped_ids:notes.append(f"Identifier column{'s' if len(dropped_ids)!=1 else ''} {', '.join(dropped_ids)} {'were' if len(dropped_ids)!=1 else 'was'} set aside: not a quantity.")
    def numeric_share(c):
        values=df[c].dropna()
        return float(pd.to_numeric(values,errors='coerce').notna().mean()) if len(values) else 0.0
    # A column named in both lists is routed by its values, not refused: a text column is a category,
    # a number column is a measure. A "numeric" column whose values are mostly text is a category too.
    for c in [c for c in cols if c in categories]:
        if numeric_share(c)>=0.5:categories.remove(c);notes.append(f'{c} holds numbers and is summarised as a measure, not a category.')
        else:cols.remove(c);notes.append(f'{c} holds text and is summarised as a category, not a measure.')
    for c in [c for c in cols if 'columns' in args and numeric_share(c)<0.5]:
        cols.remove(c)
        if c not in categories:categories.append(c)
        notes.append(f'{c} has {"no" if numeric_share(c)==0 else "mostly non-"}numeric values; it is summarised as a category instead.')
    # Over 50 categories: an identifier-like column (nearly every row distinct) is skipped; a genuine
    # long list keeps its 20 most frequent categories and pools the rest as "Other" with the count stated.
    pooled={}
    for c in list(categories):
        overall=df[c].value_counts(dropna=False)
        if len(overall)<=50:continue
        valid=df[c].dropna()
        if len(valid) and valid.nunique()>=0.9*len(valid):
            categories.remove(c);notes.append(f'{c} has {valid.nunique()} distinct values in {len(valid)} rows (an identifier or free text); it was not tabulated. Use code_transcripts or search for free text.')
            continue
        keep=[v for v in overall.index if not pd.isna(v)][:20];pooled[c]=set(keep)
        notes.append(f'{c} has {len(overall)} categories; the 20 most frequent are listed and the remaining {len(overall)-len(keep)-int(overall.index.isna().any())} are pooled as Other.')
    if not cols and not categories and not args.get('paired_changes'):raise InvalidArguments('None of the selected columns can be summarised: '+' '.join(notes))
    frame=df.copy();units=args.get('units',{});measures=[]
    for c in cols:
        frame[c]=_numeric(df,c).replace([np.inf,-np.inf],np.nan)
        measures.append({'column':c,'label':c,'unit':units.get(c,'')})   # no unit given: say nothing rather than 'units unspecified' in every line
    labels=set(cols)
    for i,pair in enumerate(args.get('paired_changes',[])):
        if not pair['label'].strip() or pair['label'] in labels:raise InvalidArguments('Each change label must be nonempty and distinct from every other measure')
        labels.add(pair['label'])
        a,b=_col(df,pair['baseline']),_col(df,pair['followup'])
        if any(identifier_column(c,ds) for c in (a,b)):raise InvalidArguments('Paired changes require measures, not identifiers')
        name='__squipi_change_'+str(i)
        while name in frame:name+='_' 
        frame[name]=(_numeric(df,b)-_numeric(df,a)).replace([np.inf,-np.inf],np.nan)
        measures.append({'column':name,'label':pair['label'],'unit':pair['unit'],'baseline':a,'followup':b})
    if not measures and not categories:raise InvalidArguments('Select numeric measures or categorical columns')
    weights=None
    if args.get('weight_column'):
        weight=_col(df,args['weight_column']);weights=pd.to_numeric(df[weight],errors='coerce')
        if (weights.dropna()<0).any() or np.isinf(weights.dropna()).any():raise InvalidArguments('Weights must be finite and nonnegative')
    grouped=list(frame.groupby(keys,dropna=False,observed=True,sort=False)) if keys else []
    if len(grouped)>50:
        # Many groups (145 countries, 28 Oct 2026) are a ranking question, not a refusal: the 15 highest and 15 lowest
        # groups by the first measure's mean (or by rank_by) are reported, the rest pooled as one row, all named.
        rank_column=next((m['column'] for m in measures if m['label']==args.get('rank_by')),measures[0]['column'] if measures else None)
        if rank_column is None:raise InvalidArguments('More than 50 groups and no numeric measure to rank them by. Choose meaningful grouping variables or filter first.')
        ranked=sorted(grouped,key=lambda kv:(pd.isna(kv[1][rank_column].mean()),-(kv[1][rank_column].mean() if not pd.isna(kv[1][rank_column].mean()) else 0)))
        top,bottom=ranked[:15],[kv for kv in ranked[15:] if not pd.isna(kv[1][rank_column].mean())][-15:]
        shown={id(kv[1]) for kv in top+bottom};rest=[kv for kv in grouped if id(kv[1]) not in shown]
        pooled_frame=pd.concat([kv[1] for kv in rest]) if rest else None
        rank_label=next(m['label'] for m in measures if m['column']==rank_column)
        notes.append(f'{len(grouped)} groups by {", ".join(keys)}: the 15 highest and 15 lowest by mean {rank_label} are listed; the other {len(rest)} are pooled as one row. Rankings use group means over all rows of each group.')
        grouped=top+bottom+([((f'Other ({len(rest)} groups)',)*len(keys) if len(keys)>1 else f'Other ({len(rest)} groups)',pooled_frame)] if rest else [])
    groups=[(None,frame)]+grouped
    records=[];freq=[];sections=[];warnings=[];provenance_hashes={}
    for ancestor in ctx.project.dataset_lineage(source.id):
        source_path=Path(ancestor.path).resolve()
        try:source_path.relative_to(ctx.project.root)
        except ValueError:continue
        provenance_path=source_path.parent/'provenance.json'
        if not provenance_path.is_file():continue
        provenance=json.loads(ctx.project.allow.check_readable(provenance_path).read_text())
        if provenance.get('output_dataset')!=ancestor.id:continue
        provenance_hashes[str(provenance_path)]=fingerprint(provenance_path)
        sources=provenance.get('sources',[]);audit=provenance.get('join_audit')
        if audit and len(sources)==2:
            sections.append({'heading':'Sample construction and linkage','text':
                f"Join: {audit['how']}; keys: {', '.join(audit['keys'])}; relationship: {audit['relationship']}. "
                f"Output rows: {audit['output_rows']}. Matched counts refer to source rows, not necessarily distinct keys. Linkage is not evidence of causal identification.",
                'table':{'columns':['Source','Input rows','Matched rows','Unmatched rows'],
                    'rows':[[s['name'],audit[side+'_rows'],audit['matched_'+side+'_rows'],audit['unmatched_'+side+'_rows']] for s,side in zip(sources,['left','right'])]}})
        else:
            sections.append({'heading':'Sample construction and appended sources','text':f"Combined dataset: {ancestor.id}. Output rows: {provenance['output_rows']}.",
                'table':{'columns':['Source','Input rows'],'rows':[[s['name'],s['rows']] for s in sources]}})
    label_maps=[_entity_labels(df,k) for k in keys] if keys else []
    if any(label_maps):notes.append('Groups are labelled with the one-to-one text column of the grouping key ('+', '.join(f"{k}: {m['column']}" for k,m in zip(keys,label_maps) if m)+'); the key value follows in brackets.')
    for key,part in groups:
        def _label(i,v):
            if pd.isna(v):return '(missing)'
            m=label_maps[i] if i<len(label_maps) else None
            if m and v in m['map']:return f"{m['map'][v]} ({v})"
            return str(v)
        group='Overall' if key is None else ' / '.join(_label(i,v) for i,v in enumerate(key if isinstance(key,tuple) else (key,)))
        group_values={} if key is None else {c:scalar(v) for c,v in zip(keys,key if isinstance(key,tuple) else (key,))}
        for measure in measures:
            row={'group':group,'scope':'overall' if key is None else 'group','group_values':group_values,'group_rows':len(part),'variable':measure['label'],'unit':measure['unit'],**numeric_summary(part[measure['column']])}
            if weights is not None:
                values=part[measure['column']];w=weights.loc[part.index];valid=values.notna()&w.notna()&(w>0)
                row.update(weighted_n=int(valid.sum()),weighted_excluded=int((~valid).sum()),
                           weighted_mean=float(np.average(values[valid],weights=w[valid])) if valid.any() else None,
                           weight_sum=float(w[valid].sum()))
            records.append(row)
            if row['n']<5:warnings.append(f"{group}, {measure['label']}: valid N={row['n']}; do not promote a tiny subgroup as a headline finding.")
        for c in categories:
            if c in pooled:
                column=part[c].astype(object).where(part[c].isna()|part[c].isin(pooled[c]),'Other (pooled categories)')
                counts=column.value_counts(dropna=False)
            else:
                counts=part[c].value_counts(dropna=False)
                if len(counts)>50:raise InvalidArguments(f'{c} has over 50 categories; inspect whether it is an identifier or define an explicit grouping')
            valid_n=int(part[c].notna().sum())
            for value,n in counts.items():
                missing=pd.isna(value)
                freq.append({'group':group,'scope':'overall' if key is None else 'group','variable':c,'category':'(missing)' if missing else str(value),'n':int(n),
                             'nonmissing_denominator':valid_n,'percent_nonmissing':None if missing or valid_n==0 else 100*int(n)/valid_n})
    out=ctx.project.new_artifact_dir('research-analysis',ds.name)
    for measure in measures:
        rows=[r for r in records if r['variable']==measure['label']]
        columns=['Group','Valid N','Missing','Mean','SD','Median']
        table=[['Overall (all rows)' if r['scope']=='overall' else r['group'],r['n'],r['missing_or_invalid'],r['mean'],r['sd'],r['median']] for r in rows]
        sections.append({'heading':f"{measure['label']}{(' (' + measure['unit'] + ')') if measure['unit'] else ''}",'table':{'columns':columns,'rows':table}})
        if weights is not None:
            sections.append({'heading':f"Weighted descriptive means: {measure['label']}",'table':{'columns':['Group','Valid N','Excluded','Weight sum','Mean'],
                             'rows':[['Overall (all rows)' if r['scope']=='overall' else r['group'],r['weighted_n'],r['weighted_excluded'],r['weight_sum'],r['weighted_mean']] for r in rows]}})
    for c in categories:
        rows=[['Overall (all rows)' if r['scope']=='overall' else r['group'],r['category'],r['n'],r['nonmissing_denominator'],r['percent_nonmissing']] for r in freq if r['variable']==c]
        for start in range(0,len(rows),150):
            sections.append({'heading':f'Category distribution: {c}'+(f' (rows {start+1}–{min(start+150,len(rows))})' if len(rows)>150 else ''),
                             'table':{'columns':['Group','Category','Count','Valid N','Percent'],'rows':rows[start:start+150]}})
    chart_paths=[]; chart_subjects=[];label_map={}
    def plot_label(value):
        value=str(value)
        # Matplotlib cannot shape every script reliably. Use explicit lookup
        # labels rather than shipping unreadable glyph boxes or changing data.
        if any(ord(c)>127 and unicodedata.category(c)[0] in ('L','M') and 'LATIN' not in unicodedata.name(c,'') for c in value):
            for code,label in label_map.items():
                if label==value:return code
            code=f'Label {len(label_map)+1}';label_map[code]=value;return code
        return value
    if args.get('charts',True):
        from .document_style import normalize_style
        from .charts import _mpl
        style=normalize_style(args.get('style'));palette=style.get('chart_palette',['#347EA6','#A86E3D'])
        plt=_mpl(style)
        requested_charts=args.get('chart_columns')
        available={m['label'] for m in measures}|set(categories)
        if requested_charts is not None:
            dropped=[c for c in requested_charts if c not in available]
            if dropped:notes.append('Not charted (not among the summarised measures): '+', '.join(dropped)+'.')
            requested_charts=[c for c in requested_charts if c in available]
            if len(set(requested_charts))!=len(requested_charts):raise InvalidArguments('chart_columns must name distinct selected numeric/change/category measures')
        chart_measures=[m for m in measures if requested_charts is None or m['label'] in requested_charts]
        for i,measure in enumerate(chart_measures[:(4 if requested_charts is not None else 2)]):
            selected=[r for r in records if r['variable']==measure['label'] and (not keys or r['group_values']) and r['mean'] is not None]
            if not selected or len(selected)>20:continue
            fig,ax=plt.subplots(figsize=(8,4.5))
            if keys:
                ax.bar(range(len(selected)),[r['mean'] for r in selected],color=palette[i%len(palette)])
                ax.set_xticks(range(len(selected)),[f"{plot_label(r['group'])}\nvalid n={r['n']}" for r in selected],rotation=25 if len(selected)>4 else 0,ha='right' if len(selected)>4 else 'center')
                ax.set_ylabel(plot_label(measure['unit']));ax.set_title(f"{'Unweighted mean' if weights is not None else 'Mean'} {plot_label(measure['label'])}")
            else:
                values=frame[measure['column']].dropna()
                ax.hist(values,bins=min(20,max(1,int(np.sqrt(len(values))))),color=palette[i%len(palette)])
                ax.set_xlabel(plot_label(measure['unit']));ax.set_ylabel('Observation count');ax.set_title(f"{plot_label(measure['label'])} distribution (valid n={len(values)})")
            fig.text(.01,.01,'Observed values only; missing/invalid values excluded. Descriptive, not causal.',fontsize=8)
            fig.tight_layout(rect=(0,.05,1,1));path=out/f'chart-{i+1}.svg'
            fig.savefig(path);fig.savefig(path.with_suffix('.png'),dpi=160);plt.close(fig)
            chart_paths.append(str(path));chart_subjects.append(measure['label']);sections.append({'heading':f"Figure {len(chart_paths)}. {measure['label']}",'text':(f"Units: {measure['unit']}. " if measure['unit'] else '')+('Unweighted group means; valid observation counts are shown below each bar.' if keys else 'Distribution of unweighted observed values; valid N is in the chart title.'),'chart_path':str(path)})
        for c in [c for c in categories if requested_charts is None or c in requested_charts][:max(0,(4 if requested_charts is not None else 2)-len(chart_paths))]:
            selected=[r for r in freq if r['variable']==c and r['scope']=='overall' and r['percent_nonmissing'] is not None]
            if not selected or len(selected)>20:continue
            fig,ax=plt.subplots(figsize=(8,4.5));ax.bar(range(len(selected)),[r['n'] for r in selected],color=palette[0])
            ax.set_xticks(range(len(selected)),[plot_label(r['category']) for r in selected],rotation=25 if len(selected)>4 else 0,ha='right' if len(selected)>4 else 'center')
            ax.set_ylabel('Observation count');ax.set_title(f"{plot_label(c)} (nonmissing n={selected[0]['nonmissing_denominator']})")
            fig.tight_layout();path=out/f'chart-category-{len(chart_paths)+1}.svg';fig.savefig(path);fig.savefig(path.with_suffix('.png'),dpi=160);plt.close(fig)
            chart_paths.append(str(path));chart_subjects.append(c);sections.append({'heading':f'Figure {len(chart_paths)}. {c}','text':'Counts of nonmissing observations. Missing values are reported in the frequency table.','chart_path':str(path)})
    if label_map:
        sections.append({'heading':'Chart label lookup','table':{'columns':['Chart label','Original label'],'rows':[[k,v] for k,v in label_map.items()]}})
        warnings.append('Some chart labels use lookup codes because this renderer cannot reliably shape every script. Exact source labels are preserved in the chart label lookup table.')
    limitations=['These summaries describe observed data; comparisons alone do not establish cause and effect.',
                 'Missing observations are reported separately, never imputed. Category percentages use the nonmissing denominator.']
    if measures:limitations.append('Nonnumeric and nonfinite values are excluded per measure. SD uses N−1 and is undefined when valid N is below two. Tables display up to four significant figures; full-precision calculations are saved in the analysis file.')
    if weights is not None:limitations.append('Weighted means are descriptive only. No survey-design-adjusted standard errors or confidence intervals were calculated; strata and cluster design require a suitable survey estimator.')
    if sampled:limitations.append(f'Only the first {len(df)} rows were analysed with explicit sampling permission; this is not a full-dataset result.')
    findings=[]
    for measure in measures[:4]:
        overall=next(r for r in records if r['variable']==measure['label'] and r['scope']=='overall')
        if overall['n']:
            findings.append(f"{measure['label']}: overall mean {overall['mean']:.4g}{(' ' + measure['unit']) if measure['unit'] else ''} (valid N={overall['n']}; missing/invalid={overall['missing_or_invalid']}).")
        comparisons=[r for r in records if r['variable']==measure['label'] and r['scope']=='group' and r['n']>=5]
        if len(comparisons)>1:
            low=min(comparisons,key=lambda r:r['mean']);high=max(comparisons,key=lambda r:r['mean'])
            findings.append(f"Among groups with at least five valid observations, {high['group']} has the highest mean ({high['mean']:.4g}; N={high['n']}) and {low['group']} the lowest ({low['mean']:.4g}; N={low['n']}) for {measure['label']}{(', in ' + measure['unit']) if measure['unit'] else ''}. This is a descriptive comparison.")
    for c in categories:
        distribution=[r for r in freq if r['variable']==c and r['scope']=='overall']
        observed=[r for r in distribution if r['percent_nonmissing'] is not None]
        if observed:
            leading=sorted(observed,key=lambda r:r['n'],reverse=True)[:3]
            missing=sum(r['n'] for r in distribution if r['percent_nonmissing'] is None)
            findings.append(f"{c}: "+'; '.join(f"{r['category']} {r['n']}/{r['nonmissing_denominator']} ({r['percent_nonmissing']:.4g}%)" for r in leading)+f". Missing responses: {missing}."+(' Up to three leading categories shown here; the full distribution is in the table.' if len(observed)>3 else ''))
    if findings:sections.insert(0,{'heading':'Key descriptive findings','text':'\n\n'.join(findings)})
    analysis_id='analysis_'+uuid.uuid4().hex[:12]
    data={'kind':'squipi_research_analysis','version':1,'analysis_id':analysis_id,'dataset_id':ds.id,'dataset_name':ds.name,'dataset_path':ds.path,'source_sha256':fingerprint(ds.path),'source_lineage':ancestors,
          'specification':args,'rows_analyzed':len(df),'population_rows':ds.rows,'sampled':sampled,'records':records,'frequencies':freq,'provenance_sha256':provenance_hashes,
          'sections':sections,'chart_paths':chart_paths,'chart_measures':chart_subjects,'chart_label_map':label_map,'asset_sha256':{str(p):fingerprint(p) for path in chart_paths for p in (Path(path),Path(path).with_suffix('.png'))},'warnings':warnings,'limitations':limitations}
    if data['source_sha256']!=source_hash:raise InvalidArguments('Source changed during analysis; recompute from a stable input')
    for source_id,signature in ancestors.items():
        if fingerprint(ctx.project.dataset(source_id).path)!=signature['sha256']:raise InvalidArguments('An upstream source changed during analysis; recompute from stable inputs')
    path=ctx.project.write_json(out/'analysis.json',data)
    index_path=ctx.project.root/'research-analyses.json'
    index=json.loads(index_path.read_text()) if index_path.exists() else {}
    index[analysis_id]={'path':str(path),'sha256':fingerprint(path),'dataset_id':ds.id}
    ctx.project.write_json(index_path,index)
    pd.DataFrame(records).drop(columns=['group_values'],errors='ignore').to_csv(out/'numeric-statistics.csv',index=False)
    return {'analysis_id':analysis_id,'dataset_id':ds.id,'rows':len(df),'sampled':sampled,'results_path':str(path),
            'computed_measures':[m['label'] for m in measures],'chart_measures':chart_subjects,
            'analysis_specification':{k:v for k,v in args.items() if k in ('columns','categorical_columns','paired_changes','group_by','units','weight_column')},
            'data_notice':'Source dataset and columns are unchanged. Paired changes are calculated measures saved in this analysis_id, not new columns in dataset_id. write_research_brief includes every selected table from this analysis.',
            'statistics':records,'frequencies':freq,'chart_paths':chart_paths,'chart_png_paths':[str(Path(p).with_suffix('.png')) for p in chart_paths],'warnings':warnings,'key_findings':findings[:8],'argument_notes':notes,
            'summary_text':f'Calculated {len(measures)} measures and {len(categories)} categorical distributions for {len(df)} rows. Use write_research_brief with this analysis_id to include every requested table and chart without recopying values.'+(' Adjustments: '+' '.join(notes) if notes else '')}


def load_analysis(project,analysis_id):
    index_path=project.root/'research-analyses.json'
    index=json.loads(index_path.read_text()) if index_path.exists() else {}
    if analysis_id not in index:raise InvalidArguments('Unknown analysis_id; run summarize_research_data first')
    item=index[analysis_id];path=Path(item['path']).resolve()
    try:path.relative_to(project.root.resolve())
    except ValueError:raise InvalidArguments('Analysis must be inside the project')
    if fingerprint(path)!=item['sha256']:raise InvalidArguments('Saved analysis changed; recompute before exporting')
    data=json.loads(path.read_text());ds=project.dataset(data['dataset_id'])
    project.allow.check_readable(ds.path)
    if fingerprint(ds.path)!=data['source_sha256']:raise InvalidArguments('Source data changed; recompute before exporting')
    for source_id,signature in data.get('source_lineage',{}).items():
        ancestor=project.dataset(source_id);project.allow.check_readable(ancestor.path)
        if fingerprint(ancestor.path)!=signature['sha256'] or ancestor.overrides!=signature['overrides']:
            raise InvalidArguments('An upstream source or column definition changed; recompute before exporting')
    for asset,digest in {**data.get('asset_sha256',{}),**data.get('provenance_sha256',{})}.items():
        p=Path(asset).resolve()
        try:p.relative_to(project.root.resolve())
        except ValueError:raise InvalidArguments('Analysis assets must be inside the project')
        if not p.is_file() or fingerprint(p)!=digest:raise InvalidArguments('An analysis chart changed or is missing, or saved source provenance changed; recompute before exporting')
    return data


def _brief_narrative(ctx,data):
    """The brief reads like an analyst's note, not a statistics dump (29 Oct 2026, user: "the word doc makes no sense,
    it doesn't tell me anything"). Three sections built from saved results only: what each variable is (lexicon plus
    observed values), what stands out per measure (highest and lowest groups with figures, spread, missingness) in
    plain words with the specialist caveat, and the latest insight scan's findings with both readings."""
    from .meaning import variable_meaning, readings
    from .datasets import read_dataset, analyse_column, SAMPLE_ROWS
    out=[]
    try:
        ds=ctx.project.dataset(data['dataset_id']);df=read_dataset(ds,nrows=SAMPLE_ROWS);n=len(df)
        cols={str(c):analyse_column(str(c),df[c],n,(ds.overrides or {}).get(str(c),{}).get('role')) for c in df.columns}
    except Exception:
        df=None;cols={}
    # The saved analysis file keeps 'records' and 'specification'; the live tool result keeps 'statistics' and
    # 'analysis_specification'. Accept both.
    spec=data.get('specification') or data.get('analysis_specification') or {}
    records=data.get('records') or data.get('statistics') or []
    measures=list(dict.fromkeys(r['variable'] for r in records if r.get('scope')=='overall' and isinstance(r.get('variable'),str))) or [m for m in data.get('computed_measures',[])]
    group_by=spec.get('group_by') or []
    names=[c for c in group_by+measures if isinstance(c,str)]
    if names:
        lines=[variable_meaning(c,cols.get(c),df[c] if df is not None and c in df.columns else None) for c in names]
        out.append({'heading':'What the data is','text':'\n'.join(lines)})
    stats=[r for r in records if r.get('scope')=='group' and not str(r.get('group','')).startswith('Other (')]
    overall={r['variable']:r for r in records if r.get('scope')=='overall'}
    paras=[]
    for m in measures:
        o=overall.get(m);rows=[r for r in stats if r.get('variable')==m and isinstance(r.get('mean'),(int,float)) and r.get('n',0)>=5]
        if not o:continue
        fmt=lambda v:f"{v:,.3g}" if isinstance(v,(int,float)) else str(v)
        text=f"{m}: the typical value is {fmt(o.get('median'))} (median) with a mean of {fmt(o.get('mean'))} across {o.get('n',0):,} valid rows"
        miss=o.get('missing_or_invalid',0)
        if miss:text+=f"; {miss:,} rows have no value, so averages describe the rows that report one"
        text+='.'
        if rows and group_by:
            rows.sort(key=lambda r:r['mean'],reverse=True);top=rows[:3];bottom=rows[-3:][::-1]
            g=' / '.join(group_by)
            text+=f" By {g}, the highest are "+', '.join(f"{r['group']} ({fmt(r['mean'])}, n={r['n']})" for r in top)
            text+=f"; the lowest are "+', '.join(f"{r['group']} ({fmt(r['mean'])}, n={r['n']})" for r in bottom)+'.'
            if len(rows)>=2:
                spread=top[0]['mean']-bottom[0]['mean']
                text+=f" The gap between the top and bottom group is {fmt(spread)}, against a spread (SD) of {fmt(o.get('sd'))} across all rows."
            rd=readings({'kind':'group_difference','outcome':m,'by':g,'numbers':{}},cols,'Association in this sample, not a causal effect.')
            text+=f"\nIn plain words: {rd['plain']}\nSpecialist reading: {rd['specialist']}"
        paras.append(text)
    if paras:out.append({'heading':'What stands out','text':'\n\n'.join(paras)})
    try:
        root=ctx.project.root/'artifacts';cands=sorted(root.glob('*_insights_*/insights.json'),key=lambda p:p.stat().st_mtime,reverse=True)
        for p in cands:
            res=json.loads(p.read_text())
            if res.get('dataset_id')!=data['dataset_id']:continue
            lines=[]
            for i,f in enumerate(res.get('findings',[])[:8],1):
                rd=f.get('reading') or {};line=f"{i}. {f.get('title','')}. {f.get('detail','')}"
                if rd.get('plain'):line+=f"\n   In plain words: {rd['plain']}"
                if rd.get('specialist'):line+=f"\n   Specialist reading: {rd['specialist']}"
                lines.append(line)
            if lines:out.append({'heading':'Patterns in the data (from the insight scan)','text':'\n\n'.join(lines)+'\n\nSource: '+str(p)})
            break
    except Exception:
        pass
    return out


@tool('write_research_brief',
 'Create and inspect a Word brief from a saved summarize_research_data analysis_id. All requested tables/charts come directly from saved calculations. '
 'Checks the actual document title, every table cell and embedded chart count. Rejects changed source data or altered analysis. '
 'Use optional context/limitations for study-design facts from the user or observed sources; never invent facts. Does not apply an exact PowerPoint template.',
 {'type':'object','properties':{'analysis_id':{'type':'string'},'title':{'type':'string','minLength':1,'maxLength':200},
  'context':{'type':'string','maxLength':4000},'limitations':{'type':'array','items':{'type':'string'},'maxItems':12},
  'filename':{'type':'string','maxLength':120},'style':STYLE_SCHEMA,'template_id':TEMPLATE_ID_SCHEMA},
  'required':['analysis_id','title'],'additionalProperties':False})
def write_research_brief(ctx,args):
    from .export import export_document,_cell
    from docx import Document
    data=load_analysis(ctx.project,args['analysis_id'])
    if args.get('filename'):
        name=args['filename']
        if not name.strip() or Path(name).name!=name or '/' in name or '\\' in name or name in ('.','..'):raise InvalidArguments('filename must be a plain name')
    sections=[{'heading':'Scope and sample','text':f"Dataset: {data.get('dataset_name',data['dataset_id'])}. Analysed rows: {data['rows_analyzed']}. "+args.get('context','')}]
    sections+=_brief_narrative(ctx,data)
    sections+=data['sections']
    sections.append({'heading':'Interpretation and limitations','text':'\n\n'.join(data['limitations']+data['warnings']+args.get('limitations',[]))})
    # Keep report rounding deliberate and consistent with the narrative. The
    # legacy generic exporter rounds >=100 to integers, losing e.g. 105.5.
    sections=json.loads(json.dumps(sections))
    for section in sections:
        if 'table' in section:
            section['table']['rows']=[[format(v,'.4g') if isinstance(v,float) and math.isfinite(v) else v for v in row] for row in section['table']['rows']]
    result=export_document(ctx,{'format':'docx','title':args['title'],'sections':sections,'style':args.get('style'),'template_id':args.get('template_id'),
                                'sources':[data['dataset_id'],data['analysis_id']]})
    path=Path(result['path'])
    if args.get('filename'):
        name=args['filename']
        if not name.lower().endswith('.docx'):name+='.docx'
        path=path.rename(path.with_name(name));result['path']=str(path)
    document=Document(path)
    expected=[s['table'] for s in sections if s.get('table')]
    actual=[[[c.text for c in row.cells] for row in table.rows] for table in document.tables]
    intended=[[[str(c) for c in t['columns']]]+[[ _cell(v) for v in row] for row in t['rows']] for t in expected]
    title_present=any(p.text==args['title'] for p in document.paragraphs)
    chart_count=len(data['chart_paths'])
    valid=title_present and actual==intended and len(document.inline_shapes)==chart_count and not result['charts_referenced_only']
    verification={'passed':valid,'title_present':title_present,'all_requested_table_cells_match':actual==intended,
                  'expected_charts':chart_count,'embedded_charts':len(document.inline_shapes),'analysis_id':data['analysis_id'],
                  'scope':'structural and numerical cell verification; not independent visual or causal validation'}
    verify_path=ctx.project.write_json(path.parent/'verification.json',verification)
    if not valid:raise InvalidArguments('The exported brief failed verification; saved draft: '+str(path))
    review_evidence={'title':args['title'],'filename':path.name,'dataset':data.get('dataset_name',data['dataset_id']),
        'dataset_id':data['dataset_id'],'rows_analyzed':data['rows_analyzed'],'analysis_id':data['analysis_id'],
        'specification':data['specification'],'verification':verification,
        'sections':[{**{k:v for k,v in s.items() if k not in ('table','chart_path')},
                     **({'table_columns':s['table']['columns'],'table_row_count':len(s['table']['rows']),
                         'table_groups':list(dict.fromkeys(str(r[0]) for r in s['table']['rows']))[:50],
                         'table_rows':[[str(_cell(v)) for v in row] for row in s['table']['rows'][:20]],
                         'rows_omitted_from_review':max(0,len(s['table']['rows'])-20)} if s.get('table') else {}),
                     **({'embedded_chart':True} if s.get('chart_path') else {})} for s in sections]}
    return {**result,'docx_path':str(path),'dataset_id':data['dataset_id'],'analysis_id':data['analysis_id'],'verification':verification,'verification_path':str(verify_path),
            'review_evidence':review_evidence,
            'summary_text':f"Created and checked {path.name}: {len(expected)} tables and {chart_count} embedded charts from the saved analysis. Visual review is still needed."}


@tool('economic_appraisal','Compute discounted costs, benefits, net present value and benefit-cost ratio from explicit per-period flows. All flows must use one stated currency and price basis. Scenario arithmetic, not evidence of causal benefits.',
 {'type':'object','properties':{'costs':{'type':'array','items':{'type':'number'},'minItems':1,'maxItems':200},
  'benefits':{'type':'array','items':{'type':'number'},'minItems':1,'maxItems':200},'discount_rate':{'type':'number','minimum':0,'maximum':1},
  'currency':{'type':'string','minLength':1},'price_basis':{'type':'string','minLength':1},'first_period':{'type':'integer','minimum':0,'maximum':1}},
 'required':['costs','benefits','discount_rate','currency','price_basis','first_period'],'additionalProperties':False})
def economic_appraisal(ctx,args):
    costs=np.asarray(args['costs'],dtype=float);benefits=np.asarray(args['benefits'],dtype=float)
    if len(costs)!=len(benefits) or not np.isfinite(costs).all() or not np.isfinite(benefits).all():raise InvalidArguments('Costs and benefits need equal-length finite flows')
    if (costs<0).any() or (benefits<0).any():raise InvalidArguments('Use nonnegative costs and benefits; classify refunds/offsets explicitly')
    periods=np.arange(len(costs))+args['first_period'];discount=(1+args['discount_rate'])**periods
    pc=float((costs/discount).sum());pb=float((benefits/discount).sum())
    return {'present_value_costs':pc,'present_value_benefits':pb,'net_present_value':pb-pc,
            'benefit_cost_ratio':pb/pc if pc>0 else None,'currency':args['currency'],'price_basis':args['price_basis'],
            'discount_rate_per_period':args['discount_rate'],'first_period':args['first_period'],
            'summary_text':'Computed scenario appraisal. Use a discount rate consistent with period length and nominal/real price basis. Benefit assumptions and attribution require separate evidence.'}


@tool('plan_research_power','Approximate equal-sized two-arm sample size for a continuous outcome using normal theory. Requires explicit meaningful difference, outcome SD, alpha and target power. Supports a stated cluster design effect and attrition allowance. Not an exact small-sample, binary-outcome or complex-design power calculation.',
 {'type':'object','properties':{'difference':{'type':'number','minimum':0},'sd':{'type':'number','minimum':0},
  'alpha':{'type':'number','minimum':0.001,'maximum':0.2},'power':{'type':'number','minimum':0.5,'maximum':0.999},
  'design_effect':{'type':'number','minimum':1,'maximum':100},'attrition':{'type':'number','minimum':0,'maximum':0.9}},
 'required':['difference','sd','alpha','power'],'additionalProperties':False})
def plan_research_power(ctx,args):
    delta=args['difference'];sd=args['sd']
    if not math.isfinite(delta) or not math.isfinite(sd) or delta<=0 or sd<=0:raise InvalidArguments('Difference and SD must be positive finite numbers in the same outcome units')
    raw=2*((stats.norm.ppf(1-args['alpha']/2)+stats.norm.ppf(args['power']))*sd/delta)**2
    per_arm=math.ceil(raw*args.get('design_effect',1)/(1-args.get('attrition',0)))
    return {'per_arm':per_arm,'total':2*per_arm,'independent_complete_cases_per_arm':math.ceil(raw),'assumptions':'Two-sided normal approximation, equal allocation, independent units before stated design effect; continuous outcome.',
            'summary_text':f'Approximate recruitment target: {per_arm} per arm, {2*per_arm} total. Cluster counts and small-sample degrees of freedom require a design-specific power analysis.'}
