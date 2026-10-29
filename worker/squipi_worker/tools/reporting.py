"""reporting_checklist: deterministic coverage check of a trial/evaluation report against a reporting guideline.

CONSORT-SPI (CONSORT 2010 extension for social and psychological interventions, Grant et al. 2018),
TREND (Des Jarlais et al. 2004, non-randomised evaluations) and STROBE (von Elm et al. 2007, v4,
observational studies) are encoded as item lists with short paraphrases, never the full copyrighted
wording. Each item carries regular expressions that indicate it is addressed; matching runs over the
document's sentences and returns quoted evidence with offsets. The result is a reading aid: "addressed"
means the report contains language on that item, not that the content is adequate or truthful. Document
text is untrusted data."""
from __future__ import annotations

import json
import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ..registry import JobContext, tool, InvalidArguments
from ..security import ensure_within

CONTENT_TRUST = "untrusted_document_data"
MAX_EVIDENCE = 3
MAX_QUOTE = 320
MONTHS = r"january|february|march|april|may|june|july|august|september|october|november|december"
YEAR = r"(?:19|20)\d\d"
UNITS = r"participants|households|individuals|students|children|patients|schools|clusters|villages|respondents|women|men|firms|teachers|clinics|communities|observations|persons|people"

# Shared pattern fragments (compiled case-insensitively)
P_STATS = r"\b(intention[- ]to[- ]treat|per[- ]protocol|ANCOVA|OLS|mixed[- ]effects|multilevel|regression (model|analysis|analyses)|logistic regression|linear regression|t[- ]tests?|analysis of (co)?variance|difference[- ]in[- ]differences|fixed effects)\b"
P_STATS_WEAK = [r"\bconfidence intervals?\b", r"\bstandard errors?\b", r"\bstatistical (analysis|analyses|methods|tests?)\b", r"\bp[- ]values?\b"]
P_ESTIMATE = r"\b(95 ?% ?(confidence interval|CI)|effect size|treatment effect|difference (in|of) means?|risk ratio|odds ratio|standardi[sz]ed (mean )?difference|cohen'?s d|regression coefficient|coefficient of)\b"
P_SUBGROUP = r"\b(subgroup|sensitivity|exploratory|moderat(ion|or)|heterogeneity|interaction) (analys[ie]s|effects?|results?|tests?)\b|\bpre[- ]specified\b|\bpost[- ]hoc\b"
P_LIMITS = r"\blimitations?\b|\b(sources? of )?(potential |possible )?bias(es)?\b|\bimprecision\b|\bcaveats?\b"
P_GENERAL = r"\b(generali[sz]abilit|external validity|transferab(le|ility)|applicab(le|ility) to other|generali[sz]e(d)? (to|beyond))"
P_FUNDING = r"\b(fund(ed|ing|er)s?|financ(ed|ial support)|grant (number|no\.?|from)|supported by|sponsor(ed|s)?)\b"
P_SETTING = r"\b(recruited|conducted|took place|carried out|implemented|collected) (in|at|across|from)\b[^.]{0,60}\b(districts?|villages?|schools?|clinics?|hospitals?|communit(y|ies)|regions?|countr(y|ies)|provinces?|cities|city|towns?|camps?|sites?)\b"
P_SETTING_WEAK = [r"\bdistricts?\b", r"\bvillages?\b", r"\bschools?\b", r"\bclinics?\b", r"\bcommunit(y|ies)\b", r"\bsettings?\b"]
P_ELIGIBLE = r"\b(eligib(le|ility)|inclusion criteria|exclusion criteria|were eligible|were excluded if)\b"
P_SAMPLE_SIZE = r"\bsample size\b|\bpower calculation\b|\b(80|90) ?% power\b|\bminimum detectable effect\b|\bpowered to detect\b"
P_NUMBERS = r"\b\d[\d,]* (" + UNITS + r") (were |was )?(randomi[sz]ed|allocated|assigned|enrolled|recruited|analy[sz]ed|included|surveyed|interviewed|completed|responded|eligible)\b|\b(randomi[sz]ed|allocated|assigned|enrolled) (n ?= ?\d|\d)"
P_FLOW = P_NUMBERS + r"|\bflow (diagram|chart)\b|\bCONSORT (diagram|flow)\b"
P_ATTRITION = r"\b(lost to follow[- ]up|attrition|dropped out|drop[- ]?outs?|withdrew|withdrawals?|excluded after (randomi[sz]ation|assignment)|non[- ]response)\b"
P_DATES = r"\b(recruit(ed|ment)|enrol(led|ment)|follow[- ]up|baseline|endline|data (were |was )?collected|surveys? (were |was )?(conducted|administered|fielded)|fieldwork)\b[^.]{0,80}\b(" + YEAR + r"|" + MONTHS + r")\b"
P_BASELINE = r"\bbaseline (characteristics|table|data|balance|values?|means?)\b|\btable \d[^.]{0,40}\bbaseline\b|\bbalance(d)? (at baseline|table)\b|\bwere (well )?balanced\b|\bcomparable at baseline\b"
P_ANALYSED = r"\b(analy[sz]ed|analysis (sample|included)|included in the (primary |main )?analysis)\b[^.]{0,60}\b\d[\d,]*\b|\b\d[\d,]* (" + UNITS + r") (were |was )?(analy[sz]ed|included in the (primary |main )?analysis)\b"
P_HARMS = r"\b(adverse (events?|effects?)|harms?|unintended (effects?|consequences?)|side[- ]effects?|safety (concerns?|monitoring|outcomes?)|negative (effects?|consequences?))\b"
P_INTERPRET = r"\b(interpret(ation|ed|ing)|in (light|view) of|consistent with (the )?(results|findings|prior evidence|previous studies)|benefits? and (harms?|costs?)|policy implications?|implications for)\b"
P_REGISTRATION = r"\b(registered|registration|registry|trial register|ISRCTN|NCT ?\d|ClinicalTrials\.gov|AEA RCT Registry|RIDIE|OSF|PACTR|CTRI|pre[- ]?registered|pre[- ]analysis plan)\b"
P_OBJECTIVES = r"\b(objectives?|aims?) (of|was|were|is|are)\b|\bwe (aimed|sought|set out|hypothesi[sz]ed)\b|\bhypothes[ie]s\b|\bresearch questions?\b|\bthis (study|trial|evaluation|paper) (asks|tests|examines|evaluates|assesses)\b"
P_BACKGROUND = r"\brationale\b|\bprior (evidence|research|studies|evaluations)\b|\bprevious (studies|research|evaluations|trials)\b|\bevidence (gap|base)\b|\bliterature\b"
P_OUTCOMES = r"\b(primary|secondary) outcomes?\b|\boutcomes? (was|were) (measured|assessed|defined|collected)\b|\boutcome measures?\b"
P_INTERVENTION = r"\bintervention (consisted|comprised|included|was delivered|delivered|involved)\b|\bcontrol (group|arm|condition) received\b|\bthe programme (consisted|comprised|provided|delivered|offered)\b|\bthe program (consisted|comprised|provided|delivered|offered)\b|\bcomparison (group|arm) received\b"
P_INTERVENTION_WEAK = [r"\bintervention\b", r"\bcomparator\b|\busual care\b|\bcontrol (group|arm|condition)\b", r"\bprogramme\b|\bprogram\b"]
P_PROVIDERS = r"\b(delivered|facilitated|implemented|administered|provided) by\b|\b(facilitators?|providers?|field officers?|community health workers?|trained staff|enumerators?)\b[^.]{0,60}\b(training|trained|experience|qualifications?)\b"
P_FIDELITY = r"\b(fidelity|adherence|as intended|take[- ]up|uptake|attendance rates?|compliance|dosage|implementation (quality|monitoring))\b"
P_BLINDING = r"\b(blind(ed|ing)?|mask(ed|ing))\b[^.]{0,60}\b(participants|assessors|enumerators|outcome|analysts|providers|staff|investigators)\b|\b(open[- ]label|not (possible|feasible) to blind|unblinded|no blinding|could not be blinded)\b"
P_UNIT = r"\bunit of (analysis|assignment|randomi[sz]ation)\b|\b(clustered|clustering|cluster[- ]robust|intra[- ]?cluster|ICC|nested within|standard errors (were )?clustered)\b"
P_ASSIGNMENT = r"\b(assign(ed|ment) (method|by|using|based on|was)|allocat(ed|ion) (by|using|based on)|non[- ]?randomi[sz]ed|quasi[- ]experimental|matched (comparison|control)|propensity|difference[- ]in[- ]differences|regression discontinuity|self[- ]select|eligibility cut[- ]?off|phased roll[- ]?out)"
P_EQUIVALENCE = r"\b(baseline (equivalence|balance|differences?)|comparable at baseline|balanced at baseline|standardi[sz]ed differences?|no (statistically )?significant differences? at baseline)\b"
P_EVIDENCE = r"\b(overall evidence|body of evidence|current evidence|weight of (the )?evidence|in the context of (existing|prior|current) (evidence|literature|research))\b|\bprior (studies|evidence|research|evaluations) (found|show|suggest|report)"
P_ETHICS = r"\b(ethic(s|al) (approval|committee|review|board|clearance)|IRB|institutional review board|informed consent)\b"
P_STAKEHOLDER = r"\b(stakeholder|community|participant|patient|public) (involvement|engagement|advisory|consultation|consulted)\b|\binvolved (stakeholders|community members|participants|beneficiaries) in\b"
P_PROTOCOL = r"\b(protocol|pre[- ]analysis plan|PAP)\b[^.]{0,80}\b(available|accessed|published|online|appendix|supplement|request|registry|registered)\b|\bfull protocol\b"
P_CHANGES = r"\b(protocol|design|methods?) (amendments?|changes?|deviations?)\b|\bchanges? (to|in) the (protocol|trial design|design|methods|eligibility)\b|\bno changes (were made )?to the (protocol|design|methods)\b|\bdeviat(ed|ions?) from the (protocol|plan)\b"
P_DESIGN_TRIAL = r"\b(parallel[- ]group|cluster[- ]randomi[sz]ed|factorial|crossover|stepped[- ]wedge|two[- ]arm|three[- ]arm|individually randomi[sz]ed|wait[- ]?list)\b|\ballocation ratio\b|\b1 ?: ?1\b"
P_SEQUENCE = r"\bsequence (was )?generated\b|\bsequence generation\b|\b(computer[- ]generated|random number (table|generator)|coin toss|lottery|public lottery|randomi[sz]ation (was )?(done|performed|carried out|conducted) (by|using|in))\b|\b(stata|excel|python|r)\b[^.]{0,30}\b(random|randomi[sz]ation)\b"
P_RESTRICTION = r"\b(block|stratified|stratification|restricted|minimi[sz]ation|pair[- ]?matched|re[- ]?randomi[sz]ation)\b[^.]{0,40}\brandomi[sz]|\brandomi[sz](ed|ation)\b[^.]{0,40}\b(blocks?|strata|stratum|stratified|pairs?)\b"
P_CONCEAL = r"\b(allocation )?conceal(ed|ment)\b|\bsealed (opaque )?envelopes?\b|\bcentral(i[sz]ed)? (randomi[sz]ation|allocation)\b|\bconcealed from\b"
P_IMPLEMENT = r"\b(generated the (allocation|random(i[sz]ation)?) sequence|enrolled participants|assigned participants|enrolled (the )?(households|students|patients)|assigned (the )?(households|students|patients))\b|\bthe (statistician|research team|field team|enumerators?|principal investigator|implementing partner|NGO) (generated|enrolled|assigned|drew|conducted the (lottery|randomi[sz]ation))\b"
P_STOPPED = r"\b(trial|study|evaluation) (was )?(ended|stopped|terminated|discontinued|halted)\b|\bearly (stopping|termination)\b|\binterim analys[ie]s\b|\bstopping (rule|guideline)s?\b"
P_BINARY = r"\b(absolute (risk|difference)|relative risk|risk difference|number needed to treat|percentage points?)\b"
P_DECLARE = r"\b(conflicts? of interest|competing interests?|declaration of interests?|nothing to declare)\b"
P_KEY = r"\b(key (results|findings)|main findings?|in summary|we find that|we found that|this study (shows|found|finds)|our (results|findings) (show|suggest|indicate))\b"
P_OBS_DESIGN = r"\b(cross[- ]sectional|cohort|case[- ]control|longitudinal|panel|observational|retrospective|prospective|repeated cross[- ]section)\b[^.]{0,30}\b(study|survey|design|analysis|data)\b|\bstudy design\b"
P_VARIABLES = r"\b(exposure|outcome|predictor|confounder|covariate|effect modifier)s?\b[^.]{0,60}\b(defined|measured|was|were|include|included)\b"
P_SOURCES = r"\b(data (source|were obtained|were collected|collection)|questionnaire|administrative (data|records)|survey instrument|enumerators?|measured (using|with|by)|self[- ]report(ed)?|interviewer[- ]administered)\b"
P_BIAS = r"\b(bias(es)?|confounding|selection effects?)\b[^.]{0,60}\b(address|minimi[sz]|reduce|account|adjust|control|limit)|\b(selection|recall|measurement|response|non[- ]response|reporting|social desirability) bias\b"
P_STUDY_SIZE = P_SAMPLE_SIZE + r"|\bstudy size\b|\b\d[\d,]* (" + UNITS + r") (were|was) (sampled|surveyed|interviewed|selected|drawn)\b"
P_QUANT = r"\b(categori[sz]ed|grouped into|quintiles?|terciles?|tertiles?|quartiles?|cut[- ]?offs?|thresholds?|log[- ]transformed|continuous variables?|winsori[sz]ed|top[- ]coded)\b"
P_DESCRIPTIVE = r"\b(characteristics of (the )?(study )?(participants|respondents|sample|households)|descriptive (statistics|table)|table \d[^.]{0,30}\b(characteristics|descriptive)|mean age|average age)\b|\b\d+(\.\d+)? ?(%|percent) (were|of (the )?(respondents|participants|sample|households))\b"
P_OUTCOME_DATA = r"\b(number of (outcome )?events|outcome (data|measures?) (are|were|was)|prevalence (of|was)|incidence (of|was)|\d+(\.\d+)? ?% (reported|had|were|experienced))\b"
P_MAIN = r"\b(unadjusted|adjusted|crude)\b[^.]{0,40}\b(estimates?|odds ratio|risk|coefficients?|association|differences?)|\b(95 ?% ?(CI|confidence interval)|odds ratio|risk ratio|prevalence ratio|regression coefficient|hazard ratio)\b"
P_INTERPRET_OBS = r"\b(cautious|plausible (explanation|mechanism)|multiplicity|similar studies|consistent with|interpret(ation|ed|ing)|in (light|view) of)\b"
P_TITLE_RANDOM = r"\brandomi[sz]ed\b|\brandom(ly)? (assign|allocat)|\btrial\b"
P_TITLE_NONRANDOM = r"\bnon[- ]?randomi[sz]ed\b|\bquasi[- ]experimental\b|\bevaluation\b|\bcomparison\b|\bimpact\b"
P_TITLE_OBS = r"\bcross[- ]sectional\b|\bcohort\b|\bcase[- ]control\b|\bsurvey\b|\bobservational\b|\blongitudinal\b"
P_ABSTRACT = r"^\s*(abstract|structured summary|executive summary|summary)\b"


def _item(id_: str, section: str, text: str, strong: Any = None, weak: Any = None, scope: str = "body") -> Dict[str, Any]:
    strong = [strong] if isinstance(strong, str) else list(strong or [])
    weak = [weak] if isinstance(weak, str) else list(weak or [])
    return {"id": id_, "section": section, "item": text, "strong": strong, "weak": weak, "scope": scope}


CONSORT_SPI = [
    _item("1a", "Title and abstract", "Title identifies the study as a randomised trial", P_TITLE_RANDOM, scope="title"),
    _item("1b", "Title and abstract", "Structured summary of design, methods, results and conclusions", P_ABSTRACT, scope="head"),
    _item("2a", "Introduction", "Scientific background, rationale and the theory of change", P_BACKGROUND, [r"\bbackground\b", r"\bintroduction\b"]),
    _item("2b", "Introduction", "Specific objectives or hypotheses", P_OBJECTIVES),
    _item("3a", "Methods: trial design", "Trial design (e.g. parallel, cluster, factorial) and allocation ratio", P_DESIGN_TRIAL, [r"\btrial design\b", r"\bcluster\b"]),
    _item("3b", "Methods: trial design", "Important changes to methods after trial start, with reasons", P_CHANGES),
    _item("4a", "Methods: participants", "Eligibility criteria for participants (and clusters/settings where applicable)", P_ELIGIBLE),
    _item("4b", "Methods: participants", "Settings and locations where the data were collected", P_SETTING, P_SETTING_WEAK),
    _item("5", "Methods: interventions", "Interventions for each arm, with enough detail to allow replication, including comparator", P_INTERVENTION, P_INTERVENTION_WEAK),
    _item("5a", "Methods: interventions", "Who delivered the intervention and how they were trained", P_PROVIDERS),
    _item("5b", "Methods: interventions", "Extent to which the intervention was delivered and received as intended (fidelity, take-up)", P_FIDELITY),
    _item("6a", "Methods: outcomes", "Pre-specified primary and secondary outcomes, how and when assessed", P_OUTCOMES),
    _item("6b", "Methods: outcomes", "Any changes to outcomes after the trial started, with reasons", r"\b(changes? to|changed) (the )?(trial |primary |secondary )?outcomes?\b|\boutcomes? (were|was) (not )?changed\b"),
    _item("7a", "Methods: sample size", "How the sample size was determined", P_SAMPLE_SIZE),
    _item("7b", "Methods: sample size", "Interim analyses and stopping guidelines, when applicable", r"\binterim analys[ie]s\b|\bstopping (rule|guideline)s?\b"),
    _item("8a", "Methods: randomisation", "Method used to generate the random allocation sequence", P_SEQUENCE),
    _item("8b", "Methods: randomisation", "Type of randomisation and any restriction (blocking, stratification)", P_RESTRICTION, [r"\bblocks?\b", r"\bstrat(a|um|ified|ification)\b"]),
    _item("9", "Methods: randomisation", "Allocation concealment mechanism", P_CONCEAL),
    _item("10", "Methods: randomisation", "Who generated the sequence, enrolled participants and assigned them", P_IMPLEMENT),
    _item("11", "Methods: blinding", "Who was blinded after assignment (participants, providers, assessors), or why blinding was not possible", P_BLINDING, [r"\bblind(ed|ing)?\b", r"\bmask(ed|ing)\b"]),
    _item("12a", "Methods: statistics", "Statistical methods for primary and secondary outcomes, including the analysis population", P_STATS, P_STATS_WEAK),
    _item("12b", "Methods: statistics", "Methods for additional analyses: subgroup, adjusted, mediation or moderation", P_SUBGROUP),
    _item("13a", "Results: participant flow", "Numbers assigned, receiving the intervention and analysed for the primary outcome, per arm", P_FLOW),
    _item("13b", "Results: participant flow", "Losses and exclusions after randomisation, with reasons", P_ATTRITION),
    _item("14a", "Results: recruitment", "Dates defining the periods of recruitment and follow-up", P_DATES),
    _item("14b", "Results: recruitment", "Why the trial ended or was stopped, when applicable", P_STOPPED),
    _item("15", "Results: baseline data", "Table of baseline characteristics for each arm", P_BASELINE),
    _item("16", "Results: numbers analysed", "Number of participants in each analysis per arm and whether analysis was by original assignment", [P_ANALYSED, r"\bintention[- ]to[- ]treat\b|\bper[- ]protocol\b|\bas (randomi[sz]ed|assigned)\b"]),
    _item("17a", "Results: outcomes", "For each outcome, results per arm, the estimated effect size and its precision (95% CI)", P_ESTIMATE),
    _item("17b", "Results: outcomes", "For binary outcomes, both absolute and relative effect sizes", P_BINARY),
    _item("18", "Results: ancillary analyses", "Results of other analyses (subgroup, adjusted), distinguishing pre-specified from exploratory", P_SUBGROUP),
    _item("19", "Results: harms", "Important harms or unintended effects in each arm", P_HARMS),
    _item("20", "Discussion", "Trial limitations: sources of bias, imprecision, multiplicity", P_LIMITS),
    _item("21", "Discussion", "Generalisability (external validity) of the findings", P_GENERAL),
    _item("22", "Discussion", "Interpretation consistent with results, balancing benefits and harms and other evidence", P_INTERPRET, [r"\bdiscussion\b", r"\bconclusions?\b"]),
    _item("23", "Other information", "Registration number and name of the trial registry", P_REGISTRATION),
    _item("24", "Other information", "Where the full protocol or pre-analysis plan can be accessed", P_PROTOCOL),
    _item("25", "Other information", "Sources of funding and role of funders", P_FUNDING),
    _item("26", "Other information", "Stakeholder involvement in design, conduct or reporting (CONSORT-SPI)", P_STAKEHOLDER),
    _item("27", "Other information", "Ethical approval and consent", P_ETHICS),
    _item("28", "Other information", "Declaration of interests", P_DECLARE),
]

TREND = [
    _item("1", "Title and abstract", "Title/abstract state how units were allocated and the study design", [P_TITLE_NONRANDOM + r"|" + P_TITLE_RANDOM], scope="title"),
    _item("1b", "Title and abstract", "Structured abstract or summary", P_ABSTRACT, scope="head"),
    _item("2", "Introduction", "Scientific background, rationale and theory behind the intervention", P_BACKGROUND, [r"\bbackground\b", r"\btheory\b"]),
    _item("3", "Methods: participants", "Eligibility criteria, recruitment method and settings", [P_ELIGIBLE, P_SETTING], P_SETTING_WEAK),
    _item("4", "Methods: interventions", "Details of the intervention and comparison condition (content, delivery, duration)", P_INTERVENTION, P_INTERVENTION_WEAK),
    _item("4b", "Methods: interventions", "Who delivered the intervention and implementation fidelity", [P_PROVIDERS, P_FIDELITY]),
    _item("5", "Methods: objectives", "Specific objectives and hypotheses", P_OBJECTIVES),
    _item("6", "Methods: outcomes", "Primary and secondary outcome measures and how/when they were collected", P_OUTCOMES),
    _item("7", "Methods: sample size", "How sample size was determined", P_SAMPLE_SIZE),
    _item("8", "Methods: assignment", "Unit of assignment and the method used to assign units to conditions", P_ASSIGNMENT, [r"\bassign(ed|ment)\b", r"\ballocat(ed|ion)\b"]),
    _item("9", "Methods: blinding", "Whether participants, providers or assessors were blinded, or why not", P_BLINDING, [r"\bblind(ed|ing)?\b", r"\bmask(ed|ing)\b"]),
    _item("10", "Methods: unit of analysis", "Unit of analysis and how clustering was handled", P_UNIT),
    _item("11", "Methods: statistics", "Statistical methods for primary outcomes, additional analyses and handling of missing data", P_STATS, P_STATS_WEAK + [r"\bmissing data\b"]),
    _item("12", "Results: participant flow", "Flow of participants through each stage per condition", P_FLOW),
    _item("12b", "Results: participant flow", "Losses, exclusions and attrition with reasons", P_ATTRITION),
    _item("13", "Results: recruitment", "Dates of recruitment and follow-up", P_DATES),
    _item("14", "Results: baseline data", "Baseline demographic and clinical characteristics per condition", P_BASELINE),
    _item("15", "Results: baseline equivalence", "Baseline comparisons of study conditions and of those lost versus retained", P_EQUIVALENCE, [r"\bbalanced?\b", r"\bequivalen(t|ce)\b"]),
    _item("16", "Results: numbers analysed", "Numbers analysed per condition and whether by original assignment", [P_ANALYSED, r"\bintention[- ]to[- ]treat\b|\bper[- ]protocol\b"]),
    _item("17", "Results: outcomes", "Summary of results per condition with effect size and confidence interval", P_ESTIMATE),
    _item("18", "Results: ancillary analyses", "Subgroup and adjusted analyses, pre-specified versus exploratory", P_SUBGROUP),
    _item("19", "Results: adverse events", "Important adverse events or unintended effects", P_HARMS),
    _item("20", "Discussion: interpretation", "Interpretation in light of hypotheses, sources of bias, multiplicity and limitations", [P_INTERPRET, P_LIMITS]),
    _item("21", "Discussion: generalisability", "Generalisability of the findings", P_GENERAL),
    _item("22", "Discussion: overall evidence", "General interpretation in the context of current evidence", P_EVIDENCE, [r"\bprior (studies|evidence|research)\b", r"\bprevious (studies|research|evaluations)\b"]),
    _item("23", "Other information", "Funding source", P_FUNDING),
]

STROBE = [
    _item("1a", "Title and abstract", "Study design indicated in the title or abstract", P_TITLE_OBS, scope="head"),
    _item("1b", "Title and abstract", "Informative and balanced summary of what was done and found", P_ABSTRACT, scope="head"),
    _item("2", "Introduction", "Scientific background and rationale", P_BACKGROUND, [r"\bbackground\b", r"\bintroduction\b"]),
    _item("3", "Introduction", "Specific objectives and any pre-specified hypotheses", P_OBJECTIVES),
    _item("4", "Methods: study design", "Key elements of the study design", P_OBS_DESIGN),
    _item("5", "Methods: setting", "Setting, locations and relevant dates (recruitment, exposure, follow-up, data collection)", [P_SETTING, P_DATES], P_SETTING_WEAK),
    _item("6", "Methods: participants", "Eligibility criteria, sources and methods of selection or sampling of participants", [P_ELIGIBLE, r"\b(sampling (frame|strategy|design)|randomly (selected|sampled)|stratified (random )?sampl|systematic sampl|multi[- ]stage|probability sampl|purposive)\b"], [r"\bsampl(e|ed|ing)\b"]),
    _item("7", "Methods: variables", "Definition of outcomes, exposures, predictors, confounders and effect modifiers", P_VARIABLES, [r"\bconfounders?\b", r"\bcovariates?\b", r"\bexposure\b"]),
    _item("8", "Methods: data sources", "Sources of data and details of measurement methods for each variable", P_SOURCES),
    _item("9", "Methods: bias", "Efforts to address potential sources of bias", P_BIAS),
    _item("10", "Methods: study size", "How the study size was arrived at", P_STUDY_SIZE),
    _item("11", "Methods: quantitative variables", "How quantitative variables were handled (groupings, transformations)", P_QUANT),
    _item("12", "Methods: statistical methods", "Statistical methods including confounding control, subgroups, missing data and weighting", [P_STATS, r"\b(survey )?weights?\b|\bweighted\b"], P_STATS_WEAK + [r"\bmissing data\b"]),
    _item("13", "Results: participants", "Numbers at each stage (eligible, included, analysed) and reasons for non-participation", [P_NUMBERS, r"\bresponse rate\b", P_ATTRITION]),
    _item("14", "Results: descriptive data", "Characteristics of participants and missing data per variable", P_DESCRIPTIVE),
    _item("15", "Results: outcome data", "Numbers of outcome events or summary measures", P_OUTCOME_DATA),
    _item("16", "Results: main results", "Unadjusted and adjusted estimates with precision (95% CI) and which confounders were adjusted for", P_MAIN),
    _item("17", "Results: other analyses", "Other analyses: subgroups, interactions, sensitivity analyses", P_SUBGROUP),
    _item("18", "Discussion: key results", "Key results summarised with reference to the objectives", P_KEY),
    _item("19", "Discussion: limitations", "Limitations, sources of bias or imprecision", P_LIMITS),
    _item("20", "Discussion: interpretation", "Cautious overall interpretation considering objectives, limitations, multiplicity and other evidence", P_INTERPRET_OBS),
    _item("21", "Discussion: generalisability", "Generalisability (external validity) of the results", P_GENERAL),
    _item("22", "Other information", "Source of funding and role of funders", P_FUNDING),
]

CHECKLISTS = {
    "CONSORT-SPI": {"items": CONSORT_SPI, "version": "CONSORT 2010 extension for social and psychological interventions (CONSORT-SPI 2018)",
                    "design": "randomised trial"},
    "TREND": {"items": TREND, "version": "TREND statement 2004 (non-randomised evaluations of behavioural and public health interventions)",
              "design": "non-randomised evaluation"},
    "STROBE": {"items": STROBE, "version": "STROBE statement v4 (2007) for observational studies", "design": "observational study"},
}
DESIGN_TO_CHECKLIST = {"randomised": "CONSORT-SPI", "randomized": "CONSORT-SPI", "non-randomised": "TREND", "non-randomized": "TREND",
                       "quasi-experimental": "TREND", "observational": "STROBE"}
QUASI_WORDS = r"\b(non[- ]?randomi[sz]ed|quasi[- ]experimental|difference[- ]in[- ]differences|propensity|matched comparison|regression discontinuity|interrupted time series|natural experiment|synthetic control|phased roll[- ]?out)\b"
RANDOM_WORDS = r"\b(randomi[sz]ed|randomly (assigned|allocated)|random (assignment|allocation)|RCT|cluster[- ]randomi[sz]ed|lottery)\b"
OBS_WORDS = r"\b(cross[- ]sectional|cohort study|case[- ]control|observational|survey|longitudinal|panel (survey|data)|census)\b"

_SCHEMA = {"type": "object", "properties": {
    "document_id": {"type": "string", "minLength": 1, "maxLength": 80, "description": "a registered document (read_document first); its extracted text is checked"},
    "text": {"type": "string", "minLength": 1, "maxLength": 2000000, "description": "report text to check instead of document_id"},
    "checklist": {"type": "string", "enum": ["CONSORT-SPI", "TREND", "STROBE", "auto"], "description": "auto picks by design words in the text (default auto)"},
    "design": {"type": "string", "enum": sorted(DESIGN_TO_CHECKLIST), "description": "override the detected design: randomised → CONSORT-SPI, non-randomised/quasi-experimental → TREND, observational → STROBE"}},
    "required": [], "additionalProperties": False}


def _sentences(text: str) -> List[Tuple[int, int, str]]:
    """(start, end, sentence) triples. Splits on . ! ? followed by whitespace (so 6.1 or 95% CI 2.3 stay whole) and on line breaks."""
    out = []
    start = 0
    for m in re.finditer(r"[.!?]+(?=\s|$)|\n+", text):
        _add_sentence(out, text, start, m.end() if m.group()[0] in ".!?" else m.start())
        start = m.end()
    _add_sentence(out, text, start, len(text))
    return out


def _add_sentence(out: List[Tuple[int, int, str]], text: str, start: int, end: int) -> None:
    chunk = text[start:end]
    stripped = chunk.strip()
    if len(stripped) >= 3:
        lead = len(chunk) - len(chunk.lstrip())
        out.append((start + lead, start + lead + len(stripped), stripped))


def detect_design(text: str) -> Tuple[str, str, Dict[str, int]]:
    counts = {"quasi": len(re.findall(QUASI_WORDS, text, re.I)), "randomised": len(re.findall(RANDOM_WORDS, text, re.I)), "observational": len(re.findall(OBS_WORDS, text, re.I))}
    if counts["quasi"]:
        return "TREND", "non-randomised design words found (e.g. quasi-experimental, difference-in-differences, matched comparison)", counts
    if counts["randomised"]:
        return "CONSORT-SPI", "randomisation words found (randomised, randomly assigned, lottery)", counts
    if counts["observational"]:
        return "STROBE", "observational design words found (cross-sectional, cohort, survey)", counts
    return "STROBE", "no design words found; defaulted to STROBE (observational) — pass checklist or design to override", counts


def _match(item: Dict[str, Any], text: str, sentences: List[Tuple[int, int, str]], title: str) -> Dict[str, Any]:
    evidence: List[Dict[str, Any]] = []
    strong_hits, weak_hits = 0, 0
    seen = set()

    def add(start: int, end: int, quote: str, kind: str) -> None:
        if (start, end) in seen or len(evidence) >= MAX_EVIDENCE:
            return
        seen.add((start, end))
        q = re.sub(r"\s+", " ", quote).strip()
        evidence.append({"quote": q[:MAX_QUOTE] + ("…" if len(q) > MAX_QUOTE else ""), "start": start, "end": end, "match": kind})

    if item["scope"] == "title":
        for rx in item["strong"]:
            m = re.search(rx, title, re.I)
            if m:
                strong_hits += 1
                add(0, len(title), title, "strong")
    elif item["scope"] == "head":
        head = text[:4000]
        for rx in item["strong"]:
            m = re.search(rx, head, re.I | re.M)
            if m:
                strong_hits += 1
                line_end = head.find("\n", m.start())
                add(m.start(), line_end if line_end > 0 else min(len(head), m.end() + 200), head[m.start():line_end if line_end > 0 else m.end() + 200], "strong")
    else:
        for kind, patterns in (("strong", item["strong"]), ("weak", item["weak"])):
            for rx in patterns:
                crx = re.compile(rx, re.I)
                hit = False
                for start, end, s in sentences:
                    if crx.search(s):
                        hit = True
                        add(start, end, s, kind)
                        if len(evidence) >= MAX_EVIDENCE:
                            break
                if hit:
                    if kind == "strong":
                        strong_hits += 1
                    else:
                        weak_hits += 1
    if strong_hits or weak_hits >= 2:
        status = "addressed"
    elif weak_hits == 1:
        status = "partially"
    else:
        status = "missing"
    return {"id": item["id"], "section": item["section"], "item": item["item"], "status": status, "evidence": evidence,
            "strong_matches": strong_hits, "weak_matches": weak_hits}


def _document_text(ctx: JobContext, document_id: str) -> Tuple[str, str]:
    try:
        doc = ctx.project.document(document_id)
    except KeyError as e:
        raise InvalidArguments(str(e)) from e
    if not doc.text_path:
        raise InvalidArguments(f"document {document_id} has no extracted text yet; run read_document on it first")
    path = ensure_within(Path(doc.text_path), ctx.project.root)
    return path.read_text("utf-8"), doc.name


def _write_docx(path: Path, name: str, version: str, title: str, rows: List[Dict[str, Any]], coverage: float, reason: str) -> None:
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.add_heading(f"{name} reporting checklist", level=1)
    p = d.add_paragraph()
    p.add_run(f"Document: {title}\n").bold = True
    p.add_run(f"Checklist: {version}. Items are paraphrased; consult the published statement for the full wording.\n")
    p.add_run(f"Coverage: {coverage:.0f}% of {len(rows)} items addressed. Checklist chosen because {reason}.\n")
    p.add_run("Status is a deterministic text match (addressed / partially / missing) with the quoted sentence that triggered it. "
              "It shows whether the report says something about an item, not whether what it says is adequate. Document text is untrusted data.")
    table = d.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Item", "Section", "What the item asks for", "Status", "Evidence (quoted)"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["id"]
        cells[1].text = row["section"]
        cells[2].text = row["item"]
        cells[3].text = row["status"]
        cells[4].text = "\n".join(f"“{e['quote']}” (chars {e['start']}–{e['end']})" for e in row["evidence"]) or "—"
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)
    missing = [r for r in rows if r["status"] == "missing"]
    d.add_heading("Missing items", level=2)
    if missing:
        for r in missing:
            d.add_paragraph(f"{r['id']} — {r['item']}", style="List Bullet")
    else:
        d.add_paragraph("Every item has at least partial coverage.")
    d.save(str(path))


@tool("reporting_checklist",
      "Check a trial or evaluation report against a reporting guideline: CONSORT-SPI (randomised social/psychological intervention trials), "
      "TREND (non-randomised evaluations) or STROBE (observational studies); 'auto' picks by design words in the text. Deterministic sentence matching "
      "marks each item addressed, partially or missing with quoted evidence and character offsets, computes coverage, lists missing items in plain words "
      "and writes a Word checklist. It judges presence of reporting language, not scientific adequacy. Give document_id (after read_document) or text.",
      _SCHEMA)
def reporting_checklist(ctx: JobContext, args: Dict[str, Any]) -> Dict[str, Any]:
    t0 = time.time()
    if bool(args.get("document_id")) == bool(args.get("text")):
        raise InvalidArguments("give exactly one of document_id or text")
    if args.get("document_id"):
        text, title = _document_text(ctx, args["document_id"])
    else:
        text, title = args["text"], "pasted text"
    if not text.strip():
        raise InvalidArguments("the document has no text to check")
    ctx.progress(10, "choosing the checklist")
    notes: List[str] = []
    wanted = args.get("checklist") or "auto"
    counts: Dict[str, int] = {}
    if args.get("design"):
        name = DESIGN_TO_CHECKLIST[args["design"]]
        reason = f"design override '{args['design']}'"
    elif wanted != "auto":
        name = wanted
        reason = "requested explicitly"
    else:
        name, reason, counts = detect_design(text)
    spec = CHECKLISTS[name]
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")[:300]
    sentences = _sentences(text)
    ctx.progress(30, "matching items")
    rows = [_match(item, text, sentences, first_line) for item in spec["items"]]
    for r in rows:
        ctx.check()
    addressed = [r for r in rows if r["status"] == "addressed"]
    partial = [r for r in rows if r["status"] == "partially"]
    missing = [r for r in rows if r["status"] == "missing"]
    coverage = round(100.0 * len(addressed) / len(rows), 1)
    ctx.progress(70, "writing the Word checklist")
    out = ctx.project.new_artifact_dir("reporting-checklist", name)
    docx_path = out / f"{name}-checklist.docx"
    _write_docx(docx_path, name, spec["version"], title, rows, coverage, reason)
    ctx.project.write_json(out / "checklist.json", {"checklist": name, "version": spec["version"], "document": title, "items": rows})
    notes.append(f"Checklist: {spec['version']}. Item wording is paraphrased, not the published text; cite the statement itself in a report.")
    notes.append("Status reflects deterministic keyword/regex matching over sentences; 'addressed' means the report contains language on the item, not that it is adequate. Review the quotes.")
    if counts:
        notes.append(f"Design words counted: {counts}.")
    missing_items = [f"{r['id']}: {r['item']}" for r in missing]
    summary = (f"{name} checklist on '{title}': {len(addressed)} of {len(rows)} items addressed ({coverage:.0f}%), {len(partial)} partially, {len(missing)} missing."
               + (f" Missing: {', '.join(r['id'] for r in missing[:12])}{'…' if len(missing) > 12 else ''}." if missing else " Every item has language in the report."))
    result = {"checklist": name, "version": spec["version"], "design_assumed": spec["design"], "checklist_reason": reason, "document_id": args.get("document_id"),
              "document_name": title, "chars": len(text), "sentences": len(sentences), "table": rows, "coverage_pct": coverage,
              "addressed": len(addressed), "partially": len(partial), "missing": len(missing), "missing_items": missing_items,
              "partial_items": [f"{r['id']}: {r['item']}" for r in partial],
              "docx_path": str(docx_path), "json_path": str(out / "checklist.json"), "artifact_dir": str(out), "notes": notes,
              "summary_text": summary, "content_trust": CONTENT_TRUST, "elapsed_s": round(time.time() - t0, 2)}
    json.dumps(result)  # the result must stay JSON-serialisable
    ctx.project.audit({"tool": "reporting_checklist", "checklist": name, "document": args.get("document_id"), "artifact": str(out)})
    ctx.progress(100, "done")
    return result
