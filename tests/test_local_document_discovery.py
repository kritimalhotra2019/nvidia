"""Real host/worker with controlled local models and isolated directory grants."""
from pathlib import Path
import os,json,subprocess,threading,sqlite3,hashlib
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pptx import Presentation
import matplotlib
matplotlib.use("Agg")
from matplotlib.backends.backend_pdf import PdfPages
import matplotlib.pyplot as plt
import argparse
p=argparse.ArgumentParser();p.add_argument('--binary',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True)
p.add_argument('--cases',nargs='+',choices=['folder_search','exact_path','new_arrival','missing'],default=['folder_search','exact_path','new_arrival','missing'])
p.add_argument('--reject-redundant-notes',action='store_true',help='Return truncated notes if the host redundantly re-reads a fully delivered fresh folder-search result. Recovery cases still generate verified notes.')
p.add_argument('--short-identity',action='store_true',help='Use a short personality.md only inside each isolated fixture, keeping the normal prompt below the recovery-compaction threshold.')
p.add_argument('--require-notes',action='store_true',help='Require a successful notes pass for the fresh folder-search case (default long context).')
a=p.parse_args()
if a.reject_redundant_notes and not a.short_identity:p.error('--reject-redundant-notes requires --short-identity; default identity may intentionally trigger context compaction and need verified notes')
if a.reject_redundant_notes and a.require_notes:p.error('Choose either rejected redundant notes or required recovery notes')
root=a.output_dir.resolve();root.mkdir(parents=True,exist_ok=False)
binary=a.binary.resolve()
target=None;original=None;provider='ollama'
requests=[];errors=[];results=[];case='';support=None;old_id='';wrong_sent=False

def steps():
 p=support/'projects/default/research.sqlite3'
 if not p.exists():return []
 with sqlite3.connect(p) as db:return [(name,json.loads(a),json.loads(r),ok) for name,a,r,ok in db.execute('SELECT tool,arguments_json,result_json,ok FROM research_steps ORDER BY sequence')]

class Peer(BaseHTTPRequestHandler):
 def log_message(self,*_):pass
 def do_POST(self):
  global wrong_sent,old_id
  try:
   body=json.loads(self.rfile.read(int(self.headers['Content-Length'])));requests.append({'case':case,'body':body})
   docs=json.loads((support/'projects/default/datasets.json').read_text())['documents']
   if isinstance(docs,dict):docs=list(docs.values())
   old_id=next(d['id'] for d in docs if d['name']=='unrelated.pptx')
   seen=steps();read=[r for t,a,r,ok in seen if t=='read_document' and ok]
   listed=[r for t,a,r,ok in seen if t=='list_local_files' and ok]
   truncated_notes=False
   if any('<source_text>' in m.get('content','') for m in body['messages']) and not body.get('tools'):
    action={'answer':'Fixture source: an evaluation discusses programme implementation, descriptive outcomes and limits of causal attribution. Source coverage is recorded by the host.'}
    truncated_notes=a.reject_redundant_notes and case=='folder_search'
   elif not wrong_sent:
    wrong_sent=True;action={'tool':'read_document','arguments':{'document_id':old_id,'offset':0,'max_chars':8000}}
   elif not listed and case!='exact_path':
    action={'tool':'list_local_files','arguments':{'path':str(target.parent),'query':'definitely-absent-evaluation-abc' if case=='missing' else 'wfp-unique','extensions':['pdf'],'limit':1}}
   elif case=='missing':action={'answer':'The complete filename search found no matching PDF. The requested source is unavailable; I have not read an older file.'}
   elif not read or (case=='exact_path' and not any(a.get('path')==str(target) for t,a,r,ok in seen if t=='read_document' and ok)):
    args={'path':str(target),'max_chars':12000}
    if case=='exact_path':args['document_id']=old_id
    action={'tool':'read_document','arguments':args}
   elif read[-1].get('has_more'):
    r=read[-1];action={'tool':'read_document','arguments':{'document_id':r['document_id'],'snapshot_id':r['snapshot_id'],'offset':r['next_offset'],'max_chars':12000}}
   else:
    if case in ['folder_search','new_arrival']:
     assert any('SYNTHETIC_END_MARKER_731' in m.get('content','') for m in body['messages']),'Final reply did not receive the original source through its last marker'
    action={'answer':'Read the requested WFP PDF with complete extracted-text coverage. The source identity is verified; this fixture does not evaluate summary quality.'}
   text=action.get('answer','');tool=action.get('tool')
   if body.get('format') or (provider=='mlx' and not body.get('stream') and 'answer' in action):text=json.dumps(action);tool=None
   self.send_response(200)
   message={'role':'assistant','content':text}
   if provider=='ollama':
    self.send_header('Content-Type','application/x-ndjson');self.end_headers()
    if tool:message['tool_calls']=[{'function':{'name':tool,'arguments':action['arguments']}}]
    self.wfile.write((json.dumps({'message':message,'done':True,'done_reason':'length' if truncated_notes else 'stop'})+'\n').encode())
   else:
    call={'id':'fixture_'+str(len(requests)),'type':'function','function':{'name':tool,'arguments':json.dumps(action['arguments'])}} if tool else None
    if body.get('stream'):
     self.send_header('Content-Type','text/event-stream');self.end_headers()
     delta={'tool_calls':[{'index':0,**call}]} if call else {'content':text}
     self.wfile.write(('data: '+json.dumps({'choices':[{'delta':delta,'finish_reason':'tool_calls' if call else 'stop'}]})+'\n\ndata: [DONE]\n\n').encode())
    else:
     self.send_header('Content-Type','application/json');self.end_headers()
     if call:message['tool_calls']=[call]
     self.wfile.write(json.dumps({'choices':[{'message':message,'finish_reason':'length' if truncated_notes else ('tool_calls' if call else 'stop')}]}).encode())

  except Exception as e:errors.append(str(e));(root/'peer-errors.json').write_text(json.dumps(errors));self.send_response(500);self.end_headers()

server=ThreadingHTTPServer(('127.0.0.1',0),Peer);threading.Thread(target=server.serve_forever,daemon=True).start()
try:
 for provider,case in [(p,c) for p in ['ollama','mlx'] for c in a.cases]:
  request_start=len(requests)
  wrong_sent=False;out=root/(os.environ.get('DISCOVERY_RUN','native3')+'-'+provider+'-'+case);support=out/'support';support.mkdir(parents=True,exist_ok=False)
  if a.short_identity:
   instructions=support/'Instructions';instructions.mkdir()
   (instructions/'personality.md').write_text('You are Squipi, a research companion. This isolated fixture tests source delivery and host completion, not personality or research reasoning.\n')
  granted=support/'granted';granted.mkdir()
  for i in range(650):(granted/f'A{i:04d}.txt').write_text('fixture')
  target=granted/'WFP-unique-evaluation.pdf'
  with PdfPages(target) as pdf:
   fig=plt.figure(figsize=(8.3,11.7));fig.text(.1,.9,'Synthetic evaluation fixture',fontsize=16);fig.text(.1,.8,'Programme implementation and descriptive outcomes.\nLimit: no causal attribution without a comparison design.\nThis is synthetic test evidence.\nSYNTHETIC_END_MARKER_731',fontsize=11);pdf.savefig(fig);plt.close(fig)
  original=hashlib.sha256(target.read_bytes()).hexdigest()
  (support/'worker-env').symlink_to(Path.home()/'Library/Application Support/Squipi/worker-env',target_is_directory=True)
  old=out/'unrelated.pptx';ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='UNRELATED OLD DOCUMENT';slide.placeholders[1].text='This is not the requested WFP evaluation.';ppt.save(old)
  goal='Find the WFP PDF in my Downloads, read it and summarise what it is about and its limitations.'
  env=os.environ|{'SQUIPI_TEST_DATA_DIR':str(support),'SQUIPI_TEST_PROVIDER':provider,'SQUIPI_TEST_DIRECTORY_GRANT':'1','SQUIPI_TEST_MLX_URL':f'http://127.0.0.1:{server.server_port}/v1','SQUIPI_TEST_OLLAMA_URL':f'http://127.0.0.1:{server.server_port}','PYTHONDONTWRITEBYTECODE':'1'}
  if case=='exact_path':goal='Read and summarise the PDF at '+str(target)
  if case=='missing':goal='Find the definitely-absent-evaluation-abc PDF in my Downloads and read it.'
  if case=='new_arrival':
   history=support/'history.json';history.write_text(json.dumps([{'role':'user','content':goal},{'role':'assistant','content':'The requested file was not present. What is its filename?'}]));env['SQUIPI_AGENT_HISTORY']=str(history);goal="the WFP PDF wasn't in Downloads earlier, but it is now"
  run=subprocess.run([str(binary),'--agent',goal,str(old)],env=env,capture_output=True,text=True,timeout=180)
  (out/'run.log').write_text(run.stdout+run.stderr);assert run.returncode==0,(case,run.stderr)
  reply=json.loads((support/'last-agent-result.json').read_text());seen=steps()
  assert not any(t=='read_document' and ok and r.get('document_id')==old_id for t,a,r,ok in seen),(case,'old source read',reply)
  assert any(t=='read_document' and not ok and 'outside this request' in str(r) for t,a,r,ok in seen),(case,'stale ID not blocked',reply)
  if case=='missing':assert reply['status']=='partial' and not any(t=='read_document' and ok for t,a,r,ok in seen),reply
  else:
   assert reply['status']=='completed',(case,reply)
   successful=[r for t,a,r,ok in seen if t=='read_document' and ok]
   assert successful and successful[-1].get('complete') is True,(case,successful[-1].keys())
   index=json.loads((support/'projects/default/datasets.json').read_text())['documents']
   paths={d['id']:d['path'] for d in index}
   assert all(paths[r['document_id']]==str(target) for r in successful),(case,'wrong path',successful[-1])
   if case!='exact_path':assert any(t=='list_local_files' and a.get('query')=='wfp-unique' and ok for t,a,r,ok in seen)
   if a.reject_redundant_notes and case=='folder_search':
    assert len(successful)==1,(case,'Completed delivered source was redundantly read',successful)
    assert not any(t=='read_document' and not ok and 'did not finish reading' in str(r) for t,a,r,ok in seen),(case,'Redundant notes failure poisoned source completion')
   if a.require_notes and case=='folder_search':
    assert any(r.get('reading_notes') for r in successful),(case,'Compacted recovery did not retain verified source notes')
  assert original==hashlib.sha256(target.read_bytes()).hexdigest()
  notes_calls=sum(any('<source_text>' in m.get('content','') for m in request['body']['messages']) for request in requests[request_start:])
  results.append({'provider':provider,'case':case,'passed':True,'status':reply['status'],'stale_document_blocked':True,'successful_document_calls':sum(t=='read_document' and ok for t,a,r,ok in seen),'notes_model_calls':notes_calls,'identity':'isolated-short' if a.short_identity else 'bundled-default'})
 assert not errors,errors
 assert original==hashlib.sha256(target.read_bytes()).hexdigest()
 (root/'native-verification.json').write_text(json.dumps({'scripted_model':True,'synthetic_sources_unchanged':True,'cases':results},indent=2));print(json.dumps(results,indent=2))
finally:
 (root/'native-requests.json').write_text(json.dumps(requests,indent=2));server.shutdown();server.server_close()
