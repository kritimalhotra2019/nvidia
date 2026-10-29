#!/usr/bin/env python3
"""Test real presentation execution with a synthetic template and loopback peer.

The native Swift Agent invokes the packaged Python worker; only inference is
scripted. One case rejects an omitted body fact and repairs the same draft. A
second returns malformed reviewer passes and must retain a partial task after
three bounded attempts. No live model, cloud, private template or real app state
is used. The driver needs only Python's standard library and an existing worker
Python environment with python-pptx installed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import signal
import sqlite3
import subprocess
import tempfile
import threading
import traceback
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET
import zipfile

from test_ollama_agent_integration import read_steps


TITLE = "Synthetic household profile"
BODY = "500 households; 25 missing income values (5%)."
MEDIAN = "Median monthly household income: INR 44,300."
LIMITS = "Income values remain missing. Descriptive results do not establish causality."
BRAND = "SYNTHETIC RESEARCH LAB"
OLD_NOTES = "SYNTHETIC_OLD_NOTES_MUST_NOT_SURVIVE"
GOAL = (f"Create exactly three PowerPoint slides titled {TITLE} using the saved default template. "
        "Use only these supplied synthetic figures: 500 households; 25 missing income values (5%); "
        "median monthly household income INR 44,300. Slide 1: title and synthetic-data notice. "
        "Slide 2: all three figures in its main body. Slide 3: explain that income values remain "
        "missing and descriptive results do not establish causality. Do not perform further "
        "analyses or ask questions.")

# Create three generous, clearly named slots per slide, fixed artwork and notes.
# The seed process uses the same worker package/interpreter as the native app.
SEED = r'''
import hashlib, json, struct, sys, zlib
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from squipi_worker.project import Project
from squipi_worker.security import AllowList
from squipi_worker.templates import library_action
output, brand, old_notes = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
project = Project(output / 'app-support/projects/default', AllowList())
def chunk(kind, data):
    return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)
logo = output / 'synthetic-logo.png'
pixels = b''.join(b'\x00' + bytes((35, 87, 170)) * 24 for _ in range(24))
logo.write_bytes(b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!2I5B', 24, 24, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(pixels)) + chunk(b'IEND', b''))
source = output / 'simple-research-template.pptx'
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
for number in range(1, 4):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    panel = slide.shapes.add_shape(1, Inches(.3), Inches(.3), Inches(9.4), Inches(.12))
    panel.name = 'fixed-decoration'
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor.from_string('2357AA')
    slide.shapes.add_picture(str(logo), Inches(9), Inches(.55), width=Inches(.35))
    for name, text, y, height, size in [
        ('content-title', '[Title]', .7, .7, 26),
        ('content-body', '[Body]', 1.8, 3.9, 20),
        ('brand-footer', brand, 6.8, .3, 10),
    ]:
        shape = slide.shapes.add_textbox(Inches(.7), Inches(y), Inches(8.1), Inches(height))
        shape.name, shape.text = name, text
        shape.text_frame.word_wrap = True
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name, run.font.size = 'Arial', Pt(size)
                run.font.color.rgb = RGBColor.from_string('14233F')
    slide.notes_slide.notes_text_frame.text = old_notes
deck.save(source)
project.allow.allow(source)
library = library_action(project, {'action':'import', 'path':str(source), 'name':'Simple Synthetic Research', 'make_default':True})
entry = library['templates'][-1]
print(json.dumps({'template_id':entry['id'], 'source':str(source), 'sha256':hashlib.sha256(source.read_bytes()).hexdigest()}))
'''


class PresentationFixture:
    def __init__(self, output: Path, mode: str, seed: dict):
        self.output, self.mode, self.seed = output, mode, seed
        self.project = output / "app-support/projects/default"
        self.requests: list[dict] = []
        self.errors: list[str] = []
        self.review_calls = 0
        self.action_calls = 0
        self.injected_error = False
        self.observed_inspection_gate = False
        self.answer_count = 0
        self.bypass_attempts = 0

    def review(self, request: dict) -> str:
        self.review_calls += 1
        assert request.get("think") is False and request["options"]["temperature"] == 0
        assert "tools" not in request, "Review must use an independent context without tools"
        assert len(request["messages"]) == 2, "Review must not inherit the actor conversation"
        evidence = json.loads(request["messages"][1]["content"].split("Candidate slide evidence:\n", 1)[1])
        slides = evidence["slides"]
        assert [slide["slide_number"] for slide in slides] == [1, 2, 3]
        body = "\n".join(item.get("text", "") for item in slides[1]["text_elements"] if item.get("region") == "body")
        assert BODY in body, "Reviewer did not receive the actual exported body evidence"
        if self.mode == "review-reject":
            # A string is not a JSON boolean. A superficially positive reviewer
            # response must fail strict parsing and never unlock completion.
            return json.dumps({"passed": "true", "reviewed_slides": [1, 2, 3], "issues": []})
        passed = MEDIAN in body
        assert passed == (self.review_calls == 2), "Expected one actual omission followed by same-draft repair"
        return json.dumps({"passed": passed, "reviewed_slides": [1, 2, 3], "issues": [] if passed else [
            "Slide 2 is missing median monthly household income INR 44,300 from its main body. Fill its content-body text field."
        ]})

    def action(self, request: dict) -> str:
        self.requests.append(request)
        if any("PRESENTATION_CONTENT_REVIEW" in message.get("content", "") for message in request["messages"]):
            return self.review(request)
        self.action_calls += 1
        assert self.action_calls <= 20, "The controlled workflow repeated rather than finishing"
        assert "format" in request, "Saved-default PPTX requests must begin and remain in structured action mode"
        assert "tools" not in request and request.get("think") is False
        assert request["options"]["temperature"] == 0
        variants = request["format"]["oneOf"]
        names = variants[0]["properties"]["tool"]["enum"]
        assert "export_document" not in names, "The actual-template route must not fall back to a generic theme export"
        messages = "\n".join(message.get("content", "") for message in request["messages"])
        assert GOAL in messages
        observed = read_steps(self.project)
        good = [step for step in observed if step["ok"]]
        reviews = [step for step in observed if step["tool"] == "presentation_content_review"]
        passed = bool(reviews and reviews[-1]["ok"])
        assert (len(variants) == 3) == passed, "A missing artifact or failed reviewer unlocked the final answer"
        begins = [step for step in good if step["tool"] == "begin_template_presentation"]
        if not begins:
            assert "begin_template_presentation" in names
            assert not {"inspect_presentation_draft", "edit_presentation_slide", "finalize_template_presentation"}.intersection(names)
            # report_dataset remains useful for analysis, but its generic PPTX
            # formats must not bypass the required actual-template workflow.
            bypasses = [step for step in observed if step["tool"] == "report_dataset"]
            for step in bypasses:
                assert not step["ok"] and "actual slides" in step["result"].get("error", ""), "Template bypass reached the worker instead of the host guard"
            if self.mode == "review-reject" and self.bypass_attempts < 2:
                if "report_dataset" not in names:
                    assert "discover_tools" in names, "The report capability must be available for discovery"
                    return self.tool("discover_tools", names=["report_dataset"])
                fmt = ["pptx", "all"][self.bypass_attempts]
                self.bypass_attempts += 1
                return self.tool("report_dataset", dataset_id="synthetic_missing_dataset", format=fmt, title=TITLE)
            assert len(bypasses) == (2 if self.mode == "review-reject" else 0)
            return self.tool("begin_template_presentation", template_id=self.seed["template_id"], slide_numbers=[1, 2, 3])
        assert len(begins) == 1, "A repair must keep the original draft"
        states = [step["result"] for step in good if step["tool"] in {
            "begin_template_presentation", "inspect_presentation_draft", "edit_presentation_slide"}]
        state = states[-1]
        base = {"draft_id": state["draft_id"], "revision": state["revision"]}
        inspections = [step["result"] for step in good if step["tool"] == "inspect_presentation_draft"]
        failures = [step for step in observed if step["tool"] == "edit_presentation_slide" and not step["ok"]]
        if failures and observed[-1] == failures[-1]:
            assert "edit_presentation_slide" not in names, "An invalid edit must force inspection before retry"
            self.observed_inspection_gate = True
            return self.tool("inspect_presentation_draft", **base, slide_number=failures[-1]["arguments"]["slide_number"])
        remaining = state["completion"]["slides_needing_edits"]
        target = remaining[0] if remaining else None
        if target is None and self.mode == "review-repair" and reviews and not passed:
            body_edits = [edit["text"] for step in good if step["tool"] == "edit_presentation_slide" and step["arguments"]["slide_number"] == 2
                          for edit in step["arguments"].get("text_edits", [])]
            if not any(MEDIAN in text for text in body_edits):
                target = 2
                assert "separate content review" in messages, "Reviewer failure did not reach the repair actor"
        if target:
            inspection = inspections[-1] if inspections else {}
            if inspection.get("slide_number") != target or inspection.get("revision") != state["revision"]:
                return self.tool("inspect_presentation_draft", **base, slide_number=target)
            if not self.injected_error:
                self.injected_error = True
                return self.tool("edit_presentation_slide", **base, slide_number=target,
                                 text_edits=[{"shape_id": 999999, "text": "Invalid target must not advance revision"}])
            slots = {shape["name"]: shape for shape in inspection["shapes"]}
            for name in ("content-title", "content-body"):
                assert slots[name]["region"] == "body" and slots[name]["text_editable"] is True
            titles = {1: TITLE, 2: "Household descriptive figures", 3: "Interpretation and limitations"}
            contents = {1: "Synthetic supplied data for a presentation test.", 2: BODY, 3: LIMITS}
            if self.mode == "review-reject" or reviews:
                contents[2] += "\n" + MEDIAN
            return self.tool("edit_presentation_slide", **base, slide_number=target, text_edits=[
                {"shape_id": slots["content-title"]["shape_id"], "text": titles[target]},
                {"shape_id": slots["content-body"]["shape_id"], "text": contents[target]},
            ])
        exports = [step for step in good if step["tool"] == "finalize_template_presentation"]
        if not exports or exports[-1]["result"]["revision"] != state["revision"] or not passed:
            assert state["completion"]["ready_to_finalize"]
            return self.tool("finalize_template_presentation", **base, title=TITLE)
        path = Path(exports[-1]["result"]["path"])
        assert path.is_file() and str(path) in messages.replace("\\/", "/")
        self.answer_count += 1
        return json.dumps({"answer": f"Created the presentation: {path}"})

    @staticmethod
    def tool(name: str, **arguments) -> str:
        return json.dumps({"tool": name, "arguments": arguments})


def handler_for(fixture: PresentationFixture):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *_):
            pass

        def send_json(self, value: dict, status: int = 200):
            data = json.dumps(value).encode()
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            path = urlsplit(self.path).path
            if path == "/api/ps":
                self.send_json({"models": []})
            elif path == "/api/tags":
                self.send_json({"models": [{"name": "fixture-ollama", "model": "fixture-ollama"}]})
            else:
                fixture.errors.append(f"Unexpected GET {path}")
                self.send_json({"error": "Unexpected fixture request"}, 404)

        def do_POST(self):
            try:
                path = urlsplit(self.path).path
                request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                if path == "/api/show":
                    self.send_json({"capabilities": ["completion", "tools"], "model_info": {"fixture.context_length": 32768}})
                    return
                assert path == "/api/chat" and request["model"] == "fixture-ollama"
                answer = fixture.action(request)
                message = {"message": {"role": "assistant", "content": answer}, "done": True,
                           "done_reason": "stop", "prompt_eval_count": 100, "eval_count": 50}
                self.send_json(message)
            except Exception as error:
                fixture.errors.append(traceback.format_exc())
                self.send_json({"error": str(error)}, 500)
    return Handler


def verify(fixture: PresentationFixture) -> dict:
    assert not fixture.errors, "; ".join(fixture.errors)
    steps = read_steps(fixture.project)
    begins = [step for step in steps if step["tool"] == "begin_template_presentation"]
    exports = [step for step in steps if step["tool"] == "finalize_template_presentation"]
    reviews = [step for step in steps if step["tool"] == "presentation_content_review"]
    assert len(begins) == 1 and all(step["ok"] for step in exports)
    assert len([step for step in steps if step["tool"] == "edit_presentation_slide" and not step["ok"]]) == 1
    assert fixture.injected_error and fixture.observed_inspection_gate
    assert fixture.bypass_attempts == (2 if fixture.mode == "review-reject" else 0)
    bypasses = [step for step in steps if step["tool"] == "report_dataset"]
    assert len(bypasses) == fixture.bypass_attempts and not any(step["ok"] for step in bypasses)
    assert {step["result"]["draft_id"] for step in exports} == {begins[0]["result"]["draft_id"]}
    expected_reviews = [False, True] if fixture.mode == "review-repair" else [False] * 3
    assert [bool(step["ok"]) for step in reviews] == expected_reviews
    assert len(exports) == len(reviews) == fixture.review_calls
    final = exports[-1]["result"]
    path, source = Path(final["path"]), Path(fixture.seed["source"])
    assert path.is_file() and path.stat().st_size > 0 and final["completion"]["ready_to_finalize"]
    assert final["template_id"] == fixture.seed["template_id"] and final["visual_review_required"] is True
    assert hashlib.sha256(source.read_bytes()).hexdigest() == fixture.seed["sha256"]
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path) as finished:
        assert finished.testzip() is None
        members = finished.namelist()
        assert not any(OLD_NOTES.encode() in finished.read(name) for name in members if name.endswith(".xml"))
        # Master/layout/theme and embedded artwork must retain source bytes;
        # slide XML must legitimately change to contain the new report.
        preserved = [name for name in original.namelist() if name.startswith(("ppt/media/", "ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/"))]
        assert any(name.startswith("ppt/media/") for name in preserved)
        for name in preserved:
            assert finished.read(name) == original.read(name), f"Template artwork or formatting changed: {name}"
        slide_names = [f"ppt/slides/slide{number}.xml" for number in range(1, 4)]
        assert len([name for name in members if name.startswith("ppt/slides/slide") and name.endswith(".xml")]) == 3
        for number, name in enumerate(slide_names, 1):
            before, after = ET.fromstring(original.read(name)), ET.fromstring(finished.read(name))
            shapes_before = before.find("p:cSld/p:spTree", ns)
            shapes_after = after.find("p:cSld/p:spTree", ns)
            assert len(shapes_before) == len(shapes_after)
            for left, right in zip(shapes_before, shapes_after):
                for query in ("p:spPr/a:xfrm", "p:spPr/a:solidFill"):
                    a, b = left.find(query, ns), right.find(query, ns)
                    assert (ET.tostring(a) if a is not None else None) == (ET.tostring(b) if b is not None else None)
            texts = {}
            for shape in after.findall("p:cSld/p:spTree/p:sp", ns):
                name_element = shape.find("p:nvSpPr/p:cNvPr", ns)
                texts[name_element.get("name")] = "\n".join(element.text or "" for element in shape.findall(".//a:t", ns))
            assert texts["brand-footer"] == BRAND
            assert not any("[" in text or "]" in text for text in texts.values())
            if number == 1:
                assert texts["content-title"] == TITLE and "Synthetic" in texts["content-body"]
            elif number == 2:
                assert BODY in texts["content-body"] and MEDIAN in texts["content-body"]
                assert "44,300" not in texts["brand-footer"]
            else:
                assert texts["content-body"] == LIMITS
    reply = json.loads((fixture.output / "app-support/last-agent-result.json").read_text())
    expected_status = "completed" if fixture.mode == "review-repair" else "partial"
    assert reply["status"] == expected_status, reply
    assert fixture.answer_count == int(fixture.mode == "review-repair")
    if fixture.mode == "review-repair":
        assert "visual inspection" in reply["text"], "Text review must not imply visual verification"
        assert exports[-1]["result"]["revision"] > exports[0]["result"]["revision"]
    else:
        assert "three bounded repair attempts" in reply["error"]
    with sqlite3.connect(f"{(fixture.project / 'research.sqlite3').as_uri()}?mode=ro", uri=True) as connection:
        tasks = connection.execute("SELECT status FROM research_tasks").fetchall()
    assert tasks == [(expected_status,)]
    return {"case": fixture.mode, "status": "passed", "task_status": expected_status,
            "model_requests": len(fixture.requests), "review_calls": fixture.review_calls,
            "same_draft": True, "inspection_after_edit_error": True, "answer_gated_on_review": True,
            "generic_report_bypasses_rejected": [step["arguments"]["format"] for step in bypasses],
            "body_facts_verified": True, "original_unchanged": True, "artwork_bytes_preserved": True,
            "presentation": str(path)}


def run_case(binary: Path, worker_python: Path, worker_source: Path, output: Path, mode: str, timeout: int) -> dict:
    output.mkdir(parents=True, exist_ok=False)
    support = output / "app-support"
    support.mkdir()
    (support / "worker-env").symlink_to(worker_python.parent.parent, target_is_directory=True)
    environment = {name: os.environ[name] for name in ["PATH", "HOME", "USER", "TMPDIR", "LANG", "LC_ALL", "TZ"] if name in os.environ}
    seed_env = {**environment, "PYTHONPATH": str(worker_source), "PYTHONDONTWRITEBYTECODE": "1"}
    seeded = subprocess.run([str(worker_python), "-c", SEED, str(output), BRAND, OLD_NOTES], env=seed_env,
                            text=True, capture_output=True, timeout=60, check=False)
    assert seeded.returncode == 0, "Could not seed synthetic template: " + seeded.stderr
    seed = json.loads(seeded.stdout)
    (output / "seed.json").write_text(json.dumps(seed, indent=2) + "\n")
    fixture = PresentationFixture(output, mode, seed)
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(fixture))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    environment.update({"SQUIPI_TEST_DATA_DIR": str(support), "SQUIPI_TEST_PROVIDER": "ollama",
                        "SQUIPI_TEST_OLLAMA_URL": f"http://127.0.0.1:{server.server_address[1]}",
                        "SQUIPI_OLLAMA_MODEL": "fixture-ollama", "PYTHONDONTWRITEBYTECODE": "1"})
    try:
        process = subprocess.Popen([str(binary), "--agent", GOAL], env=environment,
                                   text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True)
        timed_out = False
        try:
            stdout, stderr = process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            timed_out = True
            os.killpg(process.pid, signal.SIGTERM)
            try:
                stdout, stderr = process.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
                stdout, stderr = process.communicate(timeout=5)
        (output / "stdout.log").write_text(stdout)
        (output / "stderr.log").write_text(stderr)
        assert not timed_out, f"Native regression exceeded {timeout}s; its isolated process group was stopped"
        assert process.returncode == 0, f"App exited {process.returncode}; inspect {output}/stderr.log"
        result = verify(fixture)
        (output / "results.json").write_text(json.dumps(result, indent=2) + "\n")
        return result
    finally:
        server.shutdown()
        server.server_close()
        (output / "model-requests.json").write_text(json.dumps(fixture.requests, indent=2) + "\n")
        if fixture.errors:
            (output / "fixture-errors.json").write_text(json.dumps(fixture.errors, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--binary", required=True, type=Path)
    parser.add_argument("--worker-python", type=Path, default=Path.home() / "Library/Application Support/Squipi/worker-env/bin/python")
    parser.add_argument("--worker-source", type=Path, help="Defaults to the tested app's bundled worker")
    parser.add_argument("--output-dir", type=Path, help="A new directory for isolated state and retained evidence")
    parser.add_argument("--timeout", type=int, default=150, help="Bound per native app process, in seconds")
    parser.add_argument("--case", choices=["all", "review-repair", "review-reject"], default="all")
    args = parser.parse_args()
    binary = args.binary.resolve()
    worker_python = Path(os.path.abspath(args.worker_python))
    worker_source = args.worker_source.resolve() if args.worker_source else binary.parent.parent / "Resources/worker"
    assert binary.is_file() and worker_python.is_file() and os.access(worker_python, os.X_OK)
    assert (worker_source / "squipi_worker/tools/presentation_edit.py").is_file(), "Use a build with template editing support"
    output = args.output_dir.resolve() if args.output_dir else Path(tempfile.mkdtemp(prefix="squipi-presentation-integration-"))
    if args.output_dir:
        output.mkdir(parents=True, exist_ok=False)
    modes = ["review-repair", "review-reject"] if args.case == "all" else [args.case]
    results = [run_case(binary, worker_python, worker_source, output / mode, mode, args.timeout) for mode in modes]
    (output / "results.json").write_text(json.dumps(results, indent=2) + "\n")
    print(json.dumps(results, indent=2))
    print(f"PASS: native presentation editing, same-draft review repair and bounded reviewer rejection; evidence: {output}")


if __name__ == "__main__":
    main()
