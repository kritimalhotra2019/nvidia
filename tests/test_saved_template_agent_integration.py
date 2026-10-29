#!/usr/bin/env python3
"""Check saved templates through the real Swift Agent and real export worker.

A controlled loopback Ollama fixture stalls once, then returns three structured
decisions: Word export, PowerPoint export, final answer. This is a context and
execution integration test, not an inference-quality test. All documents, app
state, defaults and outputs are synthetic and isolated; no keys or live model.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import tempfile
import threading
import xml.etree.ElementTree as ET
import zipfile
from http.server import ThreadingHTTPServer

from test_ollama_agent_integration import handler_for, read_steps


GOAL = "Create a Word report and PowerPoint presentation from these supplied synthetic household figures: 40 households, mean income 2400 units per month, mean wellbeing 6.2 out of 10."
STALL = "I will prepare those files. [Status: Awaiting response from host]"
OLD_BODY = "PRIVATE_OLD_TEMPLATE_BODY_MUST_NOT_BECOME_RESEARCH"
OLD_NOTES = "PRIVATE_OLD_SLIDE_NOTES_MUST_NOT_BE_COPIED"
CHANGED = "MID_TURN_LIBRARY_EDIT_FOR_NEXT_REQUEST_ONLY"
START = "USER-SAVED LOCAL OUTPUT TEMPLATES\n"
END = "END USER-SAVED LOCAL OUTPUT TEMPLATES"
SELECTIONS = {
    "default-pinned": "Use my saved default templates.",
    "named-alternate": "Use Alternate Word Format for the Word report and Alternate Slide Format for the PowerPoint presentation.",
    "no-template": "Use no saved template for either file. Use standard formatting for both outputs.",
}

# Executed by the existing worker interpreter against the exact packaged worker
# under test. The driver therefore needs only Python's standard library.
SEED = r'''
import hashlib, json, sys, zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PColor
from lxml import etree
from squipi_worker.project import Project
from squipi_worker.security import AllowList
from squipi_worker.templates import library_action

output = Path(sys.argv[1])
old_body, old_notes = sys.argv[2:4]
project = Project(output / 'app-support/projects/default', AllowList())
source_dir = output / 'original-formats'
source_dir.mkdir()
sources = {}
ids = {}
for label, font, points, color, width, height, left in [
    ('Default', 'Georgia', 11, '0F766E', 8.5, 11, 1.1),
    ('Alternate', 'Arial', 13, '8C4A2F', 8, 10, .75),
]:
    word = source_dir / (label + '-original.docx')
    document = Document()
    for style_name in ['Normal', 'Title', 'Heading 1']:
        style = document.styles[style_name]
        style.font.name = font
        style.font.size = Pt(points if style_name == 'Normal' else points + 10)
        style.font.color.rgb = RGBColor.from_string(color)
    section = document.sections[-1]
    section.page_width, section.page_height = Inches(width), Inches(height)
    section.left_margin, section.right_margin = Inches(left), Inches(.65)
    section.header.paragraphs[0].text = label.upper() + '_BRAND_HEADER'
    section.footer.paragraphs[0].text = label.upper() + '_BRAND_FOOTER'
    document.add_heading(old_body + '_' + label, 0)
    document.add_paragraph('The old sample says 999999 households. Do not reuse it.')
    document.add_table(rows=2, cols=2).cell(0, 0).text = old_body + '_TABLE'
    document.save(word)

    deck_path = source_dir / (label + '-original.pptx')
    deck = Presentation()
    deck.slide_width = PInches(10 if label == 'Default' else 12)
    deck.slide_height = PInches(7.5 if label == 'Default' else 6.75)
    master = deck.slide_masters[0]
    master._element.cSld.set('name', label + ' Research Master')
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = PColor.from_string('F7F6F2' if label == 'Default' else 'FFF3E8')
    deck.slide_layouts[1].name = label + ' Research Content Layout'
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = old_body + '_' + label
    slide.placeholders[1].text = old_body + '_SUBTITLE'
    slide.notes_slide.notes_text_frame.text = old_notes
    content = deck.slides.add_slide(deck.slide_layouts[1])
    content.shapes.title.text = old_body + '_SECTION'
    content.placeholders[1].text = old_body + '_PARAGRAPH'
    for slide in deck.slides:
        for shape in slide.placeholders:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = font
                    paragraph.font.size = PPt(29 if shape == slide.shapes.title else 19)
                    paragraph.font.color.rgb = PColor.from_string(color)
    deck.save(deck_path)
    # Make the theme observably different as well as the master/layout.
    with zipfile.ZipFile(deck_path) as package:
        entries = [(item, package.read(item)) for item in package.infolist()]
    with zipfile.ZipFile(deck_path, 'w') as package:
        for item, body in entries:
            if item.filename == 'ppt/theme/theme1.xml':
                tree = etree.fromstring(body)
                ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
                accent = tree.find('a:themeElements/a:clrScheme/a:accent1', ns)
                for child in list(accent): accent.remove(child)
                etree.SubElement(accent, '{' + ns['a'] + '}srgbClr', val=color)
                body = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
            package.writestr(item, body)
    for fmt, source, name in [('docx', word, label + ' Word Format'), ('pptx', deck_path, label + ' Slide Format')]:
        project.allow.allow(source)
        manifest = library_action(project, {'action': 'import', 'path': str(source), 'name': name, 'make_default': label == 'Default'})
        entry = manifest['templates'][-1]
        ids[label.lower() + '_' + fmt] = entry['id']
        sources[entry['id']] = {'path': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'font': font, 'points': points, 'color': color, 'label': label,
            'word_width_twips': round(width * 1440), 'word_height_twips': round(height * 1440),
            'word_left_twips': round(left * 1440),
            'ppt_width': int(deck.slide_width), 'ppt_height': int(deck.slide_height)}
print(json.dumps({'ids': ids, 'sources': sources, 'manifest': manifest}))
'''


class TemplateFixture:
    def __init__(self, output: Path, mode: str, seed: dict):
        self.output = output
        self.project = output / "app-support/projects/default"
        self.mode = mode
        self.seed = seed
        self.requests: list[dict] = []
        self.errors: list[str] = []
        self.normal_calls = 0
        self.action_calls = 0

    def expected_id(self, fmt: str) -> str:
        if self.mode == "no-template":
            return "none"
        prefix = "alternate" if self.mode == "named-alternate" else "default"
        return self.seed["ids"][prefix + "_" + fmt]

    def action(self, request: dict) -> str:
        self.requests.append(request)
        messages = "\n".join(message.get("content", "") for message in request["messages"])
        assert messages.splitlines().count(START.strip()) == 1, "Each native/recovery request needs one template snapshot"
        section = messages.split(START, 1)[1].split(END, 1)[0]
        metadata = json.loads(section.strip().splitlines()[-1])
        assert metadata["defaults"] == self.seed["manifest"]["defaults"], "Defaults changed midway through the request"
        assert {entry["name"] for entry in metadata["templates"]} == {entry["name"] for entry in self.seed["manifest"]["templates"]}
        assert {entry["id"] for entry in metadata["templates"]} == set(self.seed["sources"])
        assert all("filename" not in entry and "path" not in entry for entry in metadata["templates"])
        assert OLD_BODY not in messages and OLD_NOTES not in messages, "Source template content leaked into the model context"
        assert CHANGED not in messages, "Template metadata must be snapshotted once per turn"
        for source in self.seed["sources"].values():
            assert source["path"] not in messages.replace("\\/", "/"), "Source paths must not enter template metadata"
        assert GOAL in messages and SELECTIONS[self.mode] in messages
        if len(self.requests) == 1 and self.mode == "default-pinned":
            # Simulate an edit in the native template UI while inference runs.
            path = self.project / "templates.json"
            edited = json.loads(path.read_text())
            edited["defaults"] = {fmt: self.seed["ids"]["alternate_" + fmt] for fmt in ["docx", "pptx"]}
            edited["templates"][-1]["summary"] += " " + CHANGED
            temporary = path.with_suffix(".next")
            temporary.write_text(json.dumps(edited))
            temporary.replace(path)
        if "format" not in request:
            self.normal_calls += 1
            assert self.normal_calls == 1 and self.action_calls == 0
            names = {tool["function"]["name"] for tool in request.get("tools", [])}
            assert {"export_document", "list_templates"}.issubset(names)
            return STALL
        self.action_calls += 1
        assert self.normal_calls == 1 and self.action_calls <= 3, "Export tools must not repeat after success"
        assert "tools" not in request and request.get("think") is False
        assert STALL not in messages, "Recovery should drop the stale awaiting-host answer"
        names = request["format"]["oneOf"][0]["properties"]["tool"]["enum"]
        assert "export_document" in names
        steps = read_steps(self.project)
        assert all(step["tool"] == "export_document" and step["ok"] for step in steps), steps
        formats = {step["result"]["format"] for step in steps}
        for fmt in ["docx", "pptx"]:
            if fmt not in formats:
                arguments = {"format": fmt, "title": "Synthetic household profile", "subtitle": "Supplied figures for this template test",
                    "sections": [{"heading": "Household overview", "text": "The supplied synthetic sample contains 40 households. Mean income is 2400 units per month and mean wellbeing is 6.2 out of 10.",
                        "table": {"columns": ["Measure", "Value"], "rows": [["Households", "40"], ["Mean monthly income (units)", "2400"], ["Mean wellbeing (0–10)", "6.2"]]}}]}
                if self.mode != "default-pinned":
                    arguments["template_id"] = self.expected_id(fmt)
                return json.dumps({"tool": "export_document", "arguments": arguments})
        paths = [step["result"]["path"] for step in steps]
        assert all(Path(path).is_file() and path in messages.replace("\\/", "/") for path in paths)
        return json.dumps({"answer": "Created the Word report and PowerPoint presentation: " + " and ".join(paths)})


def archive_xml(path: Path) -> tuple[dict[str, bytes], str]:
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist() if name.endswith(".xml")}
    return entries, "\n".join(body.decode("utf-8") for body in entries.values())


def verify(fixture: TemplateFixture) -> dict:
    assert not fixture.errors, "; ".join(fixture.errors)
    assert fixture.normal_calls == 1 and fixture.action_calls == 3
    steps = read_steps(fixture.project)
    assert [(step["tool"], step["ok"]) for step in steps] == [("export_document", 1)] * 2
    results = {}
    for step in steps:
        fmt = step["result"]["format"]
        expected = fixture.expected_id(fmt)
        assert step["arguments"]["template_id"] == expected, "The host must pin omitted defaults and preserve explicit selections"
        applied = step["result"].get("template_applied")
        assert (applied is None) if expected == "none" else applied["id"] == expected
        if applied:
            assert applied["format"] == fmt and applied["limitations"], "Actual template reuse needs explicit limits"
        path = Path(step["result"]["path"])
        entries, text = archive_xml(path)
        assert OLD_BODY not in text and OLD_NOTES not in text, "Old source body/notes survived in the generated artifact"
        visible_text = "\n".join(part for body in entries.values() for part in ET.fromstring(body).itertext())
        assert "999999 households" not in visible_text, "Old sample findings survived in generated text"
        source = fixture.seed["sources"].get(expected)
        if fmt == "docx":
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            w = "{" + ns["w"] + "}"
            document = ET.fromstring(entries["word/document.xml"])
            assert "Synthetic household profile" in " ".join(document.itertext()) and "2400" in " ".join(document.itertext())
            if source:
                assert source["label"].upper() + "_BRAND_HEADER" in text
                assert source["label"].upper() + "_BRAND_FOOTER" in text
                styles = ET.fromstring(entries["word/styles.xml"])
                normal = styles.find("w:style[@w:styleId='Normal']", ns)
                assert normal.find("w:rPr/w:rFonts", ns).get(w + "ascii") == source["font"]
                assert normal.find("w:rPr/w:sz", ns).get(w + "val") == str(source["points"] * 2)
                for style_id in ["Title", "Heading1"]:
                    assert styles.find(f"w:style[@w:styleId='{style_id}']/w:rPr/w:color", ns).get(w + "val") == source["color"]
                section = document.find("w:body/w:sectPr", ns)
                page, margins = section.find("w:pgSz", ns), section.find("w:pgMar", ns)
                assert int(page.get(w + "w")) == source["word_width_twips"]
                assert int(page.get(w + "h")) == source["word_height_twips"]
                assert int(margins.get(w + "left")) == source["word_left_twips"]
            else:
                assert "_BRAND_HEADER" not in text and "_BRAND_FOOTER" not in text
        else:
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            size = ET.fromstring(entries["ppt/presentation.xml"]).find("p:sldSz", ns)
            slides = {name: body for name, body in entries.items() if name.startswith("ppt/slides/slide")}
            assert len(slides) >= 2
            for name, body in slides.items():
                slide_text = [e.text for e in ET.fromstring(body).findall(".//a:t", ns) if e.text and e.text.strip()]
                assert slide_text, f"Generated a blank slide: {name}"
            assert any("2400" in body.decode("utf-8") for body in slides.values())
            if source:
                assert (int(size.get("cx")), int(size.get("cy"))) == (source["ppt_width"], source["ppt_height"])
                original, _ = archive_xml(Path(source["path"]))
                for name, body in original.items():
                    if name.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")):
                        assert entries.get(name) == body, f"Template master/layout/theme changed: {name}"
            else:
                assert int(size.get("cx")) == int(13.333 * 914400) and int(size.get("cy")) == int(7.5 * 914400)
                assert "Research Master" not in text and "Research Content Layout" not in text
        results[fmt] = {"path": str(path), "template_id": expected, "template_applied": applied}
    for source in fixture.seed["sources"].values():
        assert hashlib.sha256(Path(source["path"]).read_bytes()).hexdigest() == source["sha256"], "An imported original changed"
    with sqlite3.connect(f"{(fixture.project / 'research.sqlite3').as_uri()}?mode=ro", uri=True) as connection:
        tasks = connection.execute("SELECT status, summary FROM research_tasks").fetchall()
    assert len(tasks) == 1 and tasks[0][0] == "completed" and "Created the Word report" in tasks[0][1]
    return {"status": "passed", "case": fixture.mode, "native_requests": fixture.normal_calls,
            "structured_requests": fixture.action_calls, "metadata_only_context": True,
            "snapshot_stable": True, "originals_unchanged": True, "old_body_and_notes_removed": True,
            "task_status": "completed", "exports": results}


def run_case(binary: Path, worker_python: Path, worker_source: Path, output: Path, mode: str, timeout: int) -> dict:
    output.mkdir(parents=True, exist_ok=False)
    support = output / "app-support"
    support.mkdir()
    (support / "worker-env").symlink_to(worker_python.parent.parent, target_is_directory=True)
    environment = {name: os.environ[name] for name in ["PATH", "HOME", "USER", "TMPDIR", "LANG", "LC_ALL", "TZ"] if name in os.environ}
    seed_env = {**environment, "PYTHONPATH": str(worker_source), "PYTHONDONTWRITEBYTECODE": "1"}
    seeded = subprocess.run([str(worker_python), "-c", SEED, str(output), OLD_BODY, OLD_NOTES],
                            env=seed_env, text=True, capture_output=True, timeout=60, check=False)
    assert seeded.returncode == 0, "Could not seed saved templates: " + seeded.stderr
    seed = json.loads(seeded.stdout)
    (output / "seed.json").write_text(json.dumps(seed, indent=2) + "\n")
    fixture = TemplateFixture(output, mode, seed)
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler_for(fixture))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    environment.update({"SQUIPI_TEST_DATA_DIR": str(support), "SQUIPI_TEST_OLLAMA_URL": f"http://127.0.0.1:{server.server_address[1]}",
                        "SQUIPI_OLLAMA_MODEL": "fixture-ollama", "PYTHONDONTWRITEBYTECODE": "1"})
    try:
        completed = subprocess.run([str(binary), "--agent", GOAL + "\n" + SELECTIONS[mode]], env=environment,
                                   text=True, capture_output=True, timeout=timeout, check=False)
        (output / "stdout.log").write_text(completed.stdout)
        (output / "stderr.log").write_text(completed.stderr)
        assert completed.returncode == 0, f"App exited {completed.returncode}; inspect {output}/stderr.log"
        assert "FINAL: Created the Word report and PowerPoint presentation" in completed.stdout, f"No completed artifacts; inspect {output}"
        result = verify(fixture)
        (output / "results.json").write_text(json.dumps(result, indent=2) + "\n")
        return result
    finally:
        server.shutdown()
        server.server_close()
        (output / "model-requests.json").write_text(json.dumps(fixture.requests, indent=2) + "\n")
        if fixture.errors:
            (output / "fixture-errors.json").write_text(json.dumps(fixture.errors, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--binary", required=True, type=Path)
    parser.add_argument("--worker-python", type=Path, default=Path.home() / "Library/Application Support/Squipi/worker-env/bin/python")
    parser.add_argument("--worker-source", type=Path, help="Defaults to the tested app's bundled worker")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--timeout", type=int, default=180)
    args = parser.parse_args()
    binary = args.binary.resolve()
    worker_python = Path(os.path.abspath(args.worker_python))
    worker_source = args.worker_source.resolve() if args.worker_source else binary.parent.parent / "Resources/worker"
    assert binary.is_file() and worker_python.is_file() and os.access(worker_python, os.X_OK)
    assert (worker_source / "squipi_worker/templates.py").is_file(), "Use a package with saved-template support"
    output = args.output_dir.resolve() if args.output_dir else Path(tempfile.mkdtemp(prefix="squipi-template-integration-"))
    if args.output_dir:
        output.mkdir(parents=True, exist_ok=False)
    results = [run_case(binary, worker_python, worker_source, output / mode, mode, args.timeout) for mode in SELECTIONS]
    (output / "results.json").write_text(json.dumps(results, indent=2) + "\n")
    print(json.dumps(results, indent=2))
    print(f"PASS: saved-template context, pinned defaults, explicit selection and opt-out → actual DOCX/PPTX; evidence: {output}")


if __name__ == "__main__":
    main()
